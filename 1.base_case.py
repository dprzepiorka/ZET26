from __future__ import annotations

import cmath
import math
import os
import sys
import time
from copy import deepcopy
from typing import Any, Dict, Iterable, List, Sequence, Tuple

try:
    from openpyxl import Workbook, load_workbook
except Exception:
    Workbook = None
    load_workbook = None


# =============================================================================
# PARAMETRY DOMYŚLNE
# =============================================================================

POWERFACTORY_PYTHON_PATH = r"C:\Program Files\DIgSILENT\PowerFactory 2026 SP1\Python\3.14"
USER = "KE"
PROJECT_NAME = "ELVTF_ZET_2026"

EXCEL_FILE = r"input_data.xlsx"
OUT_FILE = r"results\results_base_case.xlsx"
LOG_FILE = r"results\base_case.log"

TRANSFORMER_NAME = ""
TRANSFORMER_CLASS = "ElmTr2"
TRANSFORMER_LV_SIDE = "buslv"

TRANSFORMER_EXPORT_POSITIVE = True

VOLTAGE_MIN_PU = 0.90
TARGET_VOLTAGE_PU = 1.00
VOLTAGE_MAX_PU = 1.10
U_MIN_ALLOWED_OBJ = 0.95
U_MAX_ALLOWED_OBJ = 1.05

LOADING_MAX_PERCENT = 100.0
LINE_LOADING_MAX_PERCENT = 100.0

P_STORAGE_TOTAL_MAX_KW = 90.0
P_STORAGE_PHASE_MAX_KW = 30.0
ALLOW_STORAGE_CHARGE = True
ALLOW_STORAGE_DISCHARGE = True
KEEP_INITIAL_STORAGE_Q = True

EPS = 1e-9

PHASES = ("L1", "L2", "L3")
PF_PHASE = {"L1": "A", "L2": "B", "L3": "C"}

PHASE_ATTR_P_LOAD = {"L1": "plinir", "L2": "plinis", "L3": "plinit"}
PHASE_ATTR_Q_LOAD = {"L1": "qlinir", "L2": "qlinis", "L3": "qlinit"}

EXCEL_CACHE: Dict[str, List[Dict[str, Any]]] = {}
CONTROL_CONFIG: Dict[str, Any] = {}


# =============================================================================
# LOG
# =============================================================================

def log_line(msg: str) -> None:
    ts = time.strftime("%Y-%m-%d %H:%M:%S")
    line = f"[{ts}] {msg}"
    print(line)
    try:
        log_dir = os.path.dirname(LOG_FILE)
        if log_dir:
            os.makedirs(log_dir, exist_ok=True)
        with open(LOG_FILE, "a", encoding="utf-8") as fh:
            fh.write(line + "\n")
    except Exception:
        pass


# =============================================================================
# CONFIG
# =============================================================================

def parse_bool_pl(value: Any) -> bool:
    txt = str(value).strip().lower()
    return txt in {"1", "true", "yes", "y", "tak", "prawda"}


def parse_config_value(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, (int, float, bool)):
        return value

    txt = str(value).strip()
    if not txt:
        return ""

    low = txt.lower()
    if low in {"prawda", "true", "tak", "yes", "1"}:
        return True
    if low in {"fałsz", "falsz", "false", "nie", "no", "0"}:
        return False

    try:
        if "." in txt:
            return float(txt)
        return int(txt)
    except Exception:
        return txt


def load_control_config() -> Dict[str, Any]:
    rows = excel_sheet("ControlConfig")
    cfg: Dict[str, Any] = {}

    for row in rows:
        key = str(row.get("parameter", "")).strip()
        if not key:
            continue
        cfg[key] = parse_config_value(row.get("value"))

    return cfg


def cfg_float(name: str, default: float) -> float:
    val = CONTROL_CONFIG.get(name, default)
    try:
        return float(val)
    except Exception:
        return float(default)


def cfg_int(name: str, default: int) -> int:
    val = CONTROL_CONFIG.get(name, default)
    try:
        return int(float(val))
    except Exception:
        return int(default)


def cfg_bool(name: str, default: bool) -> bool:
    val = CONTROL_CONFIG.get(name, default)
    if isinstance(val, bool):
        return val
    return parse_bool_pl(val) if val is not None else default


def apply_runtime_config() -> None:
    global VOLTAGE_MIN_PU, TARGET_VOLTAGE_PU, VOLTAGE_MAX_PU
    global U_MIN_ALLOWED_OBJ, U_MAX_ALLOWED_OBJ
    global LOADING_MAX_PERCENT
    global P_STORAGE_TOTAL_MAX_KW, P_STORAGE_PHASE_MAX_KW
    global ALLOW_STORAGE_CHARGE, ALLOW_STORAGE_DISCHARGE, KEEP_INITIAL_STORAGE_Q

    VOLTAGE_MIN_PU = cfg_float("VOLTAGE_MIN_PU", VOLTAGE_MIN_PU)
    TARGET_VOLTAGE_PU = cfg_float("TARGET_VOLTAGE_PU", TARGET_VOLTAGE_PU)
    VOLTAGE_MAX_PU = cfg_float("VOLTAGE_MAX_PU", VOLTAGE_MAX_PU)
    U_MIN_ALLOWED_OBJ = cfg_float("U_MIN_ALLOWED_OBJ", U_MIN_ALLOWED_OBJ)
    U_MAX_ALLOWED_OBJ = cfg_float("U_MAX_ALLOWED_OBJ", U_MAX_ALLOWED_OBJ)
    LOADING_MAX_PERCENT = cfg_float("LOADING_MAX_PERCENT", LOADING_MAX_PERCENT)

    P_STORAGE_TOTAL_MAX_KW = cfg_float("P_STORAGE_TOTAL_MAX_KW", P_STORAGE_TOTAL_MAX_KW)
    P_STORAGE_PHASE_MAX_KW = cfg_float("P_STORAGE_PHASE_MAX_KW", P_STORAGE_PHASE_MAX_KW)
    ALLOW_STORAGE_CHARGE = cfg_bool("ALLOW_STORAGE_CHARGE", ALLOW_STORAGE_CHARGE)
    ALLOW_STORAGE_DISCHARGE = cfg_bool("ALLOW_STORAGE_DISCHARGE", ALLOW_STORAGE_DISCHARGE)
    KEEP_INITIAL_STORAGE_Q = cfg_bool("KEEP_INITIAL_STORAGE_Q", KEEP_INITIAL_STORAGE_Q)


# =============================================================================
# POWERFACTORY
# =============================================================================

def connect_powerfactory() -> Tuple[Any, Any]:
    if POWERFACTORY_PYTHON_PATH and POWERFACTORY_PYTHON_PATH not in sys.path:
        sys.path.append(POWERFACTORY_PYTHON_PATH)

    try:
        import powerfactory  # type: ignore
    except Exception as exc:
        raise RuntimeError(
            "Nie mogę zaimportować modułu powerfactory. "
            "Uruchom skrypt w środowisku DIgSILENT PowerFactory albo popraw POWERFACTORY_PYTHON_PATH."
        ) from exc

    app = powerfactory.GetApplicationExt(USER)
    if app is None:
        raise RuntimeError("PowerFactory nie zwrócił obiektu aplikacji.")

    if PROJECT_NAME:
        app.ActivateProject(PROJECT_NAME)

    ldf = app.GetFromStudyCase("ComLdf")
    if ldf is None:
        raise RuntimeError("Nie znaleziono ComLdf w aktywnym Study Case.")

    return app, ldf


def find_element(app: Any, name: str, pf_class: str) -> Any:
    if not name:
        return None

    try:
        objs = app.GetCalcRelevantObjects(f"{name}.{pf_class}")
        if objs:
            return objs[0]
    except Exception:
        pass

    try:
        for obj in app.GetCalcRelevantObjects(f"*.{pf_class}"):
            if getattr(obj, "loc_name", None) == name:
                return obj
    except Exception:
        pass

    return None


def get_attr(obj: Any, names: Sequence[str], default: Any = None) -> Any:
    for name in names:
        try:
            val = obj.GetAttribute(name)
            if val is not None:
                return val
        except Exception:
            continue
    return default


def get_float(obj: Any, names: Sequence[str], default: float = 0.0) -> float:
    val = get_attr(obj, names, default)
    try:
        return float(val)
    except Exception:
        return float(default)


def set_attr(obj: Any, names: Sequence[str], value: Any) -> bool:
    for name in names:
        try:
            obj.SetAttribute(name, value)
            return True
        except Exception:
            continue
    return False


def run_loadflow(ldf: Any) -> None:
    rc = ldf.Execute()
    if rc not in (0, None):
        raise RuntimeError(f"Rozpływ mocy nie zbieżny albo błąd ComLdf, kod: {rc}")


def get_transformer(app: Any) -> Any:
    tr = find_element(app, TRANSFORMER_NAME, TRANSFORMER_CLASS) if TRANSFORMER_NAME else None
    if tr is not None:
        return tr

    trafos = app.GetCalcRelevantObjects(f"*.{TRANSFORMER_CLASS}")
    if not trafos:
        raise RuntimeError(f"Nie znaleziono transformatora klasy {TRANSFORMER_CLASS}.")

    return trafos[0]


# =============================================================================
# EXCEL
# =============================================================================

def _sheet_rows_from_workbook(wb: Any, sheet: str) -> List[Dict[str, Any]]:
    if sheet not in wb.sheetnames:
        return []

    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    out: List[Dict[str, Any]] = []

    for raw in rows[1:]:
        row = {headers[i]: raw[i] for i in range(min(len(headers), len(raw))) if headers[i]}
        if any(v not in (None, "") for v in row.values()):
            out.append(row)

    return out


def load_excel_cache(path: str) -> Dict[str, List[Dict[str, Any]]]:
    if load_workbook is None:
        raise RuntimeError("Do czytania Excela potrzebny jest pakiet openpyxl.")
    if not os.path.exists(path):
        raise FileNotFoundError(f"Nie znaleziono pliku Excel: {path}")

    wb = None
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        cache: Dict[str, List[Dict[str, Any]]] = {}
        for sheet in wb.sheetnames:
            cache[sheet] = _sheet_rows_from_workbook(wb, sheet)
        return cache
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


def excel_sheet(sheet: str) -> List[Dict[str, Any]]:
    return deepcopy(EXCEL_CACHE.get(sheet, []))


# =============================================================================
# USTAWIANIE MODELU
# =============================================================================

def set_loads_from_excel(app: Any) -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []

    for row in excel_sheet("Loads"):
        name = str(row.get("name", "")).strip()
        if not name:
            continue

        elm = find_element(app, name, "ElmLod")
        if elm is None:
            log_line(f"[WARN] Loads: nie znaleziono ElmLod '{name}'")
            continue

        p1 = float(row.get("P1") or 0.0)
        p2 = float(row.get("P2") or 0.0)
        p3 = float(row.get("P3") or 0.0)
        q1 = float(row.get("Q1") or 0.0)
        q2 = float(row.get("Q2") or 0.0)
        q3 = float(row.get("Q3") or 0.0)

        set_attr(elm, ["plinir"], p1)
        set_attr(elm, ["plinis"], p2)
        set_attr(elm, ["plinit"], p3)
        set_attr(elm, ["qlinir"], q1)
        set_attr(elm, ["qlinis"], q2)
        set_attr(elm, ["qlinit"], q3)

        rows_out.append(
            {
                "name": name,
                "type": "load",
                "P1_kW": p1,
                "P2_kW": p2,
                "P3_kW": p3,
                "Q1_kvar": q1,
                "Q2_kvar": q2,
                "Q3_kvar": q3,
            }
        )

    return rows_out


def set_pv_from_excel_base_case(app: Any) -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []

    for row in excel_sheet("PV"):
        name = str(row.get("name", "")).strip()
        if not name:
            continue

        elm = find_element(app, name, "ElmPvsys")
        if elm is None:
            log_line(f"[WARN] PV: nie znaleziono ElmPvsys '{name}'")
            continue

        p = float(row.get("P") or 0.0)
        smax = float(row.get("Smax") or row.get("Smax_kVA") or row.get("Sn") or 0.0)

        set_attr(elm, ["pgini"], p)
        set_attr(elm, ["qgini", "qsetp"], 0.0)
        set_attr(elm, ["av_mode"], "constq")

        rows_out.append(
            {
                "name": name,
                "type": "pv",
                "P_kW": p,
                "Q_kvar": 0.0,
                "Smax_kVA": smax,
                "control_mode": "constq",
                "cosphi": 1.0,
            }
        )

    return rows_out


def set_other_generators_from_excel(app: Any) -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []

    for sheet, cls in [("Generators", "ElmSym"), ("StatGen", "ElmGenstat")]:
        for row in excel_sheet(sheet):
            name = str(row.get("name", "")).strip()
            if not name:
                continue

            elm = find_element(app, name, cls)
            if elm is None:
                log_line(f"[WARN] {sheet}: nie znaleziono {cls} '{name}'")
                continue

            p = float(row.get("P") or 0.0)
            q = float(row.get("Q") or 0.0)

            set_attr(elm, ["pgini"], p)
            set_attr(elm, ["qgini", "qsetp"], q)

            rows_out.append(
                {
                    "name": name,
                    "type": sheet,
                    "P_kW": p,
                    "Q_kvar": q,
                }
            )

    return rows_out


def load_storage_candidates() -> List[Dict[str, Any]]:
    rows = excel_sheet("StorageCandidates")
    return [
        {
            "node": str(r.get("node", "")).strip(),
            "elem_L1": str(r.get("elem_L1", r.get("elem_A", ""))).strip(),
            "elem_L2": str(r.get("elem_L2", r.get("elem_B", ""))).strip(),
            "elem_L3": str(r.get("elem_L3", r.get("elem_C", ""))).strip(),
            "Pmax_kW": min(float(r.get("Pmax_kW") or r.get("Pmax") or 0.0), P_STORAGE_PHASE_MAX_KW),
            "Smax_kVA": float(r.get("Smax_kVA") or r.get("Smax") or 0.0),
        }
        for r in rows
    ]


def zero_storage_rows() -> List[Dict[str, Any]]:
    return [
        {"phase": "L1", "step_rel": 0.0, "p_storage_kw": 0.0, "q_storage_kvar": 0.0},
        {"phase": "L2", "step_rel": 0.0, "p_storage_kw": 0.0, "q_storage_kvar": 0.0},
        {"phase": "L3", "step_rel": 0.0, "p_storage_kw": 0.0, "q_storage_kvar": 0.0},
    ]


def choose_storage_candidate(candidates: List[Dict[str, Any]], node_name: str = "") -> Dict[str, Any]:
    if not candidates:
        return {}

    if node_name:
        for cand in candidates:
            if cand.get("node") == node_name:
                return cand

    return candidates[0]


def apply_storage(app: Any, candidates: List[Dict[str, Any]], rows: List[Dict[str, Any]], node_name: str = "") -> List[Dict[str, Any]]:
    cand = choose_storage_candidate(candidates, node_name)
    applied_rows: List[Dict[str, Any]] = []

    for row in rows:
        ph = row["phase"]
        name = cand.get(f"elem_{ph}", "")
        p = float(row.get("p_storage_kw") or 0.0)
        q = float(row.get("q_storage_kvar") or 0.0)

        if (p > 0.0) and not ALLOW_STORAGE_CHARGE:
            p = 0.0
        if (p < 0.0) and not ALLOW_STORAGE_DISCHARGE:
            p = 0.0

        p = max(-P_STORAGE_PHASE_MAX_KW, min(P_STORAGE_PHASE_MAX_KW, p))

        applied_row = {
            "node": cand.get("node", ""),
            "element": name,
            "phase": ph,
            "p_storage_kw": p,
            "q_storage_kvar": q if KEEP_INITIAL_STORAGE_Q else 0.0,
            "Pmax_kW": float(cand.get("Pmax_kW") or 0.0),
            "Smax_kVA": float(cand.get("Smax_kVA") or 0.0),
        }

        if not name:
            log_line(f"[WARN] Storage: brak przypisania elementu dla fazy {ph}")
            applied_rows.append(applied_row)
            continue

        elm_lod = find_element(app, name, "ElmLod")
        if elm_lod is not None:
            set_attr(elm_lod, [PHASE_ATTR_P_LOAD[ph], "plini"], p)
            set_attr(elm_lod, [PHASE_ATTR_Q_LOAD[ph], "qlini"], applied_row["q_storage_kvar"])
            applied_rows.append(applied_row)
            continue

        elm_gen = (
            find_element(app, name, "ElmGenstat")
            or find_element(app, name, "ElmSym")
            or find_element(app, name, "ElmPvsys")
        )
        if elm_gen is not None:
            set_attr(elm_gen, ["pgini"], -p)
            set_attr(elm_gen, ["qgini", "qsetp"], -applied_row["q_storage_kvar"])
            applied_row["p_storage_kw_in_model"] = -p
            applied_row["q_storage_kvar_in_model"] = -applied_row["q_storage_kvar"]
            applied_rows.append(applied_row)
            continue

        log_line(f"[WARN] Storage: nie znaleziono elementu '{name}' dla fazy {ph}")
        applied_rows.append(applied_row)

    total_p = sum(float(r.get("p_storage_kw") or 0.0) for r in applied_rows)
    if abs(total_p) > P_STORAGE_TOTAL_MAX_KW + EPS:
        scale = P_STORAGE_TOTAL_MAX_KW / max(abs(total_p), EPS)
        for r in applied_rows:
            r["p_storage_kw"] = float(r["p_storage_kw"]) * scale

    return applied_rows


def set_base_case_model(app: Any) -> Dict[str, List[Dict[str, Any]]]:
    log_line("Ustawianie modelu bazowego z Excela...")

    loads_data = set_loads_from_excel(app)
    pv_data = set_pv_from_excel_base_case(app)
    other_gen_data = set_other_generators_from_excel(app)

    candidates = load_storage_candidates()
    storage_data = apply_storage(app, candidates, zero_storage_rows())

    return {
        "loads_input": loads_data,
        "pv_input": pv_data,
        "other_generators_input": other_gen_data,
        "storage_input": storage_data,
    }


# =============================================================================
# ODCZYT WYNIKÓW
# =============================================================================

def collect_node_voltages(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []

    for bus in app.GetCalcRelevantObjects("*.ElmTerm"):
        row = {"node": getattr(bus, "loc_name", "")}

        for ph in PHASES:
            pf_ph = PF_PHASE[ph]
            row[f"U_{ph}_pu"] = get_float(bus, [f"m:u:{pf_ph}"], math.nan)
            row[f"angle_{ph}_deg"] = get_float(bus, [f"m:phiu:{pf_ph}"], math.nan)

        row["kU2_percent"] = get_float(bus, ["m:ubfac"], math.nan)
        rows.append(row)

    return rows


def collect_transformer_results(tr: Any) -> Dict[str, Any]:
    row = {"transformer": getattr(tr, "loc_name", "")}

    i_complex = []

    for ph in PHASES:
        pf_ph = PF_PHASE[ph]

        i_ka = get_float(tr, [f"m:I:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:I:{pf_ph}"], 0.0)
        phi_i = get_float(tr, [f"m:phii:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:phii:{pf_ph}"], math.nan)

        p = get_attr(
            tr,
            [f"m:P:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:P:{pf_ph}", f"c:P:{pf_ph}"],
            None,
        )
        q = get_attr(
            tr,
            [f"m:Q:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:Q:{pf_ph}", f"c:Q:{pf_ph}"],
            None,
        )

        i_a = float(i_ka or 0.0) * 1000.0

        row[f"I_tr_{ph}_A"] = i_a
        row[f"angle_I_tr_{ph}_deg"] = phi_i
        row[f"P_tr_{ph}_kW"] = None if p is None else float(p)
        row[f"Q_tr_{ph}_kvar"] = None if q is None else float(q)

        i_complex.append(angle_deg_to_complex(i_a, float(phi_i) if math.isfinite(float(phi_i)) else 0.0))

    row["loading_percent"] = get_float(tr, ["c:loading"], math.nan)
    row["I_neutral_A"] = abs(sum(i_complex))
    return row


def collect_branch_losses(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []

    for line in app.GetCalcRelevantObjects("*.ElmLne"):
        p_loss = get_attr(line, ["c:LossP", "c:Ploss", "c:Losses"], 0.0)
        q_loss = get_attr(line, ["c:LossQ", "c:Qloss"], 0.0)

        rows.append(
            {
                "branch": getattr(line, "loc_name", ""),
                "P_loss_kW": float(p_loss or 0.0),
                "Q_loss_kvar": float(q_loss or 0.0),
                "loading_percent": get_float(line, ["c:loading"], math.nan),
            }
        )

    return rows


def collect_pv_setpoints(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []

    for pv in app.GetCalcRelevantObjects("*.ElmPvsys"):
        p_set = get_float(pv, ["pgini"], 0.0)
        q_set = get_float(pv, ["qgini", "qsetp"], 0.0)
        q_res = get_float(pv, ["m:Q:bus1", "m:Q:buslv", "m:Q", "c:Q"], q_set)
        smax = get_float(pv, ["sgn", "sn", "snom"], 0.0)

        rows.append(
            {
                "name": getattr(pv, "loc_name", ""),
                "P_kW": p_set,
                "Q_kvar": q_res,
                "Q_set_kvar": q_set,
                "S_apparent_kVA": math.sqrt(p_set * p_set + q_res * q_res),
                "Smax_kVA": smax,
                "av_mode": get_attr(pv, ["av_mode"], ""),
            }
        )

    return rows


# =============================================================================
# POMOCNICZE OBLICZENIA
# =============================================================================

def angle_deg_to_complex(magnitude: float, angle_deg: float) -> complex:
    if not math.isfinite(magnitude):
        return complex(math.nan, math.nan)
    if not math.isfinite(angle_deg):
        return complex(magnitude, 0.0)
    return cmath.rect(magnitude, math.radians(angle_deg))


def calc_symmetrical_components(v_a: complex, v_b: complex, v_c: complex) -> Tuple[complex, complex, complex]:
    a = complex(-0.5, math.sqrt(3.0) / 2.0)
    a2 = a * a

    u0 = (v_a + v_b + v_c) / 3.0
    u1 = (v_a + a * v_b + a2 * v_c) / 3.0
    u2 = (v_a + a2 * v_b + a * v_c) / 3.0
    return u0, u1, u2


def rms_deviation(values: List[float], target: float) -> float:
    if not values:
        return math.nan
    return math.sqrt(sum((v - target) ** 2 for v in values) / len(values))


def export_import_from_transformer_phases(p_tr: List[float]) -> Tuple[float, float]:
    export_pos = sum(max(0.0, p) for p in p_tr)
    export_neg = sum(max(0.0, -p) for p in p_tr)

    if TRANSFORMER_EXPORT_POSITIVE:
        return export_pos, export_neg
    return export_neg, export_pos


# =============================================================================
# WSKAŹNIKI
# =============================================================================

def calculate_indicators(raw: Dict[str, List[Dict[str, Any]]]) -> Dict[str, Any]:
    voltages: List[float] = []
    phase_spreads: List[float] = []
    ku2_vals: List[float] = []

    alpha0_vals: List[float] = []
    alpha2_vals: List[float] = []
    u0_abs_vals: List[float] = []
    u1_abs_vals: List[float] = []
    u2_abs_vals: List[float] = []

    node_sequence_rows: List[Dict[str, Any]] = []

    for row in raw["node_voltages"]:
        vals = []
        for ph in PHASES:
            try:
                v = float(row[f"U_{ph}_pu"])
                if math.isfinite(v):
                    vals.append(v)
                    voltages.append(v)
            except Exception:
                pass

        if vals:
            phase_spreads.append(max(vals) - min(vals))

        try:
            k = float(row.get("kU2_percent"))
            if math.isfinite(k):
                ku2_vals.append(k)
        except Exception:
            pass

        try:
            va = angle_deg_to_complex(float(row["U_L1_pu"]), float(row["angle_L1_deg"]))
            vb = angle_deg_to_complex(float(row["U_L2_pu"]), float(row["angle_L2_deg"]))
            vc = angle_deg_to_complex(float(row["U_L3_pu"]), float(row["angle_L3_deg"]))

            u0, u1, u2 = calc_symmetrical_components(va, vb, vc)

            u0_abs = abs(u0)
            u1_abs = abs(u1)
            u2_abs = abs(u2)

            alpha0 = u0_abs / u1_abs if u1_abs > EPS else math.nan
            alpha2 = u2_abs / u1_abs if u1_abs > EPS else math.nan

            if math.isfinite(alpha0):
                alpha0_vals.append(alpha0)
            if math.isfinite(alpha2):
                alpha2_vals.append(alpha2)

            if math.isfinite(u0_abs):
                u0_abs_vals.append(u0_abs)
            if math.isfinite(u1_abs):
                u1_abs_vals.append(u1_abs)
            if math.isfinite(u2_abs):
                u2_abs_vals.append(u2_abs)

            node_sequence_rows.append(
                {
                    "node": row.get("node", ""),
                    "U0_abs_pu": u0_abs,
                    "U1_abs_pu": u1_abs,
                    "U2_abs_pu": u2_abs,
                    "alpha0": alpha0,
                    "alpha2": alpha2,
                }
            )
        except Exception:
            node_sequence_rows.append(
                {
                    "node": row.get("node", ""),
                    "U0_abs_pu": math.nan,
                    "U1_abs_pu": math.nan,
                    "U2_abs_pu": math.nan,
                    "alpha0": math.nan,
                    "alpha2": math.nan,
                }
            )

    tr = raw["transformer_phase_results"][0] if raw["transformer_phase_results"] else {}

    i_vals = [float(tr.get(f"I_tr_{ph}_A") or 0.0) for ph in PHASES]
    i_avg = sum(i_vals) / 3.0 if i_vals else 0.0
    i_unb = max(abs(i - i_avg) for i in i_vals) / i_avg * 100.0 if i_avg > EPS else 0.0

    fcelu_2 = float(tr.get("I_neutral_A") or 0.0)

    p_tr = [float(tr.get(f"P_tr_{ph}_kW") or 0.0) for ph in PHASES]
    q_tr = [float(tr.get(f"Q_tr_{ph}_kvar") or 0.0) for ph in PHASES]
    p_export_total, p_import_total = export_import_from_transformer_phases(p_tr)

    p_loss = sum(float(row.get("P_loss_kW") or 0.0) for row in raw["branch_losses"])
    q_loss = sum(float(row.get("Q_loss_kvar") or 0.0) for row in raw["branch_losses"])

    pv_rows = raw.get("pv_setpoints", [])
    storage_rows = raw.get("storage_setpoints", [])

    p_pv_total = sum(float(row.get("P_kW") or 0.0) for row in pv_rows)
    q_pv_total = sum(float(row.get("Q_kvar") or 0.0) for row in pv_rows)
    q_pv_abs_sum = sum(abs(float(row.get("Q_kvar") or 0.0)) for row in pv_rows)

    p_storage_L1 = float(storage_rows[0].get("p_storage_kw") or 0.0) if len(storage_rows) > 0 else 0.0
    p_storage_L2 = float(storage_rows[1].get("p_storage_kw") or 0.0) if len(storage_rows) > 1 else 0.0
    p_storage_L3 = float(storage_rows[2].get("p_storage_kw") or 0.0) if len(storage_rows) > 2 else 0.0
    p_storage_total = p_storage_L1 + p_storage_L2 + p_storage_L3
    q_storage_total = sum(float(row.get("q_storage_kvar") or 0.0) for row in storage_rows)

    udev_mean_1_00 = sum(abs(v - 1.0) for v in voltages) / len(voltages) if voltages else math.nan
    udev_rms_1_00 = rms_deviation(voltages, 1.00)
    udev_rms_1_05 = rms_deviation(voltages, 1.05)

    alpha0_mean = sum(alpha0_vals) / len(alpha0_vals) if alpha0_vals else math.nan
    alpha0_max = max(alpha0_vals) if alpha0_vals else math.nan
    alpha2_mean = sum(alpha2_vals) / len(alpha2_vals) if alpha2_vals else math.nan
    alpha2_max = max(alpha2_vals) if alpha2_vals else math.nan

    fcelu_3_voltage_component = 0.04 * (udev_rms_1_05 if math.isfinite(udev_rms_1_05) else 0.0)
    fcelu_3_alpha2_component = 0.58 * (alpha2_mean if math.isfinite(alpha2_mean) else 0.0)
    fcelu_3_alpha0_component = 0.38 * (alpha0_mean if math.isfinite(alpha0_mean) else 0.0)
    fcelu_3 = fcelu_3_voltage_component + fcelu_3_alpha2_component + fcelu_3_alpha0_component

    violations: List[str] = []

    for row in raw["node_voltages"]:
        for ph in PHASES:
            try:
                u = float(row[f"U_{ph}_pu"])
                if u < VOLTAGE_MIN_PU or u > VOLTAGE_MAX_PU:
                    violations.append(f"U:{row.get('node')}:{ph}:{u:.4f}")
            except Exception:
                pass

    try:
        tr_loading = float(tr.get("loading_percent"))
        if math.isfinite(tr_loading) and tr_loading > LOADING_MAX_PERCENT:
            violations.append(f"TR_loading:{tr_loading:.2f}")
    except Exception:
        pass

    for row in raw["branch_losses"]:
        try:
            line_loading = float(row.get("loading_percent"))
            if math.isfinite(line_loading) and line_loading > LINE_LOADING_MAX_PERCENT:
                violations.append(f"LINE_loading:{row.get('branch')}:{line_loading:.2f}")
        except Exception:
            pass

    indicators = {
        "Umax_pu": max(voltages) if voltages else math.nan,
        "Umin_pu": min(voltages) if voltages else math.nan,
        "Udev_mean_pu": udev_mean_1_00,
        "Udev_max_pu": max(abs(v - 1.0) for v in voltages) if voltages else math.nan,
        "Fcelu_1_Udev_rms_1_00": udev_rms_1_00,
        "Fcelu_1_Udev_rms_1_05": udev_rms_1_05,
        "dU_phase_max_pu": max(phase_spreads) if phase_spreads else math.nan,
        "kU2_max_percent": max(ku2_vals) if ku2_vals else math.nan,
        "alpha0_mean": alpha0_mean,
        "alpha0_max": alpha0_max,
        "alpha2_mean": alpha2_mean,
        "alpha2_max": alpha2_max,
        "Fcelu_3_weighted": fcelu_3,

        "I_tr_L1_A": i_vals[0],
        "I_tr_L2_A": i_vals[1],
        "I_tr_L3_A": i_vals[2],
        "I_unbalance_tr_percent": i_unb,
        "I_neutral_A": fcelu_2,

        "P_tr_L1_kW": p_tr[0],
        "P_tr_L2_kW": p_tr[1],
        "P_tr_L3_kW": p_tr[2],
        "P_export_total_kW": p_export_total,
        "P_import_total_kW": p_import_total,
        "P_loss_total_kW": p_loss,
        "Q_pv_abs_sum_kvar": q_pv_abs_sum,

        "P_storage_L1_kW": p_storage_L1,
        "P_storage_L2_kW": p_storage_L2,
        "P_storage_L3_kW": p_storage_L3,
        "P_storage_total_kW": p_storage_total,

        "Q_loss_total_kvar": q_loss,
        "Q_pv_total_kvar": q_pv_total,
        "P_pv_total_kW": p_pv_total,
        "Q_storage_total_kvar": q_storage_total,
        "Fcelu_2_A": fcelu_2,

        "objective_profile": "base_case",
        "J_used_total": "",
        "J_used_component_1": "",
        "J_used_component_2": "",
        "J_used_component_3": "",
        "J_used_component_4": "",
        "J_used_component_5": "",

        "constraint_violations": ";".join(violations),
    }

    raw["node_sequence_components"] = node_sequence_rows
    return indicators


def collect_results(app: Any, tr: Any, setup_data: Dict[str, List[Dict[str, Any]]]) -> Dict[str, List[Dict[str, Any]]]:
    node_voltages = collect_node_voltages(app)
    transformer_results = [collect_transformer_results(tr)]
    branch_losses = collect_branch_losses(app)
    pv_setpoints = collect_pv_setpoints(app)
    storage_input = deepcopy(setup_data["storage_input"])

    raw = {
        "node_voltages": node_voltages,
        "transformer_phase_results": transformer_results,
        "branch_losses": branch_losses,
        "pv_setpoints": pv_setpoints,
        "storage_setpoints": storage_input,
    }

    indicators = calculate_indicators(raw)

    return {
        "control_config": [{"parameter": k, "value": v} for k, v in CONTROL_CONFIG.items()],
        "loads_input": setup_data["loads_input"],
        "pv_input": setup_data["pv_input"],
        "other_generators_input": setup_data["other_generators_input"],
        "storage_input": setup_data["storage_input"],
        "node_voltages": node_voltages,
        "node_sequence_components": raw.get("node_sequence_components", []),
        "transformer_phase_results": transformer_results,
        "pv_setpoints": pv_setpoints,
        "storage_setpoints": storage_input,
        "branch_losses": branch_losses,
        "indicators": [indicators],
    }


# =============================================================================
# EXCEL EXPORT
# =============================================================================

def safe_sheet_name(name: str) -> str:
    bad = "[]:*?/\\"
    for ch in bad:
        name = name.replace(ch, "_")
    return name[:31]


def write_rows(ws: Any, rows: Iterable[Dict[str, Any]]) -> None:
    rows = list(rows)

    if not rows:
        ws.append(["empty"])
        return

    headers: List[str] = []
    for row in rows:
        for key in row.keys():
            if key not in headers:
                headers.append(key)

    ws.append(headers)

    for row in rows:
        ws.append([row.get(h, "") for h in headers])


def export_to_excel(all_tables: Dict[str, List[Dict[str, Any]]], out_file: str) -> None:
    if Workbook is None:
        raise RuntimeError("Do zapisu wyników do Excela potrzebny jest pakiet openpyxl.")

    out_dir = os.path.dirname(out_file)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, rows in all_tables.items():
        ws = wb.create_sheet(safe_sheet_name(sheet_name))
        write_rows(ws, rows)

    wb.save(out_file)


# =============================================================================
# MAIN
# =============================================================================

def run_base_case() -> None:
    global EXCEL_CACHE, CONTROL_CONFIG

    try:
        log_dir = os.path.dirname(LOG_FILE)
        if log_dir:
            os.makedirs(log_dir, exist_ok=True)
        with open(LOG_FILE, "w", encoding="utf-8") as fh:
            fh.write("=== BASE CASE LOG START ===\n")
    except Exception:
        pass

    log_line("Start base_case.py")

    EXCEL_CACHE = load_excel_cache(EXCEL_FILE)
    CONTROL_CONFIG = load_control_config()
    apply_runtime_config()

    log_line(f"Wczytano Excel: {EXCEL_FILE}")
    log_line(f"Arkusze: {', '.join(sorted(EXCEL_CACHE.keys()))}")
    log_line(
        f"Config: Vmin={VOLTAGE_MIN_PU}, Vtarget={TARGET_VOLTAGE_PU}, Vmax={VOLTAGE_MAX_PU}, "
        f"Uobj=[{U_MIN_ALLOWED_OBJ}, {U_MAX_ALLOWED_OBJ}], LoadingMax={LOADING_MAX_PERCENT}"
    )

    app, ldf = connect_powerfactory()
    tr = get_transformer(app)

    setup_data = set_base_case_model(app)

    log_line("Uruchamiam bazowy rozpływ mocy...")
    run_loadflow(ldf)

    result_tables = collect_results(app, tr, setup_data)
    export_to_excel(result_tables, OUT_FILE)

    ind = result_tables["indicators"][0]
    log_line(
        "Zakończono. "
        f"Umax={ind['Umax_pu']:.4f} pu, "
        f"Umin={ind['Umin_pu']:.4f} pu, "
        f"eksport={ind['P_export_total_kW']:.2f} kW, "
        f"import={ind['P_import_total_kW']:.2f} kW, "
        f"Fcelu_1(1.05)={ind['Fcelu_1_Udev_rms_1_05']:.6f}, "
        f"Fcelu_2={ind['Fcelu_2_A']:.4f} A, "
        f"Fcelu_3={ind['Fcelu_3_weighted']:.6f}"
    )
    log_line(f"Zapisano wyniki do: {OUT_FILE}")


if __name__ == "__main__":
    run_base_case()