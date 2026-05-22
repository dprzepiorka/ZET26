from __future__ import annotations

import cmath
import math
import os
import sys
import time
from copy import deepcopy
from typing import Any, Dict, Iterable, List, Sequence, Tuple

try:
    from openpyxl import Workbook, load_workbook
except Exception:
    Workbook = None
    load_workbook = None


# =============================================================================
# PARAMETRY DOMYŚLNE
# =============================================================================

POWERFACTORY_PYTHON_PATH = r"C:\Program Files\DIgSILENT\PowerFactory 2026 SP1\Python\3.14"
USER = "KE"
PROJECT_NAME = "ELVTF_ZET_2026"

EXCEL_FILE = r"input_data.xlsx"
OUT_FILE = r"results\results_local_transformer_flow_control.xlsx"
LOG_FILE = r"results\local_transformer_flow_control.log"

TRANSFORMER_NAME = ""
TRANSFORMER_CLASS = "ElmTr2"
TRANSFORMER_LV_SIDE = "buslv"

TRANSFORMER_EXPORT_POSITIVE = True

N_ITER = 5
MAX_RESCUE_ATTEMPTS = 10

ASYMMETRY_EXPORT_TOLERANCE_KW = 0.20
ASYMMETRY_ENABLE = True
ASYMMETRY_MAX_PASSES = 10

VOLTAGE_MIN_PU = 0.90
TARGET_VOLTAGE_PU = 1.00
VOLTAGE_MAX_PU = 1.10

U_MIN_ALLOWED_OBJ = 0.9
U_MAX_ALLOWED_OBJ = 1.1
U_TOL_PU = 0.01

P_EXPORT_REF_KW = 250 * 0.95

K_U = 100.0
K_L = 100.0
K_T = 100.0

LOADING_MAX_PERCENT = 100.0
LINE_LOADING_MAX_PERCENT = 100.0

P_STORAGE_TOTAL_MAX_KW = 90.0
P_STORAGE_PHASE_MAX_KW = 30.0
ALLOW_STORAGE_CHARGE = True
ALLOW_STORAGE_DISCHARGE = True
KEEP_INITIAL_STORAGE_Q = True

EPS = 1e-9

PHASES = ("L1", "L2", "L3")
PF_PHASE = {"L1": "A", "L2": "B", "L3": "C"}

PHASE_ATTR_P_LOAD = {"L1": "plinir", "L2": "plinis", "L3": "plinit"}
PHASE_ATTR_Q_LOAD = {"L1": "qlinir", "L2": "qlinis", "L3": "qlinit"}

EXCEL_CACHE: Dict[str, List[Dict[str, Any]]] = {}
CONTROL_CONFIG: Dict[str, Any] = {}


# =============================================================================
# LOG
# =============================================================================

def log_line(msg: str) -> None:
    ts = time.strftime("%Y-%m-%d %H:%M:%S")
    line = f"[{ts}] {msg}"
    print(line)
    try:
        log_dir = os.path.dirname(LOG_FILE)
        if log_dir:
            os.makedirs(log_dir, exist_ok=True)
        with open(LOG_FILE, "a", encoding="utf-8") as fh:
            fh.write(line + "\n")
    except Exception:
        pass


# =============================================================================
# CONFIG
# =============================================================================

def parse_bool_pl(value: Any) -> bool:
    txt = str(value).strip().lower()
    return txt in {"1", "true", "yes", "y", "tak", "prawda"}


def parse_config_value(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, (int, float, bool)):
        return value

    txt = str(value).strip()
    if not txt:
        return ""

    low = txt.lower()
    if low in {"prawda", "true", "tak", "yes", "1"}:
        return True
    if low in {"fałsz", "falsz", "false", "nie", "no", "0"}:
        return False

    try:
        if "." in txt:
            return float(txt)
        return int(txt)
    except Exception:
        return txt


def load_control_config() -> Dict[str, Any]:
    rows = excel_sheet("ControlConfig")
    cfg: Dict[str, Any] = {}

    for row in rows:
        key = str(row.get("parameter", "")).strip()
        if not key:
            continue
        cfg[key] = parse_config_value(row.get("value"))

    return cfg


def cfg_float(name: str, default: float) -> float:
    val = CONTROL_CONFIG.get(name, default)
    try:
        return float(val)
    except Exception:
        return float(default)


def cfg_int(name: str, default: int) -> int:
    val = CONTROL_CONFIG.get(name, default)
    try:
        return int(float(val))
    except Exception:
        return int(default)


def cfg_bool(name: str, default: bool) -> bool:
    val = CONTROL_CONFIG.get(name, default)
    if isinstance(val, bool):
        return val
    return parse_bool_pl(val) if val is not None else default


def apply_runtime_config() -> None:
    global VOLTAGE_MIN_PU, TARGET_VOLTAGE_PU, VOLTAGE_MAX_PU
    global U_MIN_ALLOWED_OBJ, U_MAX_ALLOWED_OBJ, U_TOL_PU
    global P_EXPORT_REF_KW
    global K_U, K_L, K_T
    global LOADING_MAX_PERCENT
    global P_STORAGE_TOTAL_MAX_KW, P_STORAGE_PHASE_MAX_KW
    global ALLOW_STORAGE_CHARGE, ALLOW_STORAGE_DISCHARGE, KEEP_INITIAL_STORAGE_Q
    global N_ITER, ASYMMETRY_EXPORT_TOLERANCE_KW, ASYMMETRY_ENABLE, ASYMMETRY_MAX_PASSES

    VOLTAGE_MIN_PU = cfg_float("VOLTAGE_MIN_PU", VOLTAGE_MIN_PU)
    TARGET_VOLTAGE_PU = cfg_float("UREF_PU", cfg_float("TARGET_VOLTAGE_PU", TARGET_VOLTAGE_PU))
    VOLTAGE_MAX_PU = cfg_float("VOLTAGE_MAX_PU", VOLTAGE_MAX_PU)

    U_MIN_ALLOWED_OBJ = cfg_float("U_MIN_ALLOWED_PU", cfg_float("U_MIN_ALLOWED_OBJ", U_MIN_ALLOWED_OBJ))
    U_MAX_ALLOWED_OBJ = cfg_float("U_MAX_ALLOWED_PU", cfg_float("U_MAX_ALLOWED_OBJ", U_MAX_ALLOWED_OBJ))
    U_TOL_PU = cfg_float("U_TOL_PU", U_TOL_PU)

    P_EXPORT_REF_KW = cfg_float("P_EXPORT_REF_KW", P_EXPORT_REF_KW)

    K_U = cfg_float("K_U", K_U)
    K_L = cfg_float("K_L", K_L)
    K_T = cfg_float("K_T", K_T)

    LOADING_MAX_PERCENT = cfg_float("LOADING_MAX_PERCENT", LOADING_MAX_PERCENT)

    P_STORAGE_TOTAL_MAX_KW = cfg_float("P_STORAGE_TOTAL_MAX_KW", P_STORAGE_TOTAL_MAX_KW)
    P_STORAGE_PHASE_MAX_KW = cfg_float("P_STORAGE_PHASE_MAX_KW", P_STORAGE_PHASE_MAX_KW)
    ALLOW_STORAGE_CHARGE = cfg_bool("ALLOW_STORAGE_CHARGE", ALLOW_STORAGE_CHARGE)
    ALLOW_STORAGE_DISCHARGE = cfg_bool("ALLOW_STORAGE_DISCHARGE", ALLOW_STORAGE_DISCHARGE)
    KEEP_INITIAL_STORAGE_Q = cfg_bool("KEEP_INITIAL_STORAGE_Q", KEEP_INITIAL_STORAGE_Q)

    N_ITER = cfg_int("LOCAL_N_ITER", cfg_int("PSO_N_ITER", N_ITER))
    ASYMMETRY_EXPORT_TOLERANCE_KW = cfg_float("ASYMMETRY_EXPORT_TOLERANCE_KW", ASYMMETRY_EXPORT_TOLERANCE_KW)
    ASYMMETRY_ENABLE = cfg_bool("ASYMMETRY_ENABLE", ASYMMETRY_ENABLE)
    ASYMMETRY_MAX_PASSES = cfg_int("ASYMMETRY_MAX_PASSES", ASYMMETRY_MAX_PASSES)


# =============================================================================
# POWERFACTORY
# =============================================================================

def connect_powerfactory() -> Tuple[Any, Any]:
    if POWERFACTORY_PYTHON_PATH and POWERFACTORY_PYTHON_PATH not in sys.path:
        sys.path.append(POWERFACTORY_PYTHON_PATH)

    try:
        import powerfactory  # type: ignore
    except Exception as exc:
        raise RuntimeError(
            "Nie mogę zaimportować modułu powerfactory. "
            "Uruchom skrypt w środowisku DIgSILENT PowerFactory albo popraw POWERFACTORY_PYTHON_PATH."
        ) from exc

    app = powerfactory.GetApplicationExt(USER)
    if app is None:
        raise RuntimeError("PowerFactory nie zwrócił obiektu aplikacji.")

    if PROJECT_NAME:
        app.ActivateProject(PROJECT_NAME)

    ldf = app.GetFromStudyCase("ComLdf")
    if ldf is None:
        raise RuntimeError("Nie znaleziono ComLdf w aktywnym Study Case.")

    return app, ldf


def find_element(app: Any, name: str, pf_class: str) -> Any:
    if not name:
        return None

    try:
        objs = app.GetCalcRelevantObjects(f"{name}.{pf_class}")
        if objs:
            return objs[0]
    except Exception:
        pass

    try:
        for obj in app.GetCalcRelevantObjects(f"*.{pf_class}"):
            if getattr(obj, "loc_name", None) == name:
                return obj
    except Exception:
        pass

    return None


def get_attr(obj: Any, names: Sequence[str], default: Any = None) -> Any:
    for name in names:
        try:
            val = obj.GetAttribute(name)
            if val is not None:
                return val
        except Exception:
            continue
    return default


def get_float(obj: Any, names: Sequence[str], default: float = 0.0) -> float:
    val = get_attr(obj, names, default)
    try:
        return float(val)
    except Exception:
        return float(default)


def set_attr(obj: Any, names: Sequence[str], value: Any) -> bool:
    for name in names:
        try:
            obj.SetAttribute(name, value)
            return True
        except Exception:
            continue
    return False


def try_run_loadflow(ldf: Any) -> Tuple[bool, Any]:
    try:
        rc = ldf.Execute()
        return rc in (0, None), rc
    except Exception as exc:
        return False, str(exc)


def get_transformer(app: Any) -> Any:
    tr = find_element(app, TRANSFORMER_NAME, TRANSFORMER_CLASS) if TRANSFORMER_NAME else None
    if tr is not None:
        return tr

    trafos = app.GetCalcRelevantObjects(f"*.{TRANSFORMER_CLASS}")
    if not trafos:
        raise RuntimeError(f"Nie znaleziono transformatora klasy {TRANSFORMER_CLASS}.")
    return trafos[0]


# =============================================================================
# EXCEL
# =============================================================================

def _sheet_rows_from_workbook(wb: Any, sheet: str) -> List[Dict[str, Any]]:
    if sheet not in wb.sheetnames:
        return []

    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    out: List[Dict[str, Any]] = []

    for raw in rows[1:]:
        row = {headers[i]: raw[i] for i in range(min(len(headers), len(raw))) if headers[i]}
        if any(v not in (None, "") for v in row.values()):
            out.append(row)

    return out


def load_excel_cache(path: str) -> Dict[str, List[Dict[str, Any]]]:
    if load_workbook is None:
        raise RuntimeError("Do czytania Excela potrzebny jest pakiet openpyxl.")
    if not os.path.exists(path):
        raise FileNotFoundError(f"Nie znaleziono pliku Excel: {path}")

    wb = None
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        cache: Dict[str, List[Dict[str, Any]]] = {}
        for sheet in wb.sheetnames:
            cache[sheet] = _sheet_rows_from_workbook(wb, sheet)
        return cache
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


def excel_sheet(sheet: str) -> List[Dict[str, Any]]:
    return deepcopy(EXCEL_CACHE.get(sheet, []))


# =============================================================================
# MODEL Z EXCELA
# =============================================================================

def set_loads_from_excel(app: Any) -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []

    for row in excel_sheet("Loads"):
        name_raw = row.get("name", "")
        if name_raw in (None, ""):
            continue
        name = str(name_raw).strip()
        if not name or name.lower() in {"none", "atribute", "attribute"}:
            continue

        elm = find_element(app, name, "ElmLod")
        if elm is None:
            log_line(f"[WARN] Loads: nie znaleziono ElmLod '{name}'")
            continue

        p1 = float(row.get("P1") or 0.0)
        p2 = float(row.get("P2") or 0.0)
        p3 = float(row.get("P3") or 0.0)
        q1 = float(row.get("Q1") or 0.0)
        q2 = float(row.get("Q2") or 0.0)
        q3 = float(row.get("Q3") or 0.0)

        set_attr(elm, ["plinir"], p1)
        set_attr(elm, ["plinis"], p2)
        set_attr(elm, ["plinit"], p3)
        set_attr(elm, ["qlinir"], q1)
        set_attr(elm, ["qlinis"], q2)
        set_attr(elm, ["qlinit"], q3)

        rows_out.append(
            {
                "name": name,
                "type": "load",
                "P1_kW": p1,
                "P2_kW": p2,
                "P3_kW": p3,
                "Q1_kvar": q1,
                "Q2_kvar": q2,
                "Q3_kvar": q3,
            }
        )

    return rows_out


def set_pv_from_excel(app: Any) -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []

    for row in excel_sheet("PV"):
        name_raw = row.get("name", "")
        if name_raw in (None, ""):
            continue
        name = str(name_raw).strip()
        if not name or name.lower() in {"none", "atribute", "attribute"}:
            continue

        elm = find_element(app, name, "ElmPvsys")
        if elm is None:
            log_line(f"[WARN] PV: nie znaleziono ElmPvsys '{name}'")
            continue

        p = float(row.get("P") or 0.0)
        smax = float(row.get("Smax") or row.get("Smax_kVA") or row.get("Sn") or 0.0)

        set_attr(elm, ["pgini"], p)
        set_attr(elm, ["qgini", "qsetp"], 0.0)

        rows_out.append(
            {
                "name": name,
                "type": "pv",
                "P_kW": p,
                "Q_kvar": 0.0,
                "Smax_kVA": smax,
                "control_mode": "to_be_set_qvchar",
            }
        )

    return rows_out


def set_other_generators_from_excel(app: Any) -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []

    for sheet, cls in [("Generators", "ElmSym"), ("StatGen", "ElmGenstat")]:
        for row in excel_sheet(sheet):
            name_raw = row.get("name", "")
            if name_raw in (None, ""):
                continue
            name = str(name_raw).strip()
            if not name or name.lower() in {"none", "atribute", "attribute"}:
                continue

            elm = find_element(app, name, cls)
            if elm is None:
                log_line(f"[WARN] {sheet}: nie znaleziono {cls} '{name}'")
                continue

            p = float(row.get("P") or 0.0)
            q = float(row.get("Q") or 0.0)

            set_attr(elm, ["pgini"], p)
            set_attr(elm, ["qgini", "qsetp"], q)

            rows_out.append(
                {
                    "name": name,
                    "type": sheet,
                    "P_kW": p,
                    "Q_kvar": q,
                }
            )

    return rows_out


def load_storage_candidates() -> List[Dict[str, Any]]:
    rows = excel_sheet("StorageCandidates")
    return [
        {
            "node": str(r.get("node", "")).strip(),
            "role": str(r.get("role", "")).strip().lower(),
            "elem_L1": str(r.get("elem_L1", r.get("elem_A", ""))).strip(),
            "elem_L2": str(r.get("elem_L2", r.get("elem_B", ""))).strip(),
            "elem_L3": str(r.get("elem_L3", r.get("elem_C", ""))).strip(),
            "Pmax_kW": min(float(r.get("Pmax_kW") or r.get("Pmax") or 0.0), P_STORAGE_PHASE_MAX_KW),
            "Smax_kVA": float(r.get("Smax_kVA") or r.get("Smax") or 0.0),
        }
        for r in rows
        if str(r.get("node", "")).strip()
    ]


def get_transformer_storage_candidate(candidates: List[Dict[str, Any]]) -> Dict[str, Any]:
    for cand in candidates:
        if cand.get("role") == "transformer":
            return cand
    raise RuntimeError("Nie znaleziono magazynu oznaczonego jako role=transformer w arkuszu StorageCandidates.")


def load_storage_transformer_steps() -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []

    def first_present(row: Dict[str, Any], keys: Sequence[str]) -> Any:
        for key in keys:
            if key in row and row.get(key) not in (None, ""):
                return row.get(key)
        return None

    for raw_row in excel_sheet("StorageTransformerControl"):
        row = {str(k).strip(): v for k, v in raw_row.items()}

        step_raw = first_present(row, ["step", "Step", "STEP"])
        p_pu_raw = first_present(
            row,
            ["P_storage_pu", "P[p.u.]", "P_storage[p.u.]", "P_pu", "p_storage_pu", "p[p.u.]"],
        )

        if step_raw in (None, "") or p_pu_raw in (None, ""):
            continue

        try:
            p_pu = float(p_pu_raw)
        except Exception:
            continue

        rows_out.append(
            {
                "step": step_raw,
                "P_storage_pu": p_pu,
            }
        )

    if not rows_out:
        raise RuntimeError(
            "Arkusz 'StorageTransformerControl' jest pusty albo nie ma poprawnych kolumn. "
            "Oczekiwane: 'step' oraz jedna z kolumn: "
            "'P_storage_pu', 'P[p.u.]', 'P_storage[p.u.]', 'P_pu'."
        )

    unique_by_p: Dict[float, Dict[str, Any]] = {}
    for row in rows_out:
        unique_by_p[float(row["P_storage_pu"])] = row

    out = list(unique_by_p.values())
    out.sort(key=lambda r: float(r["P_storage_pu"]))

    if 0.0 not in [float(r["P_storage_pu"]) for r in out]:
        out.append({"step": 0, "P_storage_pu": 0.0})
        out.sort(key=lambda r: float(r["P_storage_pu"]))

    return out


def apply_storage_to_candidate(app: Any, candidate: Dict[str, Any], rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    applied_rows: List[Dict[str, Any]] = []

    for row in rows:
        ph = row["phase"]
        name = candidate.get(f"elem_{ph}", "")
        p = float(row.get("p_storage_kw") or 0.0)
        q = float(row.get("q_storage_kvar") or 0.0)

        if (p > 0.0) and not ALLOW_STORAGE_CHARGE:
            p = 0.0
        if (p < 0.0) and not ALLOW_STORAGE_DISCHARGE:
            p = 0.0
        p = max(-P_STORAGE_PHASE_MAX_KW, min(P_STORAGE_PHASE_MAX_KW, p))

        applied_row = {
            "node": candidate.get("node", ""),
            "element": name,
            "phase": ph,
            "stage_name": row.get("stage_name", ""),
            "step_label": row.get("step_label", ""),
            "step_pu": row.get("step_pu", ""),
            "p_export_before_kw": row.get("p_export_before_kw", ""),
            "p_export_after_kw": row.get("p_export_after_kw", ""),
            "i_before_a": row.get("i_before_a", ""),
            "i_after_a": row.get("i_after_a", ""),
            "i_unbalance_before_percent": row.get("i_unbalance_before_percent", ""),
            "i_unbalance_after_percent": row.get("i_unbalance_after_percent", ""),
            "angle_before_deg": row.get("angle_before_deg", ""),
            "angle_after_deg": row.get("angle_after_deg", ""),
            "p_storage_kw": p,
            "q_storage_kvar": q if KEEP_INITIAL_STORAGE_Q else 0.0,
            "Pmax_kW": float(candidate.get("Pmax_kW") or 0.0),
            "Smax_kVA": float(candidate.get("Smax_kVA") or 0.0),
        }

        if not name:
            log_line(f"[WARN] Storage: brak przypisania elementu dla fazy {ph}")
            applied_rows.append(applied_row)
            continue

        elm_lod = find_element(app, name, "ElmLod")
        if elm_lod is not None:
            set_attr(elm_lod, [PHASE_ATTR_P_LOAD[ph], "plini"], p)
            set_attr(elm_lod, [PHASE_ATTR_Q_LOAD[ph], "qlini"], applied_row["q_storage_kvar"])
            applied_rows.append(applied_row)
            continue

        elm_gen = (
            find_element(app, name, "ElmGenstat")
            or find_element(app, name, "ElmSym")
            or find_element(app, name, "ElmPvsys")
        )
        if elm_gen is not None:
            set_attr(elm_gen, ["pgini"], -p)
            set_attr(elm_gen, ["qgini", "qsetp"], -applied_row["q_storage_kvar"])
            applied_row["p_storage_kw_in_model"] = -p
            applied_row["q_storage_kvar_in_model"] = -applied_row["q_storage_kvar"]
            applied_rows.append(applied_row)
            continue

        log_line(f"[WARN] Storage: nie znaleziono elementu '{name}' dla fazy {ph}")
        applied_rows.append(applied_row)

    total_p = sum(float(r.get("p_storage_kw") or 0.0) for r in applied_rows)
    if abs(total_p) > P_STORAGE_TOTAL_MAX_KW + EPS:
        scale = P_STORAGE_TOTAL_MAX_KW / max(abs(total_p), EPS)
        for r in applied_rows:
            r["p_storage_kw"] = float(r["p_storage_kw"]) * scale

    return applied_rows


def zero_storage_rows() -> List[Dict[str, Any]]:
    return [
        {
            "phase": "L1",
            "stage_name": "start",
            "step_label": 0,
            "step_pu": 0.0,
            "p_export_before_kw": "",
            "p_export_after_kw": "",
            "i_before_a": "",
            "i_after_a": "",
            "i_unbalance_before_percent": "",
            "i_unbalance_after_percent": "",
            "angle_before_deg": "",
            "angle_after_deg": "",
            "p_storage_kw": 0.0,
            "q_storage_kvar": 0.0,
        },
        {
            "phase": "L2",
            "stage_name": "start",
            "step_label": 0,
            "step_pu": 0.0,
            "p_export_before_kw": "",
            "p_export_after_kw": "",
            "i_before_a": "",
            "i_after_a": "",
            "i_unbalance_before_percent": "",
            "i_unbalance_after_percent": "",
            "angle_before_deg": "",
            "angle_after_deg": "",
            "p_storage_kw": 0.0,
            "q_storage_kvar": 0.0,
        },
        {
            "phase": "L3",
            "stage_name": "start",
            "step_label": 0,
            "step_pu": 0.0,
            "p_export_before_kw": "",
            "p_export_after_kw": "",
            "i_before_a": "",
            "i_after_a": "",
            "i_unbalance_before_percent": "",
            "i_unbalance_after_percent": "",
            "angle_before_deg": "",
            "angle_after_deg": "",
            "p_storage_kw": 0.0,
            "q_storage_kvar": 0.0,
        },
    ]


def set_all_pv_to_qvchar(app: Any) -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []

    for pv in app.GetCalcRelevantObjects("*.ElmPvsys"):
        set_attr(pv, ["av_mode"], "qvchar")
        rows_out.append(
            {
                "name": getattr(pv, "loc_name", ""),
                "av_mode": get_attr(pv, ["av_mode"], ""),
                "P_kW": get_float(pv, ["pgini"], 0.0),
                "Q_kvar": get_float(pv, ["qgini", "qsetp"], 0.0),
            }
        )

    return rows_out


def set_initial_model(app: Any) -> Dict[str, List[Dict[str, Any]]]:
    log_line("Ustawianie modelu z Excela...")

    loads_data = set_loads_from_excel(app)
    pv_data = set_pv_from_excel(app)
    other_gen_data = set_other_generators_from_excel(app)

    candidates = load_storage_candidates()
    transformer_candidate = get_transformer_storage_candidate(candidates)
    storage_data = apply_storage_to_candidate(app, transformer_candidate, zero_storage_rows())

    return {
        "loads_input": loads_data,
        "pv_input": pv_data,
        "other_generators_input": other_gen_data,
        "storage_input": storage_data,
        "transformer_storage_meta": [transformer_candidate],
    }


# =============================================================================
# ODCZYTY
# =============================================================================

def collect_node_voltages(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []

    for bus in app.GetCalcRelevantObjects("*.ElmTerm"):
        row = {"node": getattr(bus, "loc_name", "")}
        for ph in PHASES:
            pf_ph = PF_PHASE[ph]
            row[f"U_{ph}_pu"] = get_float(bus, [f"m:u:{pf_ph}"], math.nan)
            row[f"angle_{ph}_deg"] = get_float(bus, [f"m:phiu:{pf_ph}"], math.nan)
        row["kU2_percent"] = get_float(bus, ["m:ubfac"], math.nan)
        rows.append(row)

    return rows


def collect_transformer_results(tr: Any) -> Dict[str, Any]:
    row = {"transformer": getattr(tr, "loc_name", "")}

    i_complex = []

    for ph in PHASES:
        pf_ph = PF_PHASE[ph]

        i_ka = get_float(tr, [f"m:I:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:I:{pf_ph}"], 0.0)
        phi_i = get_float(tr, [f"m:phii:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:phii:{pf_ph}"], math.nan)
        u_pu = get_float(tr, [f"m:u:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:u:{pf_ph}"], math.nan)

        p = get_attr(tr, [f"m:P:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:P:{pf_ph}", f"c:P:{pf_ph}"], None)
        q = get_attr(tr, [f"m:Q:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:Q:{pf_ph}", f"c:Q:{pf_ph}"], None)

        i_a = float(i_ka or 0.0) * 1000.0
        p_kw = None if p is None else float(p)
        q_kvar = None if q is None else float(q)

        row[f"I_tr_{ph}_A"] = i_a
        row[f"angle_I_tr_{ph}_deg"] = phi_i
        row[f"U_tr_{ph}_pu"] = u_pu
        row[f"P_tr_{ph}_kW"] = p_kw
        row[f"Q_tr_{ph}_kvar"] = q_kvar

        i_complex.append(angle_deg_to_complex(i_a, float(phi_i) if math.isfinite(float(phi_i)) else 0.0))

    row["loading_percent"] = get_float(tr, ["c:loading"], math.nan)

    i_avg = sum(float(row.get(f"I_tr_{ph}_A") or 0.0) for ph in PHASES) / 3.0
    row["I_unbalance_tr_percent"] = (
        max(abs(float(row.get(f"I_tr_{ph}_A") or 0.0) - i_avg) for ph in PHASES) / i_avg * 100.0
        if i_avg > EPS else 0.0
    )

    row["I_neutral_A"] = abs(sum(i_complex))

    return row


def collect_branch_losses(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []

    for line in app.GetCalcRelevantObjects("*.ElmLne"):
        p_loss = get_attr(line, ["c:LossP", "c:Ploss", "c:Losses"], 0.0)
        q_loss = get_attr(line, ["c:LossQ", "c:Qloss"], 0.0)

        rows.append(
            {
                "branch": getattr(line, "loc_name", ""),
                "P_loss_kW": float(p_loss or 0.0),
                "Q_loss_kvar": float(q_loss or 0.0),
                "loading_percent": get_float(line, ["c:loading"], math.nan),
            }
        )

    return rows


def collect_pv_setpoints(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []

    for pv in app.GetCalcRelevantObjects("*.ElmPvsys"):
        p_set = get_float(pv, ["pgini"], 0.0)
        q_set = get_float(pv, ["qgini", "qsetp"], 0.0)
        q_res = get_float(pv, ["m:Q:bus1", "m:Q:buslv", "m:Q", "c:Q"], q_set)
        smax = get_float(pv, ["sgn", "sn", "snom"], 0.0)

        rows.append(
            {
                "name": getattr(pv, "loc_name", ""),
                "P_kW": p_set,
                "Q_kvar": q_res,
                "Q_set_kvar": q_set,
                "S_apparent_kVA": math.sqrt(p_set * p_set + q_res * q_res),
                "Smax_kVA": smax,
                "av_mode": get_attr(pv, ["av_mode"], ""),
            }
        )

    return rows


# =============================================================================
# POMOCNICZE OBLICZENIA
# =============================================================================

def angle_deg_to_complex(magnitude: float, angle_deg: float) -> complex:
    if not math.isfinite(magnitude):
        return complex(math.nan, math.nan)
    if not math.isfinite(angle_deg):
        return complex(magnitude, 0.0)
    return cmath.rect(magnitude, math.radians(angle_deg))


def calc_symmetrical_components(v_a: complex, v_b: complex, v_c: complex) -> Tuple[complex, complex, complex]:
    a = complex(-0.5, math.sqrt(3.0) / 2.0)
    a2 = a * a

    u0 = (v_a + v_b + v_c) / 3.0
    u1 = (v_a + a * v_b + a2 * v_c) / 3.0
    u2 = (v_a + a2 * v_b + a * v_c) / 3.0
    return u0, u1, u2


def rms_deviation(values: List[float], target: float) -> float:
    if not values:
        return math.nan
    return math.sqrt(sum((v - target) ** 2 for v in values) / len(values))


def export_import_from_transformer_phases(p_tr: List[float]) -> Tuple[float, float]:
    export_pos = sum(max(0.0, p) for p in p_tr)
    export_neg = sum(max(0.0, -p) for p in p_tr)

    if TRANSFORMER_EXPORT_POSITIVE:
        return export_pos, export_neg
    return export_neg, export_pos


def calc_all_voltage_values(raw: Dict[str, List[Dict[str, Any]]]) -> List[float]:
    vals: List[float] = []
    for row in raw["node_voltages"]:
        for ph in PHASES:
            try:
                u = float(row[f"U_{ph}_pu"])
                if math.isfinite(u):
                    vals.append(u)
            except Exception:
                pass
    return vals


def calc_penalty_u(raw: Dict[str, List[Dict[str, Any]]]) -> float:
    vals = calc_all_voltage_values(raw)
    if not vals:
        return 0.0

    total = 0.0
    for u in vals:
        total += (
            (max(0.0, u - U_MAX_ALLOWED_OBJ) / max(U_TOL_PU, EPS)) ** 2
            + (max(0.0, U_MIN_ALLOWED_OBJ - u) / max(U_TOL_PU, EPS)) ** 2
        )
    return total


def calc_penalty_lines(raw: Dict[str, List[Dict[str, Any]]]) -> float:
    total = 0.0
    for row in raw["branch_losses"]:
        try:
            loading = float(row.get("loading_percent"))
            if math.isfinite(loading):
                total += max(0.0, loading / 100.0 - 1.0) ** 2
        except Exception:
            pass
    return total


def calc_penalty_trafo(raw: Dict[str, List[Dict[str, Any]]]) -> float:
    tr = raw["transformer_phase_results"][0] if raw["transformer_phase_results"] else {}
    try:
        loading = float(tr.get("loading_percent"))
        if math.isfinite(loading):
            return max(0.0, loading / 100.0 - 1.0) ** 2
    except Exception:
        pass
    return 0.0


def calc_ju(raw: Dict[str, List[Dict[str, Any]]]) -> float:
    vals = calc_all_voltage_values(raw)
    if not vals:
        return math.nan
    return math.sqrt(sum((u - TARGET_VOLTAGE_PU) ** 2 for u in vals) / len(vals))


def calc_ja(raw: Dict[str, List[Dict[str, Any]]]) -> float:
    alpha2_vals: List[float] = []

    for row in raw.get("node_sequence_components", []):
        try:
            a2 = float(row.get("alpha2"))
            if math.isfinite(a2):
                alpha2_vals.append(max(0.0, a2))
        except Exception:
            continue

    if alpha2_vals:
        return math.sqrt(sum(v * v for v in alpha2_vals) / len(alpha2_vals))

    ku2_vals: List[float] = []
    for row in raw.get("node_voltages", []):
        try:
            ku2 = float(row.get("kU2_percent"))
            if math.isfinite(ku2):
                ku2_vals.append(max(0.0, ku2 / 100.0))
        except Exception:
            continue

    if ku2_vals:
        return math.sqrt(sum(v * v for v in ku2_vals) / len(ku2_vals))

    return math.nan


def calc_ji(i_unbalance_tr_percent: float) -> float:
    if not math.isfinite(i_unbalance_tr_percent):
        return math.nan
    return i_unbalance_tr_percent / 100.0


def calc_jp(p_export_total_kw: float) -> float:
    ref = max(P_EXPORT_REF_KW, EPS)
    return p_export_total_kw / ref


def calc_f1(ju: float, ja: float, penalty_u: float, penalty_lines: float, penalty_trafo: float) -> float:
    ju_used = 0.0 if not math.isfinite(ju) else ju
    ja_used = 0.0 if not math.isfinite(ja) else ja
    return (
        0.50 * ju_used
        + 0.50 * ja_used
        + K_U * penalty_u
        + K_L * penalty_lines
        + K_T * penalty_trafo
    )


def calc_f2(ju: float, ji: float, jp: float, penalty_u: float, penalty_lines: float, penalty_trafo: float) -> float:
    ju_used = 0.0 if not math.isfinite(ju) else ju
    ji_used = 0.0 if not math.isfinite(ji) else ji
    jp_used = 0.0 if not math.isfinite(jp) else jp
    return (
        0.33 * ju_used
        + 0.33 * ji_used
        + 0.33 * jp_used
        + K_U * penalty_u
        + K_L * penalty_lines
        + K_T * penalty_trafo
    )


# =============================================================================
# REGULATOR TRANSFORMATOROWY DWUETAPOWY
# =============================================================================

def export_kw_from_p_tr(p_tr_kw: float) -> float:
    if TRANSFORMER_EXPORT_POSITIVE:
        return max(0.0, p_tr_kw)
    return max(0.0, -p_tr_kw)


def i_unbalance_percent_from_tr_row(tr_row: Dict[str, Any]) -> float:
    i_vals = [float(tr_row.get(f"I_tr_{ph}_A") or 0.0) for ph in PHASES]
    i_avg = sum(i_vals) / 3.0 if i_vals else 0.0
    if i_avg <= EPS:
        return 0.0
    return max(abs(i - i_avg) for i in i_vals) / i_avg * 100.0


def clone_storage_rows(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    return deepcopy(rows)


def positive_steps(step_rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    rows = [r for r in step_rows if float(r["P_storage_pu"]) >= 0.0]
    rows.sort(key=lambda r: float(r["P_storage_pu"]))
    return rows


def step_index(step_rows: List[Dict[str, Any]], step_pu: float) -> int:
    vals = [float(r["P_storage_pu"]) for r in step_rows]
    for i, v in enumerate(vals):
        if abs(v - step_pu) <= 1e-12:
            return i
    best_idx = 0
    best_dist = float("inf")
    for i, v in enumerate(vals):
        d = abs(v - step_pu)
        if d < best_dist:
            best_dist = d
            best_idx = i
    return best_idx


def apply_and_run_storage(
    app: Any,
    ldf: Any,
    candidate: Dict[str, Any],
    storage_rows: List[Dict[str, Any]],
) -> Tuple[List[Dict[str, Any]], bool, Any]:
    applied_rows = apply_storage_to_candidate(app, candidate, storage_rows)
    converged, rc = try_run_loadflow(ldf)
    return applied_rows, converged, rc


def build_storage_attempt_rows(
    stage: str,
    attempt_no: int,
    candidate: Dict[str, Any],
    storage_rows: List[Dict[str, Any]],
    converged: bool,
    loadflow_result: Any,
    note: str = "",
) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for row in storage_rows:
        ph = row.get("phase", "")
        rows.append(
            {
                "stage": stage,
                "attempt_no": attempt_no,
                "storage_node": candidate.get("node", ""),
                "storage_element": candidate.get(f"elem_{ph}", ""),
                "phase": ph,
                "stage_name": row.get("stage_name", ""),
                "step_label": row.get("step_label", ""),
                "step_pu": row.get("step_pu", ""),
                "p_export_before_kw": row.get("p_export_before_kw", ""),
                "p_export_after_kw": row.get("p_export_after_kw", ""),
                "i_before_a": row.get("i_before_a", ""),
                "i_after_a": row.get("i_after_a", ""),
                "i_unbalance_before_percent": row.get("i_unbalance_before_percent", ""),
                "i_unbalance_after_percent": row.get("i_unbalance_after_percent", ""),
                "angle_before_deg": row.get("angle_before_deg", ""),
                "angle_after_deg": row.get("angle_after_deg", ""),
                "p_storage_kw": row.get("p_storage_kw", 0.0),
                "q_storage_kvar": row.get("q_storage_kvar", 0.0),
                "Pmax_kW": candidate.get("Pmax_kW", 0.0),
                "Smax_kVA": candidate.get("Smax_kVA", 0.0),
                "converged": converged,
                "loadflow_result": str(loadflow_result),
                "note": note,
            }
        )
    return rows


def optimize_single_phase_export(
    app: Any,
    ldf: Any,
    tr: Any,
    candidate: Dict[str, Any],
    base_storage_rows: List[Dict[str, Any]],
    phase: str,
    iteration_no: int,
    step_rows: List[Dict[str, Any]],
    loadflow_attempts: List[Dict[str, Any]],
    phase_search_history: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    ph_idx = PHASES.index(phase)
    pmax = float(candidate.get("Pmax_kW") or 0.0)

    tr_before = collect_transformer_results(tr)
    p_export_before = export_kw_from_p_tr(float(tr_before.get(f"P_tr_{phase}_kW") or 0.0))
    i_before = float(tr_before.get(f"I_tr_{phase}_A") or 0.0)
    angle_before = float(tr_before.get(f"angle_I_tr_{phase}_deg") or 0.0)
    i_unb_before = i_unbalance_percent_from_tr_row(tr_before)

    best_rows = clone_storage_rows(base_storage_rows)
    best_export = p_export_before

    best_rows[ph_idx]["stage_name"] = "export_search_start"
    best_rows[ph_idx]["p_export_before_kw"] = p_export_before
    best_rows[ph_idx]["p_export_after_kw"] = p_export_before
    best_rows[ph_idx]["i_before_a"] = i_before
    best_rows[ph_idx]["i_after_a"] = i_before
    best_rows[ph_idx]["i_unbalance_before_percent"] = i_unb_before
    best_rows[ph_idx]["i_unbalance_after_percent"] = i_unb_before
    best_rows[ph_idx]["angle_before_deg"] = angle_before
    best_rows[ph_idx]["angle_after_deg"] = angle_before

    if p_export_before <= EPS:
        restore_rows, conv_restore, rc_restore = apply_and_run_storage(app, ldf, candidate, best_rows)
        loadflow_attempts.extend(
            build_storage_attempt_rows(
                stage="phase_export_restore",
                attempt_no=iteration_no,
                candidate=candidate,
                storage_rows=restore_rows,
                converged=conv_restore,
                loadflow_result=rc_restore,
                note=f"{phase}: no export to reduce",
            )
        )
        if not conv_restore:
            raise RuntimeError(f"Nie udało się odtworzyć stanu bez eksportu dla fazy {phase}.")
        return best_rows

    for step_row in positive_steps(step_rows):
        step_pu = float(step_row["P_storage_pu"])

        trial_rows = clone_storage_rows(base_storage_rows)
        trial_rows[ph_idx]["stage_name"] = "export_search"
        trial_rows[ph_idx]["step_label"] = step_row["step"]
        trial_rows[ph_idx]["step_pu"] = step_pu
        trial_rows[ph_idx]["p_export_before_kw"] = p_export_before
        trial_rows[ph_idx]["i_before_a"] = i_before
        trial_rows[ph_idx]["i_unbalance_before_percent"] = i_unb_before
        trial_rows[ph_idx]["angle_before_deg"] = angle_before
        trial_rows[ph_idx]["p_storage_kw"] = step_pu * pmax
        trial_rows[ph_idx]["q_storage_kvar"] = 0.0

        applied_rows, converged, rc = apply_and_run_storage(app, ldf, candidate, trial_rows)

        if converged:
            tr_after = collect_transformer_results(tr)
            p_export_after = export_kw_from_p_tr(float(tr_after.get(f"P_tr_{phase}_kW") or 0.0))
            i_after = float(tr_after.get(f"I_tr_{phase}_A") or 0.0)
            angle_after = float(tr_after.get(f"angle_I_tr_{phase}_deg") or 0.0)
            i_unb_after = i_unbalance_percent_from_tr_row(tr_after)

            applied_rows[ph_idx]["p_export_after_kw"] = p_export_after
            applied_rows[ph_idx]["i_after_a"] = i_after
            applied_rows[ph_idx]["i_unbalance_after_percent"] = i_unb_after
            applied_rows[ph_idx]["angle_after_deg"] = angle_after
        else:
            p_export_after = math.inf
            i_after = math.nan
            angle_after = math.nan
            i_unb_after = math.inf

        loadflow_attempts.extend(
            build_storage_attempt_rows(
                stage="phase_export_search",
                attempt_no=iteration_no,
                candidate=candidate,
                storage_rows=applied_rows,
                converged=converged,
                loadflow_result=rc,
                note=f"{phase}: export search by step",
            )
        )

        phase_search_history.append(
            {
                "iteration": iteration_no,
                "phase": phase,
                "search_stage": "export_search",
                "step_label": step_row["step"],
                "step_pu": step_pu,
                "p_export_before_kw": p_export_before,
                "p_export_after_kw": p_export_after,
                "i_before_a": i_before,
                "i_after_a": i_after,
                "i_unbalance_before_percent": i_unb_before,
                "i_unbalance_after_percent": i_unb_after,
                "angle_before_deg": angle_before,
                "angle_after_deg": angle_after,
                "converged": converged,
                "loadflow_result": str(rc),
            }
        )

        if converged and p_export_after <= best_export + EPS:
            best_export = p_export_after
            best_rows = clone_storage_rows(applied_rows)
        else:
            restore_rows, conv_restore, rc_restore = apply_and_run_storage(app, ldf, candidate, best_rows)
            loadflow_attempts.extend(
                build_storage_attempt_rows(
                    stage="phase_export_restore",
                    attempt_no=iteration_no,
                    candidate=candidate,
                    storage_rows=restore_rows,
                    converged=conv_restore,
                    loadflow_result=rc_restore,
                    note=f"{phase}: restore best export state",
                )
            )
            if not conv_restore:
                raise RuntimeError(f"Nie udało się odtworzyć najlepszego stanu eksportowego dla fazy {phase}.")
            return best_rows

    restore_rows, conv_restore, rc_restore = apply_and_run_storage(app, ldf, candidate, best_rows)
    loadflow_attempts.extend(
        build_storage_attempt_rows(
            stage="phase_export_restore",
            attempt_no=iteration_no,
            candidate=candidate,
            storage_rows=restore_rows,
            converged=conv_restore,
            loadflow_result=rc_restore,
            note=f"{phase}: final restore best export state",
        )
    )
    if not conv_restore:
        raise RuntimeError(f"Nie udało się ustawić końcowego najlepszego stanu eksportowego dla fazy {phase}.")
    return best_rows


def asymmetry_trial(
    app: Any,
    ldf: Any,
    tr: Any,
    candidate: Dict[str, Any],
    base_storage_rows: List[Dict[str, Any]],
    phase: str,
    step_row: Dict[str, Any],
    export_reference: Dict[str, float],
    attempt_no: int,
    loadflow_attempts: List[Dict[str, Any]],
    phase_search_history: List[Dict[str, Any]],
) -> Tuple[bool, List[Dict[str, Any]], Dict[str, Any], float]:
    ph_idx = PHASES.index(phase)
    pmax = float(candidate.get("Pmax_kW") or 0.0)

    tr_before = collect_transformer_results(tr)
    i_unb_before = i_unbalance_percent_from_tr_row(tr_before)

    trial_rows = clone_storage_rows(base_storage_rows)
    trial_rows[ph_idx]["stage_name"] = "asymmetry_search"
    trial_rows[ph_idx]["step_label"] = step_row["step"]
    trial_rows[ph_idx]["step_pu"] = float(step_row["P_storage_pu"])
    trial_rows[ph_idx]["p_export_before_kw"] = export_kw_from_p_tr(float(tr_before.get(f"P_tr_{phase}_kW") or 0.0))
    trial_rows[ph_idx]["i_before_a"] = float(tr_before.get(f"I_tr_{phase}_A") or 0.0)
    trial_rows[ph_idx]["i_unbalance_before_percent"] = i_unb_before
    trial_rows[ph_idx]["angle_before_deg"] = float(tr_before.get(f"angle_I_tr_{phase}_deg") or 0.0)
    trial_rows[ph_idx]["p_storage_kw"] = float(step_row["P_storage_pu"]) * pmax
    trial_rows[ph_idx]["q_storage_kvar"] = 0.0

    applied_rows, converged, rc = apply_and_run_storage(app, ldf, candidate, trial_rows)

    accepted = False
    tr_after = {}
    score = math.inf

    if converged:
        tr_after = collect_transformer_results(tr)
        i_unb_after = i_unbalance_percent_from_tr_row(tr_after)

        p_export_after = {
            ph: export_kw_from_p_tr(float(tr_after.get(f"P_tr_{ph}_kW") or 0.0))
            for ph in PHASES
        }

        applied_rows[ph_idx]["p_export_after_kw"] = p_export_after[phase]
        applied_rows[ph_idx]["i_after_a"] = float(tr_after.get(f"I_tr_{phase}_A") or 0.0)
        applied_rows[ph_idx]["i_unbalance_after_percent"] = i_unb_after
        applied_rows[ph_idx]["angle_after_deg"] = float(tr_after.get(f"angle_I_tr_{phase}_deg") or 0.0)

        export_ok = True
        for ph in PHASES:
            if p_export_after[ph] > export_reference[ph] + ASYMMETRY_EXPORT_TOLERANCE_KW:
                export_ok = False
                break

        if export_ok and i_unb_after + EPS < i_unb_before:
            accepted = True
            score = i_unb_after
    else:
        i_unb_after = math.inf

    loadflow_attempts.extend(
        build_storage_attempt_rows(
            stage="phase_asymmetry_search",
            attempt_no=attempt_no,
            candidate=candidate,
            storage_rows=applied_rows,
            converged=converged,
            loadflow_result=rc,
            note=f"{phase}: asymmetry correction trial",
        )
    )

    phase_search_history.append(
        {
            "iteration": attempt_no,
            "phase": phase,
            "search_stage": "asymmetry_search",
            "step_label": step_row["step"],
            "step_pu": float(step_row["P_storage_pu"]),
            "i_unbalance_before_percent": i_unb_before,
            "i_unbalance_after_percent": i_unb_after,
            "accepted": accepted,
            "converged": converged,
            "loadflow_result": str(rc),
        }
    )

    return accepted, applied_rows, tr_after, score


def try_reduce_asymmetry(
    app: Any,
    ldf: Any,
    tr: Any,
    candidate: Dict[str, Any],
    storage_rows: List[Dict[str, Any]],
    step_rows: List[Dict[str, Any]],
    iteration_no: int,
    loadflow_attempts: List[Dict[str, Any]],
    phase_search_history: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    if not ASYMMETRY_ENABLE:
        return storage_rows

    current_rows = clone_storage_rows(storage_rows)

    for _pass in range(ASYMMETRY_MAX_PASSES):
        tr_ref = collect_transformer_results(tr)
        export_reference = {
            ph: export_kw_from_p_tr(float(tr_ref.get(f"P_tr_{ph}_kW") or 0.0))
            for ph in PHASES
        }
        current_unb = i_unbalance_percent_from_tr_row(tr_ref)

        best_candidate_rows = None
        best_candidate_score = current_unb

        for ph in PHASES:
            idx = PHASES.index(ph)
            current_step_pu = float(current_rows[idx].get("step_pu") or 0.0)
            idx_now = step_index(step_rows, current_step_pu)

            neighbor_indices = []
            if idx_now - 1 >= 0:
                neighbor_indices.append(idx_now - 1)
            if idx_now + 1 < len(step_rows):
                neighbor_indices.append(idx_now + 1)

            for nidx in neighbor_indices:
                step_row = step_rows[nidx]

                accepted, applied_rows, _tr_after, score = asymmetry_trial(
                    app=app,
                    ldf=ldf,
                    tr=tr,
                    candidate=candidate,
                    base_storage_rows=current_rows,
                    phase=ph,
                    step_row=step_row,
                    export_reference=export_reference,
                    attempt_no=iteration_no,
                    loadflow_attempts=loadflow_attempts,
                    phase_search_history=phase_search_history,
                )

                if accepted and score + EPS < best_candidate_score:
                    best_candidate_rows = clone_storage_rows(applied_rows)
                    best_candidate_score = score

                restore_rows, conv_restore, rc_restore = apply_and_run_storage(app, ldf, candidate, current_rows)
                loadflow_attempts.extend(
                    build_storage_attempt_rows(
                        stage="phase_asymmetry_restore",
                        attempt_no=iteration_no,
                        candidate=candidate,
                        storage_rows=restore_rows,
                        converged=conv_restore,
                        loadflow_result=rc_restore,
                        note=f"{ph}: restore current state after asymmetry trial",
                    )
                )
                if not conv_restore:
                    raise RuntimeError(f"Nie udało się odtworzyć stanu po próbie asymetrii dla fazy {ph}.")

        if best_candidate_rows is None:
            break

        current_rows = clone_storage_rows(best_candidate_rows)
        applied_rows, conv_apply, rc_apply = apply_and_run_storage(app, ldf, candidate, current_rows)
        loadflow_attempts.extend(
            build_storage_attempt_rows(
                stage="phase_asymmetry_accept",
                attempt_no=iteration_no,
                candidate=candidate,
                storage_rows=applied_rows,
                converged=conv_apply,
                loadflow_result=rc_apply,
                note="accept best asymmetry correction move",
            )
        )
        if not conv_apply:
            raise RuntimeError("Nie udało się zastosować zaakceptowanej korekty asymetrii.")

        log_line(
            f"Asymmetry pass accepted in iteration {iteration_no}: "
            f"I_unbalance {current_unb:.2f}% -> {best_candidate_score:.2f}%"
        )

    return current_rows


def rescue_storage_rows(candidate: Dict[str, Any], rel: float) -> List[Dict[str, Any]]:
    pmax = float(candidate.get("Pmax_kW") or 0.0)
    rows = []
    for ph in PHASES:
        rows.append(
            {
                "phase": ph,
                "stage_name": "rescue",
                "step_label": f"rescue_{rel:.2f}",
                "step_pu": rel,
                "p_export_before_kw": "",
                "p_export_after_kw": "",
                "i_before_a": "",
                "i_after_a": "",
                "i_unbalance_before_percent": "",
                "i_unbalance_after_percent": "",
                "angle_before_deg": "",
                "angle_after_deg": "",
                "p_storage_kw": rel * pmax,
                "q_storage_kvar": 0.0,
            }
        )
    return rows


# =============================================================================
# WSKAŹNIKI
# =============================================================================

def calculate_indicators(raw: Dict[str, List[Dict[str, Any]]]) -> Dict[str, Any]:
    voltages: List[float] = []
    phase_spreads: List[float] = []
    ku2_vals: List[float] = []

    alpha0_vals: List[float] = []
    alpha2_vals: List[float] = []
    u0_abs_vals: List[float] = []
    u1_abs_vals: List[float] = []
    u2_abs_vals: List[float] = []

    node_sequence_rows: List[Dict[str, Any]] = []

    for row in raw["node_voltages"]:
        vals = []
        for ph in PHASES:
            try:
                v = float(row[f"U_{ph}_pu"])
                if math.isfinite(v):
                    vals.append(v)
                    voltages.append(v)
            except Exception:
                pass

        if vals:
            phase_spreads.append(max(vals) - min(vals))

        try:
            k = float(row.get("kU2_percent"))
            if math.isfinite(k):
                ku2_vals.append(k)
        except Exception:
            pass

        try:
            va = angle_deg_to_complex(float(row["U_L1_pu"]), float(row["angle_L1_deg"]))
            vb = angle_deg_to_complex(float(row["U_L2_pu"]), float(row["angle_L2_deg"]))
            vc = angle_deg_to_complex(float(row["U_L3_pu"]), float(row["angle_L3_deg"]))

            u0, u1, u2 = calc_symmetrical_components(va, vb, vc)

            u0_abs = abs(u0)
            u1_abs = abs(u1)
            u2_abs = abs(u2)

            alpha0 = u0_abs / u1_abs if u1_abs > EPS else math.nan
            alpha2 = u2_abs / u1_abs if u1_abs > EPS else math.nan

            if math.isfinite(alpha0):
                alpha0_vals.append(alpha0)
            if math.isfinite(alpha2):
                alpha2_vals.append(alpha2)

            if math.isfinite(u0_abs):
                u0_abs_vals.append(u0_abs)
            if math.isfinite(u1_abs):
                u1_abs_vals.append(u1_abs)
            if math.isfinite(u2_abs):
                u2_abs_vals.append(u2_abs)

            node_sequence_rows.append(
                {
                    "node": row.get("node", ""),
                    "U0_abs_pu": u0_abs,
                    "U1_abs_pu": u1_abs,
                    "U2_abs_pu": u2_abs,
                    "alpha0": alpha0,
                    "alpha2": alpha2,
                }
            )
        except Exception:
            node_sequence_rows.append(
                {
                    "node": row.get("node", ""),
                    "U0_abs_pu": math.nan,
                    "U1_abs_pu": math.nan,
                    "U2_abs_pu": math.nan,
                    "alpha0": math.nan,
                    "alpha2": math.nan,
                }
            )

    tr = raw["transformer_phase_results"][0] if raw["transformer_phase_results"] else {}

    i_vals = [float(tr.get(f"I_tr_{ph}_A") or 0.0) for ph in PHASES]
    i_avg = sum(i_vals) / 3.0 if i_vals else 0.0
    i_unb = max(abs(i - i_avg) for i in i_vals) / i_avg * 100.0 if i_avg > EPS else 0.0

    p_tr = [float(tr.get(f"P_tr_{ph}_kW") or 0.0) for ph in PHASES]
    p_export_total, p_import_total = export_import_from_transformer_phases(p_tr)

    p_loss = sum(float(row.get("P_loss_kW") or 0.0) for row in raw["branch_losses"])
    q_loss = sum(float(row.get("Q_loss_kvar") or 0.0) for row in raw["branch_losses"])

    pv_rows = raw.get("pv_setpoints", [])
    storage_rows = raw.get("storage_setpoints", [])

    p_pv_total = sum(float(row.get("P_kW") or 0.0) for row in pv_rows)
    q_pv_total = sum(float(row.get("Q_kvar") or 0.0) for row in pv_rows)
    q_pv_abs_sum = sum(abs(float(row.get("Q_kvar") or 0.0)) for row in pv_rows)

    p_storage_L1 = float(storage_rows[0].get("p_storage_kw") or 0.0) if len(storage_rows) > 0 else 0.0
    p_storage_L2 = float(storage_rows[1].get("p_storage_kw") or 0.0) if len(storage_rows) > 1 else 0.0
    p_storage_L3 = float(storage_rows[2].get("p_storage_kw") or 0.0) if len(storage_rows) > 2 else 0.0
    p_storage_total = p_storage_L1 + p_storage_L2 + p_storage_L3
    q_storage_total = sum(float(row.get("q_storage_kvar") or 0.0) for row in storage_rows)

    udev_mean_1_00 = sum(abs(v - 1.0) for v in voltages) / len(voltages) if voltages else math.nan
    udev_rms_1_00 = rms_deviation(voltages, 1.00)
    udev_rms_1_05 = rms_deviation(voltages, 1.05)

    alpha0_mean = sum(alpha0_vals) / len(alpha0_vals) if alpha0_vals else math.nan
    alpha0_max = max(alpha0_vals) if alpha0_vals else math.nan
    alpha2_mean = sum(alpha2_vals) / len(alpha2_vals) if alpha2_vals else math.nan
    alpha2_max = max(alpha2_vals) if alpha2_vals else math.nan

    fcelu_3_voltage_component = 0.04 * (udev_rms_1_05 if math.isfinite(udev_rms_1_05) else 0.0)
    fcelu_3_alpha2_component = 0.58 * (alpha2_mean if math.isfinite(alpha2_mean) else 0.0)
    fcelu_3_alpha0_component = 0.38 * (alpha0_mean if math.isfinite(alpha0_mean) else 0.0)
    fcelu_3 = fcelu_3_voltage_component + fcelu_3_alpha2_component + fcelu_3_alpha0_component

    raw_for_penalty = {
        "node_voltages": raw["node_voltages"],
        "transformer_phase_results": raw["transformer_phase_results"],
        "branch_losses": raw["branch_losses"],
        "pv_setpoints": raw.get("pv_setpoints", []),
        "storage_setpoints": raw.get("storage_setpoints", []),
    }

    penalty_u = calc_penalty_u(raw_for_penalty)
    penalty_lines = calc_penalty_lines(raw_for_penalty)
    penalty_trafo = calc_penalty_trafo(raw_for_penalty)

    ju = calc_ju(raw_for_penalty)
    ja = calc_ja(raw)
    ji = calc_ji(i_unb)
    jp = calc_jp(p_export_total)

    f1 = calc_f1(ju, ja, penalty_u, penalty_lines, penalty_trafo)
    f2 = calc_f2(ju, ji, jp, penalty_u, penalty_lines, penalty_trafo)

    violations: List[str] = []

    for row in raw["node_voltages"]:
        for ph in PHASES:
            try:
                u = float(row[f"U_{ph}_pu"])
                if u < VOLTAGE_MIN_PU or u > VOLTAGE_MAX_PU:
                    violations.append(f"U:{row.get('node')}:{ph}:{u:.4f}")
            except Exception:
                pass

    try:
        tr_loading = float(tr.get("loading_percent"))
        if math.isfinite(tr_loading) and tr_loading > LOADING_MAX_PERCENT:
            violations.append(f"TR_loading:{tr_loading:.2f}")
    except Exception:
        pass

    for row in raw["branch_losses"]:
        try:
            line_loading = float(row.get("loading_percent"))
            if math.isfinite(line_loading) and line_loading > LINE_LOADING_MAX_PERCENT:
                violations.append(f"LINE_loading:{row.get('branch')}:{line_loading:.2f}")
        except Exception:
            pass

    indicators = {
        "Umax_pu": max(voltages) if voltages else math.nan,
        "Umin_pu": min(voltages) if voltages else math.nan,
        "Udev_mean_pu": udev_mean_1_00,
        "Udev_max_pu": max(abs(v - 1.0) for v in voltages) if voltages else math.nan,
        "Fcelu_1_Udev_rms_1_00": udev_rms_1_00,
        "Fcelu_1_Udev_rms_1_05": udev_rms_1_05,
        "dU_phase_max_pu": max(phase_spreads) if phase_spreads else math.nan,
        "kU2_max_percent": max(ku2_vals) if ku2_vals else math.nan,
        "alpha0_mean": alpha0_mean,
        "alpha0_max": alpha0_max,
        "alpha2_mean": alpha2_mean,
        "alpha2_max": alpha2_max,
        "Fcelu_3_weighted": fcelu_3,

        "I_tr_L1_A": i_vals[0],
        "I_tr_L2_A": i_vals[1],
        "I_tr_L3_A": i_vals[2],
        "I_unbalance_tr_percent": i_unb,
        "I_neutral_A": float(tr.get("I_neutral_A") or 0.0),

        "P_tr_L1_kW": p_tr[0],
        "P_tr_L2_kW": p_tr[1],
        "P_tr_L3_kW": p_tr[2],
        "P_export_total_kW": p_export_total,
        "P_import_total_kW": p_import_total,
        "P_loss_total_kW": p_loss,
        "Q_pv_abs_sum_kvar": q_pv_abs_sum,

        "P_storage_L1_kW": p_storage_L1,
        "P_storage_L2_kW": p_storage_L2,
        "P_storage_L3_kW": p_storage_L3,
        "P_storage_total_kW": p_storage_total,

        "Q_loss_total_kvar": q_loss,
        "Q_pv_total_kvar": q_pv_total,
        "P_pv_total_kW": p_pv_total,
        "Q_storage_total_kvar": q_storage_total,
        "Fcelu_2_A": float(tr.get("I_neutral_A") or 0.0),

        "JU": ju,
        "JA": ja,
        "JI": ji,
        "JP": jp,
        "penalty_U": penalty_u,
        "penalty_lines": penalty_lines,
        "penalty_trafo": penalty_trafo,
        "F1": f1,
        "F2": f2,

        "objective_profile": "transformer_flow_local",
        "J_used_total": "",
        "J_used_component_1": "local_export_reduction",
        "J_used_component_2": "local_current_asymmetry_reduction",
        "J_used_component_3": "heuristic_step_search",
        "J_used_component_4": "",
        "J_used_component_5": "",

        "constraint_violations": ";".join(violations),
    }

    raw["node_sequence_components"] = node_sequence_rows
    return indicators


def collect_results(
    app: Any,
    tr: Any,
    setup_data: Dict[str, List[Dict[str, Any]]],
    storage_rows: List[Dict[str, Any]],
    iteration_history: List[Dict[str, Any]],
    loadflow_attempts: List[Dict[str, Any]],
    phase_search_history: List[Dict[str, Any]],
    transformer_steps: List[Dict[str, Any]],
) -> Dict[str, List[Dict[str, Any]]]:
    node_voltages = collect_node_voltages(app)
    transformer_results = [collect_transformer_results(tr)]
    branch_losses = collect_branch_losses(app)
    pv_setpoints = collect_pv_setpoints(app)

    raw = {
        "node_voltages": node_voltages,
        "transformer_phase_results": transformer_results,
        "branch_losses": branch_losses,
        "pv_setpoints": pv_setpoints,
        "storage_setpoints": storage_rows,
    }
    indicators = calculate_indicators(raw)

    return {
        "control_config": [{"parameter": k, "value": v} for k, v in CONTROL_CONFIG.items()],
        "inputs_loads": setup_data["loads_input"],
        "inputs_pv": setup_data["pv_input"],
        "inputs_other_generators": setup_data["other_generators_input"],
        "inputs_storage_transformer_meta": setup_data["transformer_storage_meta"],
        "inputs_storage_transformer_steps": transformer_steps,
        "node_voltages": node_voltages,
        "node_sequence_components": raw.get("node_sequence_components", []),
        "transformer_phase_results": transformer_results,
        "pv_setpoints": pv_setpoints,
        "storage_setpoints": storage_rows,
        "iteration_history": iteration_history,
        "phase_search_history": phase_search_history,
        "loadflow_attempts": loadflow_attempts,
        "branch_losses": branch_losses,
        "indicators": [indicators],
    }


# =============================================================================
# EXCEL EXPORT
# =============================================================================

def safe_sheet_name(name: str) -> str:
    bad = "[]:*?/\\"
    for ch in bad:
        name = name.replace(ch, "_")
    return name[:31]


def write_rows(ws: Any, rows: Iterable[Dict[str, Any]]) -> None:
    rows = list(rows)
    if not rows:
        ws.append(["empty"])
        return

    headers: List[str] = []
    for row in rows:
        for key in row.keys():
            if key not in headers:
                headers.append(key)

    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])


def export_to_excel(all_tables: Dict[str, List[Dict[str, Any]]], out_file: str) -> None:
    if Workbook is None:
        raise RuntimeError("Do zapisu wyników do Excela potrzebny jest pakiet openpyxl.")

    out_dir = os.path.dirname(out_file)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, rows in all_tables.items():
        ws = wb.create_sheet(safe_sheet_name(sheet_name))
        write_rows(ws, rows)

    wb.save(out_file)


# =============================================================================
# MAIN
# =============================================================================

def run_local_transformer_flow_control() -> None:
    global EXCEL_CACHE, CONTROL_CONFIG

    try:
        log_dir = os.path.dirname(LOG_FILE)
        if log_dir:
            os.makedirs(log_dir, exist_ok=True)
        with open(LOG_FILE, "w", encoding="utf-8") as fh:
            fh.write("=== LOCAL TRANSFORMER FLOW CONTROL LOG START ===\n")
    except Exception:
        pass

    log_line("Start local_transformer_flow_control.py")

    EXCEL_CACHE = load_excel_cache(EXCEL_FILE)
    CONTROL_CONFIG = load_control_config()
    apply_runtime_config()

    log_line(f"Wczytano Excel: {EXCEL_FILE}")
    log_line(f"Arkusze: {', '.join(sorted(EXCEL_CACHE.keys()))}")
    log_line(
        f"Config: Vmin={VOLTAGE_MIN_PU}, Uref={TARGET_VOLTAGE_PU}, Vmax={VOLTAGE_MAX_PU}, "
        f"Uallowed=[{U_MIN_ALLOWED_OBJ}, {U_MAX_ALLOWED_OBJ}], "
        f"Utol={U_TOL_PU}, PexportRef={P_EXPORT_REF_KW}, "
        f"K=({K_U}, {K_L}, {K_T}), Iter={N_ITER}"
    )

    transformer_steps = load_storage_transformer_steps()
    log_line(f"Wczytano kroki StorageTransformerControl: {[(r['step'], r['P_storage_pu']) for r in transformer_steps]}")

    app, ldf = connect_powerfactory()
    tr = get_transformer(app)
    log_line(f"Odczyt parametrów transformatora ze strony: {TRANSFORMER_LV_SIDE}")

    setup_data = set_initial_model(app)
    transformer_candidate = setup_data["transformer_storage_meta"][0]

    loadflow_attempts: List[Dict[str, Any]] = []
    phase_search_history: List[Dict[str, Any]] = []

    pv_after_mode_set = set_all_pv_to_qvchar(app)
    log_line(f"Przełączono PV na qvchar: {len(pv_after_mode_set)} szt.")

    current_storage_rows = apply_storage_to_candidate(app, transformer_candidate, zero_storage_rows())

    log_line("Uruchamiam pierwszy rozpływ kontrolny...")
    converged, result_code = try_run_loadflow(ldf)

    loadflow_attempts.extend(
        build_storage_attempt_rows(
            stage="initial",
            attempt_no=0,
            candidate=transformer_candidate,
            storage_rows=current_storage_rows,
            converged=converged,
            loadflow_result=result_code,
            note="first loadflow after switching PV to qvchar and storage=0",
        )
    )

    rescue_history: List[Dict[str, Any]] = []

    if not converged:
        log_line("Pierwszy rozpływ rozbieżny. Uruchamiam pętlę ratunkową.")
        for attempt in range(1, MAX_RESCUE_ATTEMPTS + 1):
            rel = attempt / 10.0
            current_storage_rows = apply_storage_to_candidate(
                app,
                transformer_candidate,
                rescue_storage_rows(transformer_candidate, rel),
            )

            converged, result_code = try_run_loadflow(ldf)

            rescue_history.append(
                {
                    "iteration": 0,
                    "stage": "rescue",
                    "attempt": attempt,
                    "rel_charge": rel,
                    "p_phase_kw": rel * float(transformer_candidate.get("Pmax_kW") or 0.0),
                    "converged": converged,
                    "loadflow_result": str(result_code),
                    "storage_L1": transformer_candidate.get("elem_L1", ""),
                    "storage_L2": transformer_candidate.get("elem_L2", ""),
                    "storage_L3": transformer_candidate.get("elem_L3", ""),
                }
            )

            loadflow_attempts.extend(
                build_storage_attempt_rows(
                    stage="rescue",
                    attempt_no=attempt,
                    candidate=transformer_candidate,
                    storage_rows=current_storage_rows,
                    converged=converged,
                    loadflow_result=result_code,
                    note=f"rescue charging level = {rel:.2f} of Pmax",
                )
            )

            log_line(
                f"Rescue attempt {attempt}: "
                f"charge={rel:.2f} * Pmax, "
                f"L1={current_storage_rows[0]['p_storage_kw']:.2f} kW, "
                f"L2={current_storage_rows[1]['p_storage_kw']:.2f} kW, "
                f"L3={current_storage_rows[2]['p_storage_kw']:.2f} kW, "
                f"converged={converged}, result={result_code}"
            )

            if converged:
                break

    if not converged:
        log_line("[ERROR] Nie udało się uzyskać zbieżności po pętli ratunkowej.")
        result_tables = collect_results(
            app,
            tr,
            setup_data,
            current_storage_rows,
            rescue_history,
            loadflow_attempts,
            phase_search_history,
            transformer_steps,
        )
        result_tables["errors"] = [{"message": "Brak zbieżności po pętli ratunkowej."}]
        export_to_excel(result_tables, OUT_FILE)
        raise RuntimeError("Nie udało się uzyskać zbieżności nawet po pętli ratunkowej.")

    iteration_history: List[Dict[str, Any]] = rescue_history[:]

    for iter_no in range(1, N_ITER + 1):
        log_line(f"--- Iteracja transformatorowej regulacji lokalnej {iter_no} ---")

        for phase in PHASES:
            current_storage_rows = optimize_single_phase_export(
                app=app,
                ldf=ldf,
                tr=tr,
                candidate=transformer_candidate,
                base_storage_rows=current_storage_rows,
                phase=phase,
                iteration_no=iter_no,
                step_rows=transformer_steps,
                loadflow_attempts=loadflow_attempts,
                phase_search_history=phase_search_history,
            )

            idx = PHASES.index(phase)
            log_line(
                f"Iter {iter_no}, {phase}, export stage: "
                f"step={current_storage_rows[idx]['step_label']}, "
                f"Ppu={current_storage_rows[idx]['step_pu']}, "
                f"P={current_storage_rows[idx]['p_storage_kw']:.2f} kW, "
                f"P_export_after={current_storage_rows[idx].get('p_export_after_kw')}"
            )

        current_storage_rows = try_reduce_asymmetry(
            app=app,
            ldf=ldf,
            tr=tr,
            candidate=transformer_candidate,
            storage_rows=current_storage_rows,
            step_rows=transformer_steps,
            iteration_no=iter_no,
            loadflow_attempts=loadflow_attempts,
            phase_search_history=phase_search_history,
        )

        tr_row_final = collect_transformer_results(tr)

        p_export_ph = {
            ph: export_kw_from_p_tr(float(tr_row_final.get(f"P_tr_{ph}_kW") or 0.0))
            for ph in PHASES
        }

        p_import_ph = {
            ph: max(0.0, -float(tr_row_final.get(f"P_tr_{ph}_kW") or 0.0))
            if TRANSFORMER_EXPORT_POSITIVE
            else max(0.0, float(tr_row_final.get(f"P_tr_{ph}_kW") or 0.0))
            for ph in PHASES
        }

        i_vals = [float(tr_row_final.get(f"I_tr_{ph}_A") or 0.0) for ph in PHASES]
        i_avg = sum(i_vals) / 3.0 if i_vals else 0.0
        i_unb = max(abs(i - i_avg) for i in i_vals) / i_avg * 100.0 if i_avg > EPS else 0.0

        iteration_entry: Dict[str, Any] = {
            "iteration": iter_no,
            "P_tr_L1_kW": tr_row_final.get("P_tr_L1_kW"),
            "P_tr_L2_kW": tr_row_final.get("P_tr_L2_kW"),
            "P_tr_L3_kW": tr_row_final.get("P_tr_L3_kW"),
            "P_export_L1_kW": p_export_ph["L1"],
            "P_export_L2_kW": p_export_ph["L2"],
            "P_export_L3_kW": p_export_ph["L3"],
            "P_export_total_kW": sum(p_export_ph.values()),
            "P_import_L1_kW": p_import_ph["L1"],
            "P_import_L2_kW": p_import_ph["L2"],
            "P_import_L3_kW": p_import_ph["L3"],
            "P_import_total_kW": sum(p_import_ph.values()),
            "I_tr_L1_A": tr_row_final.get("I_tr_L1_A"),
            "I_tr_L2_A": tr_row_final.get("I_tr_L2_A"),
            "I_tr_L3_A": tr_row_final.get("I_tr_L3_A"),
            "I_unbalance_percent": i_unb,
            "I_neutral_A": tr_row_final.get("I_neutral_A"),
            "angle_I_tr_L1_deg": tr_row_final.get("angle_I_tr_L1_deg"),
            "angle_I_tr_L2_deg": tr_row_final.get("angle_I_tr_L2_deg"),
            "angle_I_tr_L3_deg": tr_row_final.get("angle_I_tr_L3_deg"),
            "loadflow_status": "0",
        }

        for row in current_storage_rows:
            ph = row["phase"]
            iteration_entry[f"P_storage_{ph}_kW"] = row.get("p_storage_kw", 0.0)
            iteration_entry[f"step_{ph}"] = row.get("step_label", "")
            iteration_entry[f"step_pu_{ph}"] = row.get("step_pu", "")
            iteration_entry[f"P_export_after_{ph}_kW"] = row.get("p_export_after_kw", "")
            iteration_entry[f"I_after_{ph}_A"] = row.get("i_after_a", "")

        iteration_history.append(iteration_entry)

        log_line(
            f"Iteracja {iter_no}: "
            f"P_export=({p_export_ph['L1']:.2f}, {p_export_ph['L2']:.2f}, {p_export_ph['L3']:.2f}) kW, "
            f"P_import=({p_import_ph['L1']:.2f}, {p_import_ph['L2']:.2f}, {p_import_ph['L3']:.2f}) kW, "
            f"I_unbalance={i_unb:.2f}%, "
            f"P_storage=({current_storage_rows[0]['p_storage_kw']:.2f}, "
            f"{current_storage_rows[1]['p_storage_kw']:.2f}, "
            f"{current_storage_rows[2]['p_storage_kw']:.2f}) kW"
        )

    result_tables = collect_results(
        app,
        tr,
        setup_data,
        current_storage_rows,
        iteration_history,
        loadflow_attempts,
        phase_search_history,
        transformer_steps,
    )
    export_to_excel(result_tables, OUT_FILE)

    ind = result_tables["indicators"][0]
    log_line(
        "Zakończono. "
        f"Umax={ind['Umax_pu']:.4f} pu, "
        f"Umin={ind['Umin_pu']:.4f} pu, "
        f"P_storage_total={ind['P_storage_total_kW']:.2f} kW, "
        f"eksport={ind['P_export_total_kW']:.2f} kW, "
        f"import={ind['P_import_total_kW']:.2f} kW, "
        f"I_neutral={ind['I_neutral_A']:.2f} A, "
        f"F1={ind['F1']:.6f}, "
        f"F2={ind['F2']:.6f}, "
        f"Fcelu_3={ind['Fcelu_3_weighted']:.6f}"
    )
    log_line(f"Zapisano wyniki do: {OUT_FILE}")


if __name__ == "__main__":
    run_local_transformer_flow_control()