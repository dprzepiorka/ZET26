"""
Jednoplikowy skrypt do porównania sterowania asymetrią w sieci nN APS.

Uruchamianie: tak jak wcześniejsze skrypty, z poziomu środowiska Pythona
DIgSILENT PowerFactory albo zewnętrznie po ustawieniu ścieżki do modułu
`powerfactory`. Skrypt:
  1) czyta nastawy początkowe z Excela,
  2) ustawia model PowerFactory,
  3) uruchamia wybrany przypadek lub wszystkie cztery przypadki,
  4) zapisuje wyniki do jednego pliku Excel.

Wybór przypadku jest w stałej CASE_TO_RUN na początku pliku:
  "base_no_control", "pso_global", "local_qu_storage_tr",
  "local_qu_storage_end" albo "all".
"""
from __future__ import annotations

import math
import os
import random
import sys
import time
import traceback
from copy import deepcopy
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    import numpy as np
except Exception:
    np = None  # PSO wymaga numpy; przypadki bez PSO mogą działać bez niego.

try:
    from openpyxl import Workbook, load_workbook
except Exception:
    Workbook = None
    load_workbook = None

try:
    from PSO import PSO
except Exception:
    PSO = None  # PSO wymaga zależności numerycznych dostępnych w środowisku docelowym.

# =============================================================================
# KONFIGURACJA - zwykle zmieniasz tylko tę sekcję
# =============================================================================
POWERFACTORY_PYTHON_PATH = r"C:\Program Files\DIgSILENT\PowerFactory 2026 SP1\Python\3.14"
USER = "KE"
PROJECT_NAME = "ELVTF_ZET_2026"

EXCEL_FILE = "dane.xlsx"
OUT_FILE = "wyniki_Codex.xlsx"

# "all" albo jeden z: base_no_control, pso_global, local_qu_storage_tr, local_qu_storage_end
CASE_TO_RUN = "base_no_control"

# Jeśli puste, skrypt bierze pierwszy ElmTr2 z aktywnego study case.
TRANSFORMER_NAME = ""
TRANSFORMER_CLASS = "ElmTr2"
TRANSFORMER_LV_SIDE = "buslv"

# Konwencja dla odczytanego P transformatora: True oznacza, że dodatnie P_tr_Lx
# jest eksportem z nN do SN. Jeśli w modelu PowerFactory znak jest odwrotny,
# ustaw False.
TRANSFORMER_EXPORT_POSITIVE = True

# Magazyn: dodatnie P_storage = ładowanie, czyli pobór mocy z sieci nN.
S_STORAGE_TOTAL_KVA = 60.0
P_STORAGE_TOTAL_MAX_KW = 60.0
S_STORAGE_PHASE_KVA = 20.0
P_STORAGE_PHASE_MAX_KW = 20.0
P_STORAGE_START_KW = 1.0
STORAGE_STEPS = [0, 25, 50, 75, 100]

# Q(U) falowników PV: Q > 0 oznacza indukcyjny pobór mocy biernej przez falownik.
QU_U_MIN_FULL = 0.95
QU_U_MIN_DEADBAND = 0.97
QU_U_MAX_DEADBAND = 1.03
QU_U_MAX_FULL = 1.08

# Reguła magazynu przy transformatorze.
TR_I_UNBALANCE_THRESHOLD_1 = 0.05
TR_I_UNBALANCE_THRESHOLD_2 = 0.10

# Reguła magazynu w węźle krytycznym.
END_U_START_PU = 1.03
END_U_STEP_25_PU = 1.04
END_U_STEP_50_PU = 1.05
END_U_STEP_75_PU = 1.06
END_U_UNBALANCE_BOOST_PU = 0.005

# Ograniczenia do raportowania naruszeń.
VOLTAGE_MIN_PU = 0.90
VOLTAGE_MAX_PU = 1.10
LOADING_MAX_PERCENT = 100.0

# PSO - optymalizuje tylko Q_PV oraz P_storage_L1/L2/L3.
PSO_N_PARTICLES = 40
PSO_N_ITER = 80
PSO_W = 0.70
PSO_C1 = 1.50
PSO_C2 = 1.50
PSO_RANDOM_SEED = 42
PSO_EVAL_DELAY = 0.01

# Wagi funkcji celu PSO, normalizowane względem base_no_control.
W_V = 0.20
W_VU = 0.30
W_IU = 0.25
W_EXP = 0.20
W_PEN = 0.05
EPS = 1e-9

PHASES = ("L1", "L2", "L3")
PF_PHASE = {"L1": "A", "L2": "B", "L3": "C"}
PHASE_ATTR_P_LOAD = {"L1": "plinir", "L2": "plinis", "L3": "plinit"}
PHASE_ATTR_Q_LOAD = {"L1": "qlinir", "L2": "qlinis", "L3": "qlinit"}

# =============================================================================
# NARZĘDZIA POWERFACTORY I EXCEL
# =============================================================================

def connect_powerfactory() -> Tuple[Any, Any]:
    if POWERFACTORY_PYTHON_PATH and POWERFACTORY_PYTHON_PATH not in sys.path:
        sys.path.append(POWERFACTORY_PYTHON_PATH)
    try:
        import powerfactory  # type: ignore
    except Exception as exc:
        raise RuntimeError(
            "Nie mogę zaimportować modułu powerfactory. Uruchom skrypt w środowisku "
            "DIgSILENT PowerFactory albo popraw POWERFACTORY_PYTHON_PATH."
        ) from exc
    app = powerfactory.GetApplicationExt(USER)
    if app is None:
        raise RuntimeError("PowerFactory nie zwrócił obiektu aplikacji.")
    if PROJECT_NAME:
        app.ActivateProject(PROJECT_NAME)
    ldf = app.GetFromStudyCase("ComLdf")
    if ldf is None:
        raise RuntimeError("Nie znaleziono ComLdf w aktywnym Study Case.")
    return app, ldf


def read_excel_sheet(path: str, sheet: str) -> List[Dict[str, Any]]:
    if load_workbook is None:
        raise RuntimeError("Do czytania Excela potrzebny jest pakiet openpyxl.")
    if not os.path.exists(path):
        raise FileNotFoundError(f"Nie znaleziono pliku Excel: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    out: List[Dict[str, Any]] = []
    for raw in rows[1:]:
        row = {headers[i]: raw[i] for i in range(min(len(headers), len(raw))) if headers[i]}
        if any(v not in (None, "") for v in row.values()):
            out.append(row)
    return out


def find_element(app: Any, name: str, pf_class: str) -> Any:
    if not name:
        return None
    try:
        objs = app.GetCalcRelevantObjects(f"{name}.{pf_class}")
        if objs:
            return objs[0]
    except Exception:
        pass
    try:
        for obj in app.GetCalcRelevantObjects(f"*.{pf_class}"):
            if getattr(obj, "loc_name", None) == name:
                return obj
    except Exception:
        pass
    return None


def get_attr(obj: Any, names: Sequence[str], default: Any = None) -> Any:
    for name in names:
        try:
            val = obj.GetAttribute(name)
            if val is not None:
                return val
        except Exception:
            continue
    return default


def get_float(obj: Any, names: Sequence[str], default: float = 0.0) -> float:
    val = get_attr(obj, names, default)
    try:
        return float(val)
    except Exception:
        return float(default)


def set_attr(obj: Any, names: Sequence[str], value: float) -> bool:
    for name in names:
        try:
            obj.SetAttribute(name, float(value))
            return True
        except Exception:
            continue
    return False


def run_loadflow(ldf: Any) -> None:
    rc = ldf.Execute()
    if rc not in (0, None):
        raise RuntimeError(f"Rozpływ mocy nie zbieżny albo błąd ComLdf, kod: {rc}")


def set_initial_model_from_excel(app: Any) -> None:
    """Ustawia obciążenia, generatory, PV i magazyny/statgeny z dane.xlsx."""
    for row in read_excel_sheet(EXCEL_FILE, "Loads"):
        elm = find_element(app, str(row.get("name", "")).strip(), "ElmLod")
        if elm is None:
            continue
        set_attr(elm, ["plinir"], float(row.get("P1") or 0.0))
        set_attr(elm, ["plinis"], float(row.get("P2") or 0.0))
        set_attr(elm, ["plinit"], float(row.get("P3") or 0.0))
        set_attr(elm, ["qlinir"], float(row.get("Q1") or 0.0))
        set_attr(elm, ["qlinis"], float(row.get("Q2") or 0.0))
        set_attr(elm, ["qlinit"], float(row.get("Q3") or 0.0))

    for sheet, cls in [("Generators", "ElmSym"), ("PV", "ElmPvsys"), ("StatGen", "ElmGenstat")]:
        for row in read_excel_sheet(EXCEL_FILE, sheet):
            elm = find_element(app, str(row.get("name", "")).strip(), cls)
            if elm is None:
                continue
            set_attr(elm, ["pgini"], float(row.get("P") or 0.0))
            set_attr(elm, ["qgini"], float(row.get("Q") or 0.0))


def load_storage_candidates() -> List[Dict[str, str]]:
    rows = read_excel_sheet(EXCEL_FILE, "StorageCandidates")
    return [
        {
            "node": str(r.get("node", "")).strip(),
            "elem_L1": str(r.get("elem_L1", r.get("elem_A", ""))).strip(),
            "elem_L2": str(r.get("elem_L2", r.get("elem_B", ""))).strip(),
            "elem_L3": str(r.get("elem_L3", r.get("elem_C", ""))).strip(),
        }
        for r in rows
    ]


def get_transformer(app: Any) -> Any:
    tr = find_element(app, TRANSFORMER_NAME, TRANSFORMER_CLASS) if TRANSFORMER_NAME else None
    if tr is not None:
        return tr
    trafos = app.GetCalcRelevantObjects(f"*.{TRANSFORMER_CLASS}")
    if not trafos:
        raise RuntimeError(f"Nie znaleziono transformatora klasy {TRANSFORMER_CLASS}.")
    return trafos[0]

# =============================================================================
# ODCZYTY MODELU I WSKAŹNIKI
# =============================================================================

def infer_phase_from_name(name: str) -> str:
    upper = name.upper()
    if upper.endswith(("F2", "L2", "_B", "-B", ".B")):
        return "L2"
    if upper.endswith(("F3", "L3", "_C", "-C", ".C")):
        return "L3"
    return "L1"


def collect_pv_objects(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    try:
        pvs = app.GetCalcRelevantObjects("*.ElmPvsys")
    except Exception:
        pvs = []
    for pv in pvs:
        name = getattr(pv, "loc_name", "PV")
        phase = infer_phase_from_name(name)
        p_kw = get_float(pv, ["pgini"], 0.0)
        q_kvar = get_float(pv, ["qgini"], 0.0)
        s_inv = get_float(pv, ["sgn", "sn", "snom"], max(abs(p_kw), 1.0))
        u_pu = get_float(pv, [f"m:u:{PF_PHASE[phase]}", "m:u"], 1.0)
        rows.append({"name": name, "phase": phase, "object": pv, "p_kw": p_kw, "q_kvar": q_kvar, "s_inv_kva": s_inv, "u_pu": u_pu})
    return rows


def collect_node_voltages(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for bus in app.GetCalcRelevantObjects("*.ElmTerm"):
        row = {"node": getattr(bus, "loc_name", "")}
        for ph in PHASES:
            pf_ph = PF_PHASE[ph]
            row[f"U_{ph}_pu"] = get_float(bus, [f"m:u:{pf_ph}"], math.nan)
            row[f"angle_{ph}_deg"] = get_float(bus, [f"m:phiu:{pf_ph}"], math.nan)
        row["kU2_percent"] = get_float(bus, ["m:ubfac"], math.nan)
        rows.append(row)
    return rows


def collect_transformer_results(tr: Any) -> Dict[str, Any]:
    row = {"transformer": getattr(tr, "loc_name", "")}
    for ph in PHASES:
        pf_ph = PF_PHASE[ph]
        i_ka = get_float(tr, [f"m:I:{TRANSFORMER_LV_SIDE}:{pf_ph}"], 0.0)
        p = get_attr(tr, [f"m:P:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:P:{pf_ph}", f"c:P:{pf_ph}"], None)
        q = get_attr(tr, [f"m:Q:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:Q:{pf_ph}", f"c:Q:{pf_ph}"], None)
        row[f"I_tr_{ph}_A"] = float(i_ka or 0.0) * 1000.0
        row[f"P_tr_{ph}_kW"] = None if p is None else (float(p) * 1000.0 if abs(float(p)) < 10.0 else float(p))
        row[f"Q_tr_{ph}_kvar"] = None if q is None else (float(q) * 1000.0 if abs(float(q)) < 10.0 else float(q))
    row["loading_percent"] = get_float(tr, ["c:loading"], math.nan)
    return row


def collect_branch_losses(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for line in app.GetCalcRelevantObjects("*.ElmLne"):
        rows.append(
            {
                "branch": getattr(line, "loc_name", ""),
                "P_loss_kW": 1000.0 * get_float(line, ["c:LossP"], 0.0),
                "Q_loss_kvar": 1000.0 * get_float(line, ["c:LossQ"], 0.0),
                "loading_percent": get_float(line, ["c:loading"], math.nan),
            }
        )
    return rows


def collect_results(app: Any, tr: Any, pv_rows: List[Dict[str, Any]], storage_rows: List[Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
    clean_pv = [{k: v for k, v in row.items() if k != "object"} for row in pv_rows]
    return {
        "node_voltages": collect_node_voltages(app),
        "transformer_phase_results": [collect_transformer_results(tr)],
        "pv_setpoints": clean_pv,
        "storage_setpoints": deepcopy(storage_rows),
        "branch_losses": collect_branch_losses(app),
    }


def calculate_indicators(raw: Dict[str, List[Dict[str, Any]]]) -> Dict[str, Any]:
    voltages = []
    phase_spreads = []
    ku2 = []
    for row in raw["node_voltages"]:
        vals = []
        for ph in PHASES:
            try:
                v = float(row[f"U_{ph}_pu"])
                if math.isfinite(v):
                    vals.append(v)
                    voltages.append(v)
            except Exception:
                pass
        if vals:
            phase_spreads.append(max(vals) - min(vals))
        try:
            k = float(row.get("kU2_percent"))
            if math.isfinite(k):
                ku2.append(k)
        except Exception:
            pass

    tr = raw["transformer_phase_results"][0] if raw["transformer_phase_results"] else {}
    i_vals = [float(tr.get(f"I_tr_{ph}_A") or 0.0) for ph in PHASES]
    i_avg = sum(i_vals) / 3.0 if i_vals else 0.0
    i_unb = max(abs(i - i_avg) for i in i_vals) / i_avg * 100.0 if i_avg > EPS else 0.0
    p_tr = [float(tr.get(f"P_tr_{ph}_kW") or 0.0) for ph in PHASES]
    p_export = sum(max(0.0, p if TRANSFORMER_EXPORT_POSITIVE else -p) for p in p_tr)
    p_loss = sum(float(row.get("P_loss_kW") or 0.0) for row in raw["branch_losses"])
    q_pv = [abs(float(row.get("q_kvar") or 0.0)) for row in raw["pv_setpoints"]]
    p_storage = [0.0, 0.0, 0.0]
    for idx, row in enumerate(raw["storage_setpoints"][:3]):
        p_storage[idx] = float(row.get("p_storage_kw") or 0.0)

    violations = []
    for row in raw["node_voltages"]:
        for ph in PHASES:
            try:
                u = float(row[f"U_{ph}_pu"])
                if u < VOLTAGE_MIN_PU or u > VOLTAGE_MAX_PU:
                    violations.append(f"U:{row.get('node')}:{ph}:{u:.4f}")
            except Exception:
                pass
    try:
        loading = float(tr.get("loading_percent"))
        if math.isfinite(loading) and loading > LOADING_MAX_PERCENT:
            violations.append(f"TR_loading:{loading:.2f}")
    except Exception:
        pass

    return {
        "Umax_pu": max(voltages) if voltages else math.nan,
        "Umin_pu": min(voltages) if voltages else math.nan,
        "Udev_mean_pu": sum(abs(v - 1.0) for v in voltages) / len(voltages) if voltages else math.nan,
        "Udev_max_pu": max(abs(v - 1.0) for v in voltages) if voltages else math.nan,
        "dU_phase_max_pu": max(phase_spreads) if phase_spreads else math.nan,
        "kU2_max_percent": max(ku2) if ku2 else math.nan,
        "kU2_mean_percent": sum(ku2) / len(ku2) if ku2 else math.nan,
        "I_tr_L1_A": i_vals[0],
        "I_tr_L2_A": i_vals[1],
        "I_tr_L3_A": i_vals[2],
        "I_unbalance_tr_percent": i_unb,
        "P_tr_L1_kW": p_tr[0],
        "P_tr_L2_kW": p_tr[1],
        "P_tr_L3_kW": p_tr[2],
        "P_export_total_kW": p_export,
        "P_loss_total_kW": p_loss,
        "Q_pv_abs_sum_kvar": sum(q_pv),
        "Q_pv_abs_max_kvar": max(q_pv) if q_pv else 0.0,
        "P_storage_L1_kW": p_storage[0],
        "P_storage_L2_kW": p_storage[1],
        "P_storage_L3_kW": p_storage[2],
        "P_storage_total_kW": sum(p_storage),
        "constraint_violations": ";".join(violations),
    }

# =============================================================================
# STEROWANIE PV I MAGAZYNEM
# =============================================================================

def qmax_available(p_pv_kw: float, s_inv_kva: float) -> float:
    p = abs(float(p_pv_kw))
    s = float(s_inv_kva)
    if p > s + 1e-9:
        raise ValueError(f"PV P={p:.3f} kW przekracza S_INV={s:.3f} kVA. To nie jest curtailment, tylko błąd danych.")
    return math.sqrt(max(0.0, s * s - p * p))


def qu_curve(u_pu: float, qmax: float) -> float:
    u = float(u_pu)
    if u <= QU_U_MIN_FULL:
        return -qmax
    if u < QU_U_MIN_DEADBAND:
        return -qmax * (QU_U_MIN_DEADBAND - u) / (QU_U_MIN_DEADBAND - QU_U_MIN_FULL)
    if u <= QU_U_MAX_DEADBAND:
        return 0.0
    if u < QU_U_MAX_FULL:
        return qmax * (u - QU_U_MAX_DEADBAND) / (QU_U_MAX_FULL - QU_U_MAX_DEADBAND)
    return qmax


def set_pv_q_zero(pv_rows: List[Dict[str, Any]]) -> None:
    for row in pv_rows:
        row["q_kvar"] = 0.0
        set_attr(row["object"], ["qgini", "qsetp"], 0.0)


def apply_local_qu(pv_rows: List[Dict[str, Any]]) -> None:
    for row in pv_rows:
        obj = row["object"]
        phase = row["phase"]
        row["u_pu"] = get_float(obj, [f"m:u:{PF_PHASE[phase]}", "m:u"], row.get("u_pu", 1.0))
        row["p_kw"] = get_float(obj, ["pgini"], row.get("p_kw", 0.0))
        row["s_inv_kva"] = get_float(obj, ["sgn", "sn", "snom"], row.get("s_inv_kva", abs(row["p_kw"])))
        qmax = qmax_available(row["p_kw"], row["s_inv_kva"])
        q = max(-qmax, min(qmax, qu_curve(row["u_pu"], qmax)))
        row["q_kvar"] = q
        set_attr(obj, ["qgini", "qsetp"], q)


def zero_storage_rows() -> List[Dict[str, Any]]:
    return [{"phase": ph, "step": 0, "p_storage_kw": 0.0, "q_storage_kvar": 0.0} for ph in PHASES]


def validate_storage(rows: List[Dict[str, Any]]) -> None:
    total_abs = 0.0
    for row in rows:
        p = float(row.get("p_storage_kw") or 0.0)
        q = float(row.get("q_storage_kvar") or 0.0)
        if abs(p) > P_STORAGE_PHASE_MAX_KW + 1e-9:
            raise ValueError(f"Magazyn {row['phase']} przekracza P fazowe: {p} kW")
        if p * p + q * q > S_STORAGE_PHASE_KVA * S_STORAGE_PHASE_KVA + 1e-9:
            raise ValueError(f"Magazyn {row['phase']} przekracza S fazowe: P={p}, Q={q}")
        total_abs += abs(p)
    if total_abs > P_STORAGE_TOTAL_MAX_KW + 1e-9:
        raise ValueError(f"Magazyn przekracza moc całkowitą: {total_abs} kW")


def choose_storage_candidate(candidates: List[Dict[str, str]], node_name: Optional[str] = None) -> Dict[str, str]:
    if not candidates:
        return {}
    if node_name:
        for cand in candidates:
            if cand.get("node") == node_name:
                return cand
    return candidates[0]


def apply_storage(app: Any, candidates: List[Dict[str, str]], rows: List[Dict[str, Any]], node_name: Optional[str] = None) -> None:
    validate_storage(rows)
    cand = choose_storage_candidate(candidates, node_name)
    for row in rows:
        ph = row["phase"]
        name = cand.get(f"elem_{ph}", "")
        if not name:
            continue
        p = float(row.get("p_storage_kw") or 0.0)  # + ładowanie
        q = float(row.get("q_storage_kvar") or 0.0)
        elm_lod = find_element(app, name, "ElmLod")
        if elm_lod is not None:
            set_attr(elm_lod, [PHASE_ATTR_P_LOAD[ph], "plini"], p)
            set_attr(elm_lod, [PHASE_ATTR_Q_LOAD[ph], "qlini"], q)
            continue
        elm_gen = find_element(app, name, "ElmGenstat") or find_element(app, name, "ElmSym") or find_element(app, name, "ElmPvsys")
        if elm_gen is not None:
            set_attr(elm_gen, ["pgini"], -p)
            set_attr(elm_gen, ["qgini"], -q)


def export_step(p_export_kw: float) -> int:
    if p_export_kw <= P_STORAGE_START_KW:
        return 0
    if p_export_kw <= 0.25 * P_STORAGE_PHASE_MAX_KW:
        return 1
    if p_export_kw <= 0.50 * P_STORAGE_PHASE_MAX_KW:
        return 2
    if p_export_kw <= 0.75 * P_STORAGE_PHASE_MAX_KW:
        return 3
    return 4


def transformer_storage_rule(tr_row: Dict[str, Any]) -> List[Dict[str, Any]]:
    for ph in PHASES:
        if tr_row.get(f"P_tr_{ph}_kW") is None:
            raise ValueError(
                "local_qu_storage_tr wymaga fazowych mocy czynnych transformatora P_tr_L1/L2/L3. "
                "Sam moduł prądu nie wystarcza do wykrycia kierunku przepływu."
            )
    i_vals = [float(tr_row.get(f"I_tr_{ph}_A") or 0.0) for ph in PHASES]
    i_avg = sum(i_vals) / 3.0 if i_vals else 0.0
    rows: List[Dict[str, Any]] = []
    for ph, i in zip(PHASES, i_vals):
        p_tr = float(tr_row[f"P_tr_{ph}_kW"])
        p_export = p_tr if TRANSFORMER_EXPORT_POSITIVE else -p_tr
        step = export_step(p_export)
        d_i = (i - i_avg) / i_avg if i_avg > EPS else 0.0
        if d_i > TR_I_UNBALANCE_THRESHOLD_2:
            step += 2
        elif d_i > TR_I_UNBALANCE_THRESHOLD_1:
            step += 1
        step = int(max(-4, min(4, step)))
        p = step / 4.0 * P_STORAGE_PHASE_MAX_KW
        rows.append({"phase": ph, "step": step, "p_storage_kw": p, "q_storage_kvar": 0.0})
    validate_storage(rows)
    return rows


def select_critical_node(base_raw: Dict[str, List[Dict[str, Any]]]) -> str:
    best_node = None
    best_u = -math.inf
    for row in base_raw["node_voltages"]:
        vals = [float(row.get(f"U_{ph}_pu") or -math.inf) for ph in PHASES]
        u = max(vals)
        if u > best_u:
            best_u = u
            best_node = row["node"]
    if best_node is None:
        raise RuntimeError("Nie udało się wybrać węzła krytycznego z wyników bazowych.")
    return str(best_node)


def node_voltage_row(raw: Dict[str, List[Dict[str, Any]]], node: str) -> Dict[str, Any]:
    for row in raw["node_voltages"]:
        if row.get("node") == node:
            return row
    raise RuntimeError(f"Nie znaleziono napięć węzła krytycznego: {node}")


def end_node_storage_rule(v_row: Dict[str, Any]) -> List[Dict[str, Any]]:
    u_vals = [float(v_row.get(f"U_{ph}_pu") or 1.0) for ph in PHASES]
    u_avg = sum(u_vals) / 3.0
    rows: List[Dict[str, Any]] = []
    for ph, u in zip(PHASES, u_vals):
        if u <= END_U_START_PU:
            step = 0
        elif u <= END_U_STEP_25_PU:
            step = 1
        elif u <= END_U_STEP_50_PU:
            step = 2
        elif u <= END_U_STEP_75_PU:
            step = 3
        else:
            step = 4
        if u - u_avg > END_U_UNBALANCE_BOOST_PU:
            step += 1
        step = int(max(0, min(4, step)))
        p = step / 4.0 * P_STORAGE_PHASE_MAX_KW
        rows.append({"phase": ph, "step": step, "p_storage_kw": p, "q_storage_kvar": 0.0})
    validate_storage(rows)
    return rows

def require_pso_dependencies() -> None:
    if np is None or PSO is None:
        raise RuntimeError("Przypadek pso_global wymaga pakietu numpy i dostępnego pliku PSO.py.")


# =============================================================================
# PRZYPADKI BADAWCZE
# =============================================================================

def reset_model_for_case(app: Any, ldf: Any, candidates: List[Dict[str, str]]) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    set_initial_model_from_excel(app)
    pv_rows = collect_pv_objects(app)
    set_pv_q_zero(pv_rows)
    storage_rows = zero_storage_rows()
    apply_storage(app, candidates, storage_rows)
    run_loadflow(ldf)
    return pv_rows, storage_rows


def run_base_no_control(app: Any, ldf: Any, tr: Any, candidates: List[Dict[str, str]]) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    pv_rows, storage_rows = reset_model_for_case(app, ldf, candidates)
    raw = collect_results(app, tr, pv_rows, storage_rows)
    return raw, calculate_indicators(raw)


def run_local_qu_storage_tr(app: Any, ldf: Any, tr: Any, candidates: List[Dict[str, str]]) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    pv_rows, _ = reset_model_for_case(app, ldf, candidates)
    apply_local_qu(pv_rows)
    run_loadflow(ldf)
    raw_before_storage = collect_results(app, tr, pv_rows, zero_storage_rows())
    storage_rows = transformer_storage_rule(raw_before_storage["transformer_phase_results"][0])
    apply_storage(app, candidates, storage_rows)
    run_loadflow(ldf)
    raw = collect_results(app, tr, pv_rows, storage_rows)
    return raw, calculate_indicators(raw)


def run_local_qu_storage_end(app: Any, ldf: Any, tr: Any, candidates: List[Dict[str, str]], base_raw: Dict[str, Any]) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    critical_node = select_critical_node(base_raw)
    pv_rows, _ = reset_model_for_case(app, ldf, candidates)
    apply_local_qu(pv_rows)
    run_loadflow(ldf)
    raw_before_storage = collect_results(app, tr, pv_rows, zero_storage_rows())
    storage_rows = end_node_storage_rule(node_voltage_row(raw_before_storage, critical_node))
    apply_storage(app, candidates, storage_rows, node_name=critical_node)
    run_loadflow(ldf)
    raw = collect_results(app, tr, pv_rows, storage_rows)
    raw["case_meta"] = [{"critical_node": critical_node}]
    return raw, calculate_indicators(raw)


def pso_variables(pv_rows: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], Any, Any]:
    require_pso_dependencies()
    variables: List[Dict[str, Any]] = []
    lb: List[float] = []
    ub: List[float] = []
    for idx, row in enumerate(pv_rows):
        qmax = qmax_available(row["p_kw"], row["s_inv_kva"])
        variables.append({"kind": "pv_q", "index": idx, "qmax": qmax, "name": row["name"], "phase": row["phase"]})
        lb.append(-qmax)
        ub.append(qmax)
    for ph in PHASES:
        variables.append({"kind": "storage_p", "phase": ph})
        lb.append(-P_STORAGE_PHASE_MAX_KW)
        ub.append(P_STORAGE_PHASE_MAX_KW)
    return variables, np.array(lb, dtype=float), np.array(ub, dtype=float)


def apply_pso_vector(app: Any, candidates: List[Dict[str, str]], x: np.ndarray, variables: List[Dict[str, Any]], pv_rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    storage_rows = zero_storage_rows()
    for value, var in zip(x, variables):
        if var["kind"] == "pv_q":
            q = max(-var["qmax"], min(var["qmax"], float(value)))
            row = pv_rows[var["index"]]
            row["q_kvar"] = q
            set_attr(row["object"], ["qgini", "qsetp"], q)
        else:
            idx = PHASES.index(var["phase"])
            p = max(-P_STORAGE_PHASE_MAX_KW, min(P_STORAGE_PHASE_MAX_KW, float(value)))
            storage_rows[idx]["p_storage_kw"] = p
    apply_storage(app, candidates, storage_rows)
    return storage_rows


def pso_objective_value(ind: Dict[str, Any], base_ind: Dict[str, Any]) -> float:
    def norm(key: str) -> float:
        return abs(float(ind.get(key) or 0.0)) / max(abs(float(base_ind.get(key) or 0.0)), EPS)

    penalty = 1.0 if ind.get("constraint_violations") else 0.0
    return W_V * norm("Udev_mean_pu") + W_VU * norm("kU2_mean_percent") + W_IU * norm("I_unbalance_tr_percent") + W_EXP * norm("P_export_total_kW") + W_PEN * penalty


def run_pso_global(app: Any, ldf: Any, tr: Any, candidates: List[Dict[str, str]], base_ind: Dict[str, Any]) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    require_pso_dependencies()
    random.seed(PSO_RANDOM_SEED)
    np.random.seed(PSO_RANDOM_SEED)
    pv_rows, _ = reset_model_for_case(app, ldf, candidates)
    variables, lb, ub = pso_variables(pv_rows)

    def obj(x: np.ndarray) -> float:
        try:
            storage_rows = apply_pso_vector(app, candidates, np.asarray(x, dtype=float), variables, pv_rows)
            run_loadflow(ldf)
            raw = collect_results(app, tr, pv_rows, storage_rows)
            ind = calculate_indicators(raw)
            return pso_objective_value(ind, base_ind)
        except Exception:
            with open("pso_eval_errors.log", "a", encoding="utf-8") as fh:
                fh.write(traceback.format_exc() + "\n")
            return float("inf")
        finally:
            # Każda kolejna ocena nadpisuje wszystkie zmienne, więc nie trzeba odtwarzać stanu.
            if PSO_EVAL_DELAY:
                time.sleep(PSO_EVAL_DELAY)

    opt = PSO(obj, PSO_N_PARTICLES, len(variables), lb, ub, PSO_N_ITER, PSO_W, PSO_C1, PSO_C2)
    res = opt.optimize()
    storage_rows = apply_pso_vector(app, candidates, np.asarray(res["gbest"], dtype=float), variables, pv_rows)
    run_loadflow(ldf)
    raw = collect_results(app, tr, pv_rows, storage_rows)
    raw["pso_best"] = [{"gbest_val": res["gbest_val"], "variables": str(variables), "gbest": str(list(map(float, res["gbest"])))}]
    return raw, calculate_indicators(raw)

# =============================================================================
# EKSPORT DO EXCELA
# =============================================================================

def safe_sheet_name(name: str) -> str:
    bad = "[]:*?/\\"
    for ch in bad:
        name = name.replace(ch, "_")
    return name[:31]


def write_rows(ws: Any, rows: Iterable[Dict[str, Any]]) -> None:
    rows = list(rows)
    if not rows:
        ws.append(["empty"])
        return
    headers: List[str] = []
    for row in rows:
        for key in row.keys():
            if key not in headers:
                headers.append(key)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])


def indicators_as_rows(all_indicators: Dict[str, Dict[str, Any]]) -> List[Dict[str, Any]]:
    keys: List[str] = []
    for ind in all_indicators.values():
        for key in ind:
            if key not in keys:
                keys.append(key)
    rows = []
    for case, ind in all_indicators.items():
        row = {"case": case}
        row.update({key: ind.get(key, "") for key in keys})
        rows.append(row)
    return rows


def comparison_to_base(all_indicators: Dict[str, Dict[str, Any]]) -> List[Dict[str, Any]]:
    minimized = ["Udev_mean_pu", "Udev_max_pu", "dU_phase_max_pu", "kU2_max_percent", "kU2_mean_percent", "I_unbalance_tr_percent", "P_export_total_kW", "P_loss_total_kW"]
    base = all_indicators.get("base_no_control", {})
    rows = []
    for case, ind in all_indicators.items():
        if case == "base_no_control":
            continue
        for key in minimized:
            b = float(base.get(key) or 0.0)
            x = float(ind.get(key) or 0.0)
            imp = (b - x) / b * 100.0 if abs(b) > EPS else ""
            rows.append({"case": case, "indicator": key, "base": b, "case_value": x, "improvement_vs_base_percent": imp})
    return rows


def comparison_to_pso(all_indicators: Dict[str, Dict[str, Any]]) -> List[Dict[str, Any]]:
    minimized = ["Udev_mean_pu", "Udev_max_pu", "dU_phase_max_pu", "kU2_max_percent", "kU2_mean_percent", "I_unbalance_tr_percent", "P_export_total_kW", "P_loss_total_kW"]
    base = all_indicators.get("base_no_control", {})
    pso = all_indicators.get("pso_global", {})
    rows = []
    for case in ["local_qu_storage_tr", "local_qu_storage_end"]:
        if case not in all_indicators:
            continue
        ind = all_indicators[case]
        for key in minimized:
            b = float(base.get(key) or 0.0)
            p = float(pso.get(key) or 0.0)
            x = float(ind.get(key) or 0.0)
            den = b - p
            eff = (b - x) / den * 100.0 if abs(den) > EPS else ""
            rows.append({"case": case, "indicator": key, "base": b, "pso": p, "local": x, "effectiveness_vs_pso_percent": eff})
    return rows


def export_to_excel(all_raw: Dict[str, Dict[str, Any]], all_indicators: Dict[str, Dict[str, Any]]) -> None:
    if Workbook is None:
        raise RuntimeError("Do zapisu wyników do Excela potrzebny jest pakiet openpyxl.")
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("indicators_all_cases")
    write_rows(ws, indicators_as_rows(all_indicators))

    ws = wb.create_sheet("comparison_to_base")
    write_rows(ws, comparison_to_base(all_indicators))

    if "pso_global" in all_indicators:
        ws = wb.create_sheet("comparison_to_pso")
        write_rows(ws, comparison_to_pso(all_indicators))

    for case, raw in all_raw.items():
        prefix = {
            "base_no_control": "base",
            "pso_global": "pso",
            "local_qu_storage_tr": "tr",
            "local_qu_storage_end": "end",
        }.get(case, case[:8])
        for table, rows in raw.items():
            ws = wb.create_sheet(safe_sheet_name(f"{prefix}_{table}"))
            write_rows(ws, rows)

    wb.save(OUT_FILE)

# =============================================================================
# MAIN
# =============================================================================

def run_study() -> None:
    app, ldf = connect_powerfactory()
    tr = get_transformer(app)
    candidates = load_storage_candidates()
    cases = ["base_no_control", "pso_global", "local_qu_storage_tr", "local_qu_storage_end"] if CASE_TO_RUN == "all" else [CASE_TO_RUN]

    all_raw: Dict[str, Dict[str, Any]] = {}
    all_indicators: Dict[str, Dict[str, Any]] = {}

    # Baza jest potrzebna do normalizacji PSO i wyboru węzła krytycznego, więc licz ją zawsze,
    # jeśli uruchamiany jest PSO albo local_qu_storage_end.
    need_base = CASE_TO_RUN == "all" or CASE_TO_RUN in {"pso_global", "local_qu_storage_end"}
    base_raw = None
    base_ind = None
    if need_base and "base_no_control" not in cases:
        base_raw, base_ind = run_base_no_control(app, ldf, tr, candidates)

    for case in cases:
        print(f"Uruchamiam przypadek: {case}")
        if case == "base_no_control":
            raw, ind = run_base_no_control(app, ldf, tr, candidates)
            base_raw, base_ind = raw, ind
        elif case == "pso_global":
            if base_ind is None:
                base_raw, base_ind = run_base_no_control(app, ldf, tr, candidates)
            raw, ind = run_pso_global(app, ldf, tr, candidates, base_ind)
        elif case == "local_qu_storage_tr":
            raw, ind = run_local_qu_storage_tr(app, ldf, tr, candidates)
        elif case == "local_qu_storage_end":
            if base_raw is None:
                base_raw, base_ind = run_base_no_control(app, ldf, tr, candidates)
            raw, ind = run_local_qu_storage_end(app, ldf, tr, candidates, base_raw)
        else:
            raise ValueError(f"Nieznany przypadek CASE_TO_RUN: {case}")
        all_raw[case] = raw
        all_indicators[case] = ind
        print(f"  Umax={ind['Umax_pu']:.4f} pu, eksport={ind['P_export_total_kW']:.2f} kW, naruszenia={ind['constraint_violations']}")

    # Jeśli liczono bazę pomocniczo, ale użytkownik uruchamia pojedynczy PSO/end, dopisz bazę do zestawień.
    if base_raw is not None and base_ind is not None and "base_no_control" not in all_raw:
        all_raw = {"base_no_control": base_raw, **all_raw}
        all_indicators = {"base_no_control": base_ind, **all_indicators}

    export_to_excel(all_raw, all_indicators)
    print(f"Zapisano wyniki do: {OUT_FILE}")


if __name__ == "__main__":
    run_study()
