from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Iterable, List

from openpyxl import Workbook, load_workbook


DEFAULT_INPUT_SHEETS = ["study_config", "loads", "pv_units", "storage", "object_map", "control_params"]


def _sheet_rows_to_dicts(ws) -> List[Dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    items: List[Dict[str, Any]] = []
    for raw in rows[1:]:
        if raw is None:
            continue
        row = {headers[idx]: value for idx, value in enumerate(raw) if idx < len(headers) and headers[idx]}
        if any(v is not None and str(v).strip() != "" for v in row.values()):
            items.append(row)
    return items


def _coerce_number(v: Any) -> Any:
    if isinstance(v, (int, float)):
        return v
    if v is None:
        return v
    s = str(v).strip()
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return v


def read_input_excel(path: str) -> Dict[str, Any]:
    wb = load_workbook(path, data_only=True)
    data: Dict[str, Any] = {
        "study_config": {},
        "loads": [],
        "pv_units": [],
        "storage": [],
        "object_map": [],
        "control_params": {},
    }

    if "study_config" in wb.sheetnames:
        rows = _sheet_rows_to_dicts(wb["study_config"])
        for r in rows:
            key = r.get("key")
            value = r.get("value")
            if key is not None:
                data["study_config"][str(key)] = _coerce_number(value)

    for name in ["loads", "pv_units", "storage", "object_map"]:
        if name in wb.sheetnames:
            data[name] = _sheet_rows_to_dicts(wb[name])

    if "control_params" in wb.sheetnames:
        rows = _sheet_rows_to_dicts(wb["control_params"])
        for r in rows:
            group = str(r.get("group", "general"))
            key = r.get("key")
            value = r.get("value")
            if key is None:
                continue
            data["control_params"].setdefault(group, {})[str(key)] = _coerce_number(value)

    return data


def _write_rows(ws, rows: Iterable[Dict[str, Any]]) -> None:
    rows = list(rows)
    if not rows:
        ws.append(["info"])
        ws.append(["no_data"])
        return
    headers: List[str] = []
    for r in rows:
        for k in r.keys():
            if k not in headers:
                headers.append(k)
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])


def _safe_sheet_name(name: str) -> str:
    forbidden = set("[]:*?/\\")
    clean = "".join("_" if c in forbidden else c for c in name)
    return clean[:31]


def write_case_results_excel(path: str, sheets: Dict[str, List[Dict[str, Any]]]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(_safe_sheet_name(sheet_name))
        _write_rows(ws, rows)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_summary_excel(path: str, summary_sheets: Dict[str, List[Dict[str, Any]]]) -> None:
    write_case_results_excel(path, summary_sheets)


def create_sample_input_excel(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "study_config"
    ws.append(["key", "value"])
    ws.append(["case_name", "all"])
    ws.append(["output_dir", "results"])
    ws.append(["output_file", "results.xlsx"])
    ws.append(["demo_mode", "true"])

    ws = wb.create_sheet("loads")
    ws.append(["object_id", "node_id", "L1_kW", "L2_kW", "L3_kW"])
    ws.append(["load_1", "N1", 6.0, 7.0, 7.0])

    ws = wb.create_sheet("pv_units")
    ws.append(["object_id", "node_id", "L1_kW", "L2_kW", "L3_kW", "S_inv_kVA"])
    ws.append(["pv_1", "N2", 28.0, 27.0, 25.0, 30.0])

    ws = wb.create_sheet("storage")
    ws.append(["object_id", "node_id", "location", "S_phase_kVA", "P_phase_max_kW"])
    ws.append(["storage_1", "TR_NODE", "transformer", 20.0, 20.0])

    ws = wb.create_sheet("object_map")
    ws.append(["object_id", "pf_path", "node_id", "object_type"])
    ws.append(["TR", "Network Model\\Grid\\TR", "TR_NODE", "transformer"])
    ws.append(["N1", "Network Model\\Grid\\Busbar N1", "N1", "node"])
    ws.append(["N2", "Network Model\\Grid\\Busbar N2", "N2", "node"])

    ws = wb.create_sheet("control_params")
    ws.append(["group", "key", "value"])
    ws.append(["storage", "s_total_kva", 60.0])
    ws.append(["storage", "p_total_max_kw", 60.0])
    ws.append(["storage", "s_phase_kva", 20.0])
    ws.append(["storage", "p_phase_max_kw", 20.0])
    ws.append(["storage", "p_start_kw", 1.0])
    ws.append(["transformer_rule", "current_unbalance_threshold_1", 0.05])
    ws.append(["transformer_rule", "current_unbalance_threshold_2", 0.10])
    ws.append(["end_node_rule", "u_start_pu", 1.03])
    ws.append(["end_node_rule", "u_step_25_pu", 1.04])
    ws.append(["end_node_rule", "u_step_50_pu", 1.05])
    ws.append(["end_node_rule", "u_step_75_pu", 1.06])
    ws.append(["end_node_rule", "u_unbalance_boost_pu", 0.005])
    ws.append(["qu_curve", "u_min_full", 0.95])
    ws.append(["qu_curve", "u_min_deadband", 0.97])
    ws.append(["qu_curve", "u_max_deadband", 1.03])
    ws.append(["qu_curve", "u_max_full", 1.08])

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
