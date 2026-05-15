from __future__ import annotations

import cmath
import math
import os
import random
import sys
import time
from copy import deepcopy
from typing import Any, Dict, Iterable, List, Sequence, Tuple

import numpy as np

try:
    from openpyxl import Workbook, load_workbook
except Exception:
    Workbook = None
    load_workbook = None

from PSO import PSO


# =============================================================================
# PARAMETRY DOMYŚLNE
# =============================================================================

POWERFACTORY_PYTHON_PATH = r"C:\Program Files\DIgSILENT\PowerFactory 2026 SP1\Python\3.14"
USER = "KE"
PROJECT_NAME = "ELVTF_ZET_2026"

EXCEL_FILE = r"input_data.xlsx"
OUT_FILE = r"results\results_pso_global_critical.xlsx"
LOG_FILE = r"results\pso_global_critical.log"

TRANSFORMER_NAME = ""
TRANSFORMER_CLASS = "ElmTr2"
TRANSFORMER_LV_SIDE = "buslv"

TRANSFORMER_EXPORT_POSITIVE = True

VOLTAGE_MIN_PU = 0.90
TARGET_VOLTAGE_PU = 1.00
VOLTAGE_MAX_PU = 1.10

U_MIN_ALLOWED_OBJ = 0.95
U_MAX_ALLOWED_OBJ = 1.05
U_TOL_PU = 0.01

P_EXPORT_REF_KW = 100.0

K_U = 100.0
K_L = 100.0
K_T = 100.0

LOADING_MAX_PERCENT = 100.0
LINE_LOADING_MAX_PERCENT = 100.0

P_STORAGE_TOTAL_MAX_KW = 90.0
P_STORAGE_PHASE_MAX_KW = 30.0
ALLOW_STORAGE_CHARGE = True
ALLOW_STORAGE_DISCHARGE = True
KEEP_INITIAL_STORAGE_Q = True

EPS = 1e-9
TG_PHI_MAX = 0.33

# PSO
N_PARTICLES = 20
N_ITER = 10
PSO_W = 0.7
PSO_C1 = 1.5
PSO_C2 = 1.5
PSO_RANDOM_SEED = 42
PSO_AUTOSAVE_EVERY = 0
PSO_AUTOSAVE_PATH = "results/pso_global_critical_checkpoint.npz"
PSO_EVAL_DELAY = 0.0
NONCONVERGENCE_PENALTY = 1e6

PHASES = ("L1", "L2", "L3")
PF_PHASE = {"L1": "A", "L2": "B", "L3": "C"}

PHASE_ATTR_P_LOAD = {"L1": "plinir", "L2": "plinis", "L3": "plinit"}
PHASE_ATTR_Q_LOAD = {"L1": "qlinir", "L2": "qlinis", "L3": "qlinit"}

EXCEL_CACHE: Dict[str, List[Dict[str, Any]]] = {}
CONTROL_CONFIG: Dict[str, Any] = {}

PSO_CONVERGENCE_ROWS: List[Dict[str, Any]] = []
PSO_PARTICLE_EVAL_ROWS: List[Dict[str, Any]] = []
CURRENT_ITER_EVALS: List[Dict[str, Any]] = []
PARTICLE_EVAL_COUNTER = 0


# =============================================================================
# LOG
# =============================================================================

def log_line(msg: str) -> None:
    ts = time.strftime("%Y-%m-%d %H:%M:%S")
    line = f"[{ts}] {msg}"
    print(line)
    try:
        log_dir = os.path.dirname(LOG_FILE)
        if log_dir:
            os.makedirs(log_dir, exist_ok=True)
        with open(LOG_FILE, "a", encoding="utf-8") as fh:
            fh.write(line + "\n")
    except Exception:
        pass


# =============================================================================
# CONFIG
# =============================================================================

def parse_bool_pl(value: Any) -> bool:
    txt = str(value).strip().lower()
    return txt in {"1", "true", "yes", "y", "tak", "prawda"}


def parse_config_value(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, (int, float, bool)):
        return value

    txt = str(value).strip()
    if not txt:
        return ""

    low = txt.lower()
    if low in {"prawda", "true", "tak", "yes", "1"}:
        return True
    if low in {"fałsz", "falsz", "false", "nie", "no", "0"}:
        return False

    try:
        if "." in txt:
            return float(txt)
        return int(txt)
    except Exception:
        return txt


def load_control_config() -> Dict[str, Any]:
    rows = excel_sheet("ControlConfig")
    cfg: Dict[str, Any] = {}

    for row in rows:
        key = str(row.get("parameter", "")).strip()
        if not key:
            continue
        cfg[key] = parse_config_value(row.get("value"))

    return cfg


def cfg_float(name: str, default: float) -> float:
    val = CONTROL_CONFIG.get(name, default)
    try:
        return float(val)
    except Exception:
        return float(default)


def cfg_int(name: str, default: int) -> int:
    val = CONTROL_CONFIG.get(name, default)
    try:
        return int(float(val))
    except Exception:
        return int(default)


def cfg_bool(name: str, default: bool) -> bool:
    val = CONTROL_CONFIG.get(name, default)
    if isinstance(val, bool):
        return val
    return parse_bool_pl(val) if val is not None else default


def apply_runtime_config() -> None:
    global VOLTAGE_MIN_PU, TARGET_VOLTAGE_PU, VOLTAGE_MAX_PU
    global U_MIN_ALLOWED_OBJ, U_MAX_ALLOWED_OBJ, U_TOL_PU
    global P_EXPORT_REF_KW
    global K_U, K_L, K_T
    global LOADING_MAX_PERCENT
    global P_STORAGE_TOTAL_MAX_KW, P_STORAGE_PHASE_MAX_KW
    global ALLOW_STORAGE_CHARGE, ALLOW_STORAGE_DISCHARGE, KEEP_INITIAL_STORAGE_Q
    global NONCONVERGENCE_PENALTY
    global N_PARTICLES, N_ITER, PSO_W, PSO_C1, PSO_C2, PSO_RANDOM_SEED

    VOLTAGE_MIN_PU = cfg_float("VOLTAGE_MIN_PU", VOLTAGE_MIN_PU)
    TARGET_VOLTAGE_PU = cfg_float("UREF_PU", cfg_float("TARGET_VOLTAGE_PU", TARGET_VOLTAGE_PU))
    VOLTAGE_MAX_PU = cfg_float("VOLTAGE_MAX_PU", VOLTAGE_MAX_PU)

    U_MIN_ALLOWED_OBJ = cfg_float("U_MIN_ALLOWED_PU", cfg_float("U_MIN_ALLOWED_OBJ", U_MIN_ALLOWED_OBJ))
    U_MAX_ALLOWED_OBJ = cfg_float("U_MAX_ALLOWED_PU", cfg_float("U_MAX_ALLOWED_OBJ", U_MAX_ALLOWED_OBJ))
    U_TOL_PU = cfg_float("U_TOL_PU", U_TOL_PU)

    P_EXPORT_REF_KW = cfg_float("P_EXPORT_REF_KW", P_EXPORT_REF_KW)

    K_U = cfg_float("K_U", K_U)
    K_L = cfg_float("K_L", K_L)
    K_T = cfg_float("K_T", K_T)

    LOADING_MAX_PERCENT = cfg_float("LOADING_MAX_PERCENT", LOADING_MAX_PERCENT)

    P_STORAGE_TOTAL_MAX_KW = cfg_float("P_STORAGE_TOTAL_MAX_KW", P_STORAGE_TOTAL_MAX_KW)
    P_STORAGE_PHASE_MAX_KW = cfg_float("P_STORAGE_PHASE_MAX_KW", P_STORAGE_PHASE_MAX_KW)
    ALLOW_STORAGE_CHARGE = cfg_bool("ALLOW_STORAGE_CHARGE", ALLOW_STORAGE_CHARGE)
    ALLOW_STORAGE_DISCHARGE = cfg_bool("ALLOW_STORAGE_DISCHARGE", ALLOW_STORAGE_DISCHARGE)
    KEEP_INITIAL_STORAGE_Q = cfg_bool("KEEP_INITIAL_STORAGE_Q", KEEP_INITIAL_STORAGE_Q)

    NONCONVERGENCE_PENALTY = cfg_float("NONCONVERGENCE_PENALTY", NONCONVERGENCE_PENALTY)

    N_PARTICLES = cfg_int("PSO_N_PARTICLES", N_PARTICLES)
    N_ITER = cfg_int("PSO_N_ITER", N_ITER)
    PSO_W = cfg_float("PSO_W", PSO_W)
    PSO_C1 = cfg_float("PSO_C1", PSO_C1)
    PSO_C2 = cfg_float("PSO_C2", PSO_C2)
    PSO_RANDOM_SEED = cfg_int("PSO_RANDOM_SEED", PSO_RANDOM_SEED)


# =============================================================================
# POWERFACTORY
# =============================================================================

def connect_powerfactory() -> Tuple[Any, Any]:
    if POWERFACTORY_PYTHON_PATH and POWERFACTORY_PYTHON_PATH not in sys.path:
        sys.path.append(POWERFACTORY_PYTHON_PATH)

    try:
        import powerfactory  # type: ignore
    except Exception as exc:
        raise RuntimeError(
            "Nie mogę zaimportować modułu powerfactory. "
            "Uruchom skrypt w środowisku DIgSILENT PowerFactory albo popraw POWERFACTORY_PYTHON_PATH."
        ) from exc

    app = powerfactory.GetApplicationExt(USER)
    if app is None:
        raise RuntimeError("PowerFactory nie zwrócił obiektu aplikacji.")

    if PROJECT_NAME:
        app.ActivateProject(PROJECT_NAME)

    ldf = app.GetFromStudyCase("ComLdf")
    if ldf is None:
        raise RuntimeError("Nie znaleziono ComLdf w aktywnym Study Case.")

    return app, ldf


def find_element(app: Any, name: str, pf_class: str) -> Any:
    if not name:
        return None

    try:
        objs = app.GetCalcRelevantObjects(f"{name}.{pf_class}")
        if objs:
            return objs[0]
    except Exception:
        pass

    try:
        for obj in app.GetCalcRelevantObjects(f"*.{pf_class}"):
            if getattr(obj, "loc_name", None) == name:
                return obj
    except Exception:
        pass

    return None


def get_attr(obj: Any, names: Sequence[str], default: Any = None) -> Any:
    for name in names:
        try:
            val = obj.GetAttribute(name)
            if val is not None:
                return val
        except Exception:
            continue
    return default


def get_float(obj: Any, names: Sequence[str], default: float = 0.0) -> float:
    val = get_attr(obj, names, default)
    try:
        return float(val)
    except Exception:
        return float(default)


def set_attr(obj: Any, names: Sequence[str], value: Any) -> bool:
    for name in names:
        try:
            obj.SetAttribute(name, value)
            return True
        except Exception:
            continue
    return False


def try_run_loadflow(ldf: Any) -> Tuple[bool, Any]:
    try:
        rc = ldf.Execute()
        return rc in (0, None), rc
    except Exception as exc:
        return False, str(exc)


def get_transformer(app: Any) -> Any:
    tr = find_element(app, TRANSFORMER_NAME, TRANSFORMER_CLASS) if TRANSFORMER_NAME else None
    if tr is not None:
        return tr

    trafos = app.GetCalcRelevantObjects(f"*.{TRANSFORMER_CLASS}")
    if not trafos:
        raise RuntimeError(f"Nie znaleziono transformatora klasy {TRANSFORMER_CLASS}.")
    return trafos[0]


# =============================================================================
# EXCEL
# =============================================================================

def _sheet_rows_from_workbook(wb: Any, sheet: str) -> List[Dict[str, Any]]:
    if sheet not in wb.sheetnames:
        return []

    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    out: List[Dict[str, Any]] = []

    for raw in rows[1:]:
        row = {headers[i]: raw[i] for i in range(min(len(headers), len(raw))) if headers[i]}
        if any(v not in (None, "") for v in row.values()):
            out.append(row)

    return out


def load_excel_cache(path: str) -> Dict[str, List[Dict[str, Any]]]:
    if load_workbook is None:
        raise RuntimeError("Do czytania Excela potrzebny jest pakiet openpyxl.")
    if not os.path.exists(path):
        raise FileNotFoundError(f"Nie znaleziono pliku Excel: {path}")

    wb = None
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        cache: Dict[str, List[Dict[str, Any]]] = {}
        for sheet in wb.sheetnames:
            cache[sheet] = _sheet_rows_from_workbook(wb, sheet)
        return cache
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


def excel_sheet(sheet: str) -> List[Dict[str, Any]]:
    return deepcopy(EXCEL_CACHE.get(sheet, []))


# =============================================================================
# USTAWIANIE MODELU
# =============================================================================

def set_loads_from_excel(app: Any) -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []

    for row in excel_sheet("Loads"):
        name_raw = row.get("name", "")
        if name_raw in (None, ""):
            continue
        name = str(name_raw).strip()
        if not name or name.lower() in {"none", "atribute", "attribute"}:
            continue

        elm = find_element(app, name, "ElmLod")
        if elm is None:
            log_line(f"[WARN] Loads: nie znaleziono ElmLod '{name}'")
            continue

        p1 = float(row.get("P1") or 0.0)
        p2 = float(row.get("P2") or 0.0)
        p3 = float(row.get("P3") or 0.0)
        q1 = float(row.get("Q1") or 0.0)
        q2 = float(row.get("Q2") or 0.0)
        q3 = float(row.get("Q3") or 0.0)

        set_attr(elm, ["plinir"], p1)
        set_attr(elm, ["plinis"], p2)
        set_attr(elm, ["plinit"], p3)
        set_attr(elm, ["qlinir"], q1)
        set_attr(elm, ["qlinis"], q2)
        set_attr(elm, ["qlinit"], q3)

        rows_out.append(
            {
                "name": name,
                "type": "load",
                "P1_kW": p1,
                "P2_kW": p2,
                "P3_kW": p3,
                "Q1_kvar": q1,
                "Q2_kvar": q2,
                "Q3_kvar": q3,
            }
        )

    return rows_out


def set_pv_from_excel(app: Any) -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []

    for row in excel_sheet("PV"):
        name_raw = row.get("name", "")
        if name_raw in (None, ""):
            continue
        name = str(name_raw).strip()
        if not name or name.lower() in {"none", "atribute", "attribute"}:
            continue

        elm = find_element(app, name, "ElmPvsys")
        if elm is None:
            log_line(f"[WARN] PV: nie znaleziono ElmPvsys '{name}'")
            continue

        p = float(row.get("P") or 0.0)
        smax = float(row.get("Smax") or row.get("Smax_kVA") or row.get("Sn") or 0.0)

        set_attr(elm, ["pgini"], p)
        set_attr(elm, ["qgini", "qsetp"], 0.0)
        set_attr(elm, ["av_mode"], "constq")

        rows_out.append(
            {
                "name": name,
                "type": "pv",
                "P_kW": p,
                "Q_kvar": 0.0,
                "Smax_kVA": smax,
                "control_mode": "constq",
            }
        )

    return rows_out


def set_other_generators_from_excel(app: Any) -> List[Dict[str, Any]]:
    rows_out: List[Dict[str, Any]] = []

    for sheet, cls in [("Generators", "ElmSym"), ("StatGen", "ElmGenstat")]:
        for row in excel_sheet(sheet):
            name_raw = row.get("name", "")
            if name_raw in (None, ""):
                continue
            name = str(name_raw).strip()
            if not name or name.lower() in {"none", "atribute", "attribute"}:
                continue

            elm = find_element(app, name, cls)
            if elm is None:
                log_line(f"[WARN] {sheet}: nie znaleziono {cls} '{name}'")
                continue

            p = float(row.get("P") or 0.0)
            q = float(row.get("Q") or 0.0)

            set_attr(elm, ["pgini"], p)
            set_attr(elm, ["qgini", "qsetp"], q)

            rows_out.append(
                {
                    "name": name,
                    "type": sheet,
                    "P_kW": p,
                    "Q_kvar": q,
                }
            )

    return rows_out


def load_storage_candidates() -> List[Dict[str, Any]]:
    rows = excel_sheet("StorageCandidates")
    return [
        {
            "node": str(r.get("node", "")).strip(),
            "role": str(r.get("role", "")).strip().lower(),
            "elem_L1": str(r.get("elem_L1", r.get("elem_A", ""))).strip(),
            "elem_L2": str(r.get("elem_L2", r.get("elem_B", ""))).strip(),
            "elem_L3": str(r.get("elem_L3", r.get("elem_C", ""))).strip(),
            "Pmax_kW": min(float(r.get("Pmax_kW") or r.get("Pmax") or 0.0), P_STORAGE_PHASE_MAX_KW),
            "Smax_kVA": float(r.get("Smax_kVA") or r.get("Smax") or 0.0),
        }
        for r in rows
        if str(r.get("node", "")).strip()
    ]


def get_critical_storage_candidate(candidates: List[Dict[str, Any]]) -> Dict[str, Any]:
    for cand in candidates:
        if cand.get("role") == "critical":
            return cand
    raise RuntimeError("Nie znaleziono magazynu oznaczonego jako role=critical w arkuszu StorageCandidates.")


def zero_storage_rows() -> List[Dict[str, Any]]:
    return [
        {"phase": "L1", "p_storage_kw": 0.0, "q_storage_kvar": 0.0},
        {"phase": "L2", "p_storage_kw": 0.0, "q_storage_kvar": 0.0},
        {"phase": "L3", "p_storage_kw": 0.0, "q_storage_kvar": 0.0},
    ]


def apply_storage_to_candidate(app: Any, candidate: Dict[str, Any], rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    applied_rows: List[Dict[str, Any]] = []

    for row in rows:
        ph = row["phase"]
        name = candidate.get(f"elem_{ph}", "")
        p = float(row.get("p_storage_kw") or 0.0)
        q = float(row.get("q_storage_kvar") or 0.0)

        if (p > 0.0) and not ALLOW_STORAGE_CHARGE:
            p = 0.0
        if (p < 0.0) and not ALLOW_STORAGE_DISCHARGE:
            p = 0.0
        p = max(-P_STORAGE_PHASE_MAX_KW, min(P_STORAGE_PHASE_MAX_KW, p))

        applied_row = {
            "node": candidate.get("node", ""),
            "element": name,
            "phase": ph,
            "p_storage_kw": p,
            "q_storage_kvar": q if KEEP_INITIAL_STORAGE_Q else 0.0,
            "Pmax_kW": float(candidate.get("Pmax_kW") or 0.0),
            "Smax_kVA": float(candidate.get("Smax_kVA") or 0.0),
        }

        if not name:
            log_line(f"[WARN] Storage: brak przypisania elementu dla fazy {ph}")
            applied_rows.append(applied_row)
            continue

        elm_lod = find_element(app, name, "ElmLod")
        if elm_lod is not None:
            set_attr(elm_lod, [PHASE_ATTR_P_LOAD[ph], "plini"], p)
            set_attr(elm_lod, [PHASE_ATTR_Q_LOAD[ph], "qlini"], applied_row["q_storage_kvar"])
            applied_rows.append(applied_row)
            continue

        elm_gen = (
            find_element(app, name, "ElmGenstat")
            or find_element(app, name, "ElmSym")
            or find_element(app, name, "ElmPvsys")
        )
        if elm_gen is not None:
            set_attr(elm_gen, ["pgini"], -p)
            set_attr(elm_gen, ["qgini", "qsetp"], -applied_row["q_storage_kvar"])
            applied_row["p_storage_kw_in_model"] = -p
            applied_row["q_storage_kvar_in_model"] = -applied_row["q_storage_kvar"]
            applied_rows.append(applied_row)
            continue

        log_line(f"[WARN] Storage: nie znaleziono elementu '{name}' dla fazy {ph}")
        applied_rows.append(applied_row)

    total_p = sum(float(r.get("p_storage_kw") or 0.0) for r in applied_rows)
    if abs(total_p) > P_STORAGE_TOTAL_MAX_KW + EPS:
        scale = P_STORAGE_TOTAL_MAX_KW / max(abs(total_p), EPS)
        for r in applied_rows:
            r["p_storage_kw"] = float(r["p_storage_kw"]) * scale

    return applied_rows


def set_initial_model(app: Any) -> Dict[str, List[Dict[str, Any]]]:
    log_line("Ustawianie modelu z Excela...")

    loads_data = set_loads_from_excel(app)
    pv_data = set_pv_from_excel(app)
    other_gen_data = set_other_generators_from_excel(app)

    candidates = load_storage_candidates()
    critical_candidate = get_critical_storage_candidate(candidates)
    storage_data = apply_storage_to_candidate(app, critical_candidate, zero_storage_rows())

    return {
        "loads_input": loads_data,
        "pv_input": pv_data,
        "other_generators_input": other_gen_data,
        "storage_input": storage_data,
        "critical_storage_meta": [critical_candidate],
    }


# =============================================================================
# PV FOR PSO
# =============================================================================

def compute_pv_q_limit(p_kw: float, smax_kva: float, tg_phi_max: float = TG_PHI_MAX) -> float:
    q_from_s = math.sqrt(max(0.0, smax_kva * smax_kva - p_kw * p_kw)) if smax_kva > EPS else 0.0
    q_from_pf = abs(p_kw) * tg_phi_max
    if smax_kva > EPS:
        return min(q_from_s, q_from_pf)
    return q_from_pf


def collect_pv_specs_for_pso(app: Any) -> List[Dict[str, Any]]:
    specs: List[Dict[str, Any]] = []

    for row in excel_sheet("PV"):
        name_raw = row.get("name", "")
        if name_raw in (None, ""):
            continue
        name = str(name_raw).strip()
        if not name or name.lower() in {"none", "atribute", "attribute"}:
            continue

        elm = find_element(app, name, "ElmPvsys")
        if elm is None:
            continue

        p_kw = float(row.get("P") or 0.0)
        smax_kva = float(row.get("Smax") or row.get("Smax_kVA") or row.get("Sn") or 0.0)
        qmax = compute_pv_q_limit(p_kw, smax_kva)

        specs.append(
            {
                "name": name,
                "element": elm,
                "P_kW": p_kw,
                "Smax_kVA": smax_kva,
                "Qmax_kvar": qmax,
                "Qmin_kvar": -qmax,
            }
        )

    return specs


def apply_pv_q_setpoints(pv_specs: List[Dict[str, Any]], q_values: Sequence[float]) -> List[Dict[str, Any]]:
    applied: List[Dict[str, Any]] = []

    for spec, q in zip(pv_specs, q_values):
        qmax = float(spec["Qmax_kvar"])
        q_limited = max(-qmax, min(qmax, float(q)))

        elm = spec["element"]
        set_attr(elm, ["qgini", "qsetp"], q_limited)
        set_attr(elm, ["av_mode"], "constq")

        applied.append(
            {
                "name": spec["name"],
                "P_kW": spec["P_kW"],
                "Q_kvar": q_limited,
                "Q_requested_kvar": float(q),
                "Qmax_kvar": qmax,
                "Smax_kVA": spec["Smax_kVA"],
                "av_mode": get_attr(elm, ["av_mode"], ""),
            }
        )

    return applied


# =============================================================================
# ODCZYT WYNIKÓW
# =============================================================================

def collect_node_voltages(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []

    for bus in app.GetCalcRelevantObjects("*.ElmTerm"):
        row = {"node": getattr(bus, "loc_name", "")}

        for ph in PHASES:
            pf_ph = PF_PHASE[ph]
            row[f"U_{ph}_pu"] = get_float(bus, [f"m:u:{pf_ph}"], math.nan)
            row[f"angle_{ph}_deg"] = get_float(bus, [f"m:phiu:{pf_ph}"], math.nan)

        row["kU2_percent"] = get_float(bus, ["m:ubfac"], math.nan)
        rows.append(row)

    return rows


def collect_transformer_results(tr: Any) -> Dict[str, Any]:
    row = {"transformer": getattr(tr, "loc_name", "")}

    i_complex = []

    for ph in PHASES:
        pf_ph = PF_PHASE[ph]

        i_ka = get_float(tr, [f"m:I:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:I:{pf_ph}"], 0.0)
        phi_i = get_float(tr, [f"m:phii:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:phii:{pf_ph}"], math.nan)

        p = get_attr(
            tr,
            [f"m:P:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:P:{pf_ph}", f"c:P:{pf_ph}"],
            None,
        )
        q = get_attr(
            tr,
            [f"m:Q:{TRANSFORMER_LV_SIDE}:{pf_ph}", f"m:Q:{pf_ph}", f"c:Q:{pf_ph}"],
            None,
        )

        i_a = float(i_ka or 0.0) * 1000.0

        row[f"I_tr_{ph}_A"] = i_a
        row[f"angle_I_tr_{ph}_deg"] = phi_i
        row[f"P_tr_{ph}_kW"] = None if p is None else float(p)
        row[f"Q_tr_{ph}_kvar"] = None if q is None else float(q)

        i_complex.append(angle_deg_to_complex(i_a, float(phi_i) if math.isfinite(float(phi_i)) else 0.0))

    row["loading_percent"] = get_float(tr, ["c:loading"], math.nan)
    row["I_neutral_A"] = abs(sum(i_complex))
    return row


def collect_branch_losses(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []

    for line in app.GetCalcRelevantObjects("*.ElmLne"):
        p_loss = get_attr(line, ["c:LossP", "c:Ploss", "c:Losses"], 0.0)
        q_loss = get_attr(line, ["c:LossQ", "c:Qloss"], 0.0)

        rows.append(
            {
                "branch": getattr(line, "loc_name", ""),
                "P_loss_kW": float(p_loss or 0.0),
                "Q_loss_kvar": float(q_loss or 0.0),
                "loading_percent": get_float(line, ["c:loading"], math.nan),
            }
        )

    return rows


def collect_pv_setpoints(app: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []

    for pv in app.GetCalcRelevantObjects("*.ElmPvsys"):
        p_set = get_float(pv, ["pgini"], 0.0)
        q_set = get_float(pv, ["qgini", "qsetp"], 0.0)
        q_res = get_float(pv, ["m:Q:bus1", "m:Q:buslv", "m:Q", "c:Q"], q_set)
        smax = get_float(pv, ["sgn", "sn", "snom"], 0.0)

        rows.append(
            {
                "name": getattr(pv, "loc_name", ""),
                "P_kW": p_set,
                "Q_kvar": q_res,
                "Q_set_kvar": q_set,
                "S_apparent_kVA": math.sqrt(p_set * p_set + q_res * q_res),
                "Smax_kVA": smax,
                "av_mode": get_attr(pv, ["av_mode"], ""),
            }
        )

    return rows


# =============================================================================
# POMOCNICZE OBLICZENIA
# =============================================================================

def angle_deg_to_complex(magnitude: float, angle_deg: float) -> complex:
    if not math.isfinite(magnitude):
        return complex(math.nan, math.nan)
    if not math.isfinite(angle_deg):
        return complex(magnitude, 0.0)
    return cmath.rect(magnitude, math.radians(angle_deg))


def calc_symmetrical_components(v_a: complex, v_b: complex, v_c: complex) -> Tuple[complex, complex, complex]:
    a = complex(-0.5, math.sqrt(3.0) / 2.0)
    a2 = a * a

    u0 = (v_a + v_b + v_c) / 3.0
    u1 = (v_a + a * v_b + a2 * v_c) / 3.0
    u2 = (v_a + a2 * v_b + a * v_c) / 3.0
    return u0, u1, u2


def rms_deviation(values: List[float], target: float) -> float:
    if not values:
        return math.nan
    return math.sqrt(sum((v - target) ** 2 for v in values) / len(values))


def export_import_from_transformer_phases(p_tr: List[float]) -> Tuple[float, float]:
    export_pos = sum(max(0.0, p) for p in p_tr)
    export_neg = sum(max(0.0, -p) for p in p_tr)

    if TRANSFORMER_EXPORT_POSITIVE:
        return export_pos, export_neg
    return export_neg, export_pos


def calc_all_voltage_values(raw: Dict[str, List[Dict[str, Any]]]) -> List[float]:
    vals: List[float] = []
    for row in raw["node_voltages"]:
        for ph in PHASES:
            try:
                u = float(row[f"U_{ph}_pu"])
                if math.isfinite(u):
                    vals.append(u)
            except Exception:
                pass
    return vals


def calc_penalty_u(raw: Dict[str, List[Dict[str, Any]]]) -> float:
    vals = calc_all_voltage_values(raw)
    if not vals:
        return 0.0

    total = 0.0
    for u in vals:
        total += (
            (max(0.0, u - U_MAX_ALLOWED_OBJ) / max(U_TOL_PU, EPS)) ** 2
            + (max(0.0, U_MIN_ALLOWED_OBJ - u) / max(U_TOL_PU, EPS)) ** 2
        )
    return total


def calc_penalty_lines(raw: Dict[str, List[Dict[str, Any]]]) -> float:
    total = 0.0
    for row in raw["branch_losses"]:
        try:
            loading = float(row.get("loading_percent"))
            if math.isfinite(loading):
                total += max(0.0, loading / 100.0 - 1.0) ** 2
        except Exception:
            pass
    return total


def calc_penalty_trafo(raw: Dict[str, List[Dict[str, Any]]]) -> float:
    tr = raw["transformer_phase_results"][0] if raw["transformer_phase_results"] else {}
    try:
        loading = float(tr.get("loading_percent"))
        if math.isfinite(loading):
            return max(0.0, loading / 100.0 - 1.0) ** 2
    except Exception:
        pass
    return 0.0


def calc_ju(raw: Dict[str, List[Dict[str, Any]]]) -> float:
    vals = calc_all_voltage_values(raw)
    if not vals:
        return math.nan
    return math.sqrt(sum((u - TARGET_VOLTAGE_PU) ** 2 for u in vals) / len(vals))


def calc_ja(alpha2_max: float, ku2_max_percent: float) -> float:
    if math.isfinite(alpha2_max):
        return alpha2_max
    if math.isfinite(ku2_max_percent):
        return ku2_max_percent / 100.0
    return math.nan


def calc_ji(i_unbalance_tr_percent: float) -> float:
    if not math.isfinite(i_unbalance_tr_percent):
        return math.nan
    return i_unbalance_tr_percent / 100.0


def calc_jp(p_export_total_kw: float) -> float:
    ref = max(P_EXPORT_REF_KW, EPS)
    return p_export_total_kw / ref


def calc_f1(ju: float, ja: float, penalty_u: float, penalty_lines: float, penalty_trafo: float) -> float:
    ju_used = 0.0 if not math.isfinite(ju) else ju
    ja_used = 0.0 if not math.isfinite(ja) else ja
    return (
        0.70 * ju_used
        + 0.30 * ja_used
        + K_U * penalty_u
        + K_L * penalty_lines
        + K_T * penalty_trafo
    )


def calc_f2(ju: float, ji: float, jp: float, penalty_u: float, penalty_lines: float, penalty_trafo: float) -> float:
    ju_used = 0.0 if not math.isfinite(ju) else ju
    ji_used = 0.0 if not math.isfinite(ji) else ji
    jp_used = 0.0 if not math.isfinite(jp) else jp
    return (
        0.20 * ju_used
        + 0.35 * ji_used
        + 0.45 * jp_used
        + K_U * penalty_u
        + K_L * penalty_lines
        + K_T * penalty_trafo
    )


# =============================================================================
# WSKAŹNIKI RAPORTOWE
# =============================================================================

def calculate_indicators(raw: Dict[str, List[Dict[str, Any]]]) -> Dict[str, Any]:
    voltages: List[float] = []
    phase_spreads: List[float] = []
    ku2_vals: List[float] = []

    alpha0_vals: List[float] = []
    alpha2_vals: List[float] = []
    u0_abs_vals: List[float] = []
    u1_abs_vals: List[float] = []
    u2_abs_vals: List[float] = []

    node_sequence_rows: List[Dict[str, Any]] = []

    for row in raw["node_voltages"]:
        vals = []
        for ph in PHASES:
            try:
                v = float(row[f"U_{ph}_pu"])
                if math.isfinite(v):
                    vals.append(v)
                    voltages.append(v)
            except Exception:
                pass

        if vals:
            phase_spreads.append(max(vals) - min(vals))

        try:
            k = float(row.get("kU2_percent"))
            if math.isfinite(k):
                ku2_vals.append(k)
        except Exception:
            pass

        try:
            va = angle_deg_to_complex(float(row["U_L1_pu"]), float(row["angle_L1_deg"]))
            vb = angle_deg_to_complex(float(row["U_L2_pu"]), float(row["angle_L2_deg"]))
            vc = angle_deg_to_complex(float(row["U_L3_pu"]), float(row["angle_L3_deg"]))

            u0, u1, u2 = calc_symmetrical_components(va, vb, vc)

            u0_abs = abs(u0)
            u1_abs = abs(u1)
            u2_abs = abs(u2)

            alpha0 = u0_abs / u1_abs if u1_abs > EPS else math.nan
            alpha2 = u2_abs / u1_abs if u1_abs > EPS else math.nan

            if math.isfinite(alpha0):
                alpha0_vals.append(alpha0)
            if math.isfinite(alpha2):
                alpha2_vals.append(alpha2)

            if math.isfinite(u0_abs):
                u0_abs_vals.append(u0_abs)
            if math.isfinite(u1_abs):
                u1_abs_vals.append(u1_abs)
            if math.isfinite(u2_abs):
                u2_abs_vals.append(u2_abs)

            node_sequence_rows.append(
                {
                    "node": row.get("node", ""),
                    "U0_abs_pu": u0_abs,
                    "U1_abs_pu": u1_abs,
                    "U2_abs_pu": u2_abs,
                    "alpha0": alpha0,
                    "alpha2": alpha2,
                }
            )
        except Exception:
            node_sequence_rows.append(
                {
                    "node": row.get("node", ""),
                    "U0_abs_pu": math.nan,
                    "U1_abs_pu": math.nan,
                    "U2_abs_pu": math.nan,
                    "alpha0": math.nan,
                    "alpha2": math.nan,
                }
            )

    tr = raw["transformer_phase_results"][0] if raw["transformer_phase_results"] else {}

    i_vals = [float(tr.get(f"I_tr_{ph}_A") or 0.0) for ph in PHASES]
    i_avg = sum(i_vals) / 3.0 if i_vals else 0.0
    i_unb = max(abs(i - i_avg) for i in i_vals) / i_avg * 100.0 if i_avg > EPS else 0.0
    i_neutral = float(tr.get("I_neutral_A") or 0.0)

    p_tr = [float(tr.get(f"P_tr_{ph}_kW") or 0.0) for ph in PHASES]
    q_tr = [float(tr.get(f"Q_tr_{ph}_kvar") or 0.0) for ph in PHASES]
    p_export_total, p_import_total = export_import_from_transformer_phases(p_tr)

    p_loss = sum(float(row.get("P_loss_kW") or 0.0) for row in raw["branch_losses"])
    q_loss = sum(float(row.get("Q_loss_kvar") or 0.0) for row in raw["branch_losses"])

    pv_rows = raw.get("pv_setpoints", [])
    storage_rows = raw.get("storage_setpoints", [])

    p_pv_total = sum(float(row.get("P_kW") or 0.0) for row in pv_rows)
    q_pv_total = sum(float(row.get("Q_kvar") or 0.0) for row in pv_rows)
    q_pv_abs_sum = sum(abs(float(row.get("Q_kvar") or 0.0)) for row in pv_rows)

    p_storage_L1 = float(storage_rows[0].get("p_storage_kw") or 0.0) if len(storage_rows) > 0 else 0.0
    p_storage_L2 = float(storage_rows[1].get("p_storage_kw") or 0.0) if len(storage_rows) > 1 else 0.0
    p_storage_L3 = float(storage_rows[2].get("p_storage_kw") or 0.0) if len(storage_rows) > 2 else 0.0
    p_storage_total = p_storage_L1 + p_storage_L2 + p_storage_L3
    q_storage_total = sum(float(row.get("q_storage_kvar") or 0.0) for row in storage_rows)

    udev_mean_1_00 = sum(abs(v - 1.0) for v in voltages) / len(voltages) if voltages else math.nan
    udev_rms_1_00 = rms_deviation(voltages, 1.00)
    udev_rms_1_05 = rms_deviation(voltages, 1.05)

    alpha0_mean = sum(alpha0_vals) / len(alpha0_vals) if alpha0_vals else math.nan
    alpha0_max = max(alpha0_vals) if alpha0_vals else math.nan
    alpha2_mean = sum(alpha2_vals) / len(alpha2_vals) if alpha2_vals else math.nan
    alpha2_max = max(alpha2_vals) if alpha2_vals else math.nan

    fcelu_3_voltage_component = 0.04 * (udev_rms_1_05 if math.isfinite(udev_rms_1_05) else 0.0)
    fcelu_3_alpha2_component = 0.58 * (alpha2_mean if math.isfinite(alpha2_mean) else 0.0)
    fcelu_3_alpha0_component = 0.38 * (alpha0_mean if math.isfinite(alpha0_mean) else 0.0)
    fcelu_3 = fcelu_3_voltage_component + fcelu_3_alpha2_component + fcelu_3_alpha0_component

    penalty_u = calc_penalty_u(raw)
    penalty_lines = calc_penalty_lines(raw)
    penalty_trafo = calc_penalty_trafo(raw)

    ju = calc_ju(raw)
    ja = calc_ja(alpha2_max, max(ku2_vals) if ku2_vals else math.nan)
    ji = calc_ji(i_unb)
    jp = calc_jp(p_export_total)

    f1 = calc_f1(ju, ja, penalty_u, penalty_lines, penalty_trafo)
    f2 = calc_f2(ju, ji, jp, penalty_u, penalty_lines, penalty_trafo)

    violations: List[str] = []

    for row in raw["node_voltages"]:
        for ph in PHASES:
            try:
                u = float(row[f"U_{ph}_pu"])
                if u < VOLTAGE_MIN_PU or u > VOLTAGE_MAX_PU:
                    violations.append(f"U:{row.get('node')}:{ph}:{u:.4f}")
            except Exception:
                pass

    try:
        tr_loading = float(tr.get("loading_percent"))
        if math.isfinite(tr_loading) and tr_loading > LOADING_MAX_PERCENT:
            violations.append(f"TR_loading:{tr_loading:.2f}")
    except Exception:
        pass

    for row in raw["branch_losses"]:
        try:
            line_loading = float(row.get("loading_percent"))
            if math.isfinite(line_loading) and line_loading > LINE_LOADING_MAX_PERCENT:
                violations.append(f"LINE_loading:{row.get('branch')}:{line_loading:.2f}")
        except Exception:
            pass

    indicators = {
        "Umax_pu": max(voltages) if voltages else math.nan,
        "Umin_pu": min(voltages) if voltages else math.nan,
        "Udev_mean_pu": udev_mean_1_00,
        "Udev_max_pu": max(abs(v - 1.0) for v in voltages) if voltages else math.nan,
        "Fcelu_1_Udev_rms_1_00": udev_rms_1_00,
        "Fcelu_1_Udev_rms_1_05": udev_rms_1_05,
        "dU_phase_max_pu": max(phase_spreads) if phase_spreads else math.nan,
        "kU2_max_percent": max(ku2_vals) if ku2_vals else math.nan,
        "alpha0_mean": alpha0_mean,
        "alpha0_max": alpha0_max,
        "alpha2_mean": alpha2_mean,
        "alpha2_max": alpha2_max,
        "Fcelu_3_weighted": fcelu_3,

        "I_tr_L1_A": i_vals[0],
        "I_tr_L2_A": i_vals[1],
        "I_tr_L3_A": i_vals[2],
        "I_unbalance_tr_percent": i_unb,
        "I_neutral_A": i_neutral,
        "Fcelu_2_A": i_neutral,

        "P_tr_L1_kW": p_tr[0],
        "P_tr_L2_kW": p_tr[1],
        "P_tr_L3_kW": p_tr[2],
        "Q_tr_L1_kvar": q_tr[0],
        "Q_tr_L2_kvar": q_tr[1],
        "Q_tr_L3_kvar": q_tr[2],
        "P_export_total_kW": p_export_total,
        "P_import_total_kW": p_import_total,

        "P_loss_total_kW": p_loss,
        "Q_loss_total_kvar": q_loss,

        "P_pv_total_kW": p_pv_total,
        "Q_pv_total_kvar": q_pv_total,
        "Q_pv_abs_sum_kvar": q_pv_abs_sum,

        "P_storage_L1_kW": p_storage_L1,
        "P_storage_L2_kW": p_storage_L2,
        "P_storage_L3_kW": p_storage_L3,
        "P_storage_total_kW": p_storage_total,
        "Q_storage_total_kvar": q_storage_total,

        "JU": ju,
        "JA": ja,
        "JI": ji,
        "JP": jp,
        "penalty_U": penalty_u,
        "penalty_lines": penalty_lines,
        "penalty_trafo": penalty_trafo,
        "F1": f1,
        "F2": f2,

        "constraint_violations": ";".join(violations),
    }

    raw["node_sequence_components"] = node_sequence_rows
    return indicators


# =============================================================================
# FUNKCJE CELU
# =============================================================================

def calculate_used_objective_critical(indicators: Dict[str, Any]) -> Dict[str, float]:
    ju = float(indicators.get("JU") or 0.0)
    ja = float(indicators.get("JA") or 0.0)
    penalty_u = float(indicators.get("penalty_U") or 0.0)
    penalty_lines = float(indicators.get("penalty_lines") or 0.0)
    penalty_trafo = float(indicators.get("penalty_trafo") or 0.0)

    f1 = calc_f1(ju, ja, penalty_u, penalty_lines, penalty_trafo)
    ji = float(indicators.get("JI") or 0.0)
    jp = float(indicators.get("JP") or 0.0)
    f2 = calc_f2(ju, ji, jp, penalty_u, penalty_lines, penalty_trafo)

    return {
        "J_used_total": f1,
        "J_used_component_1": ju,
        "J_used_component_2": ja,
        "J_used_component_3": penalty_u,
        "J_used_component_4": penalty_lines,
        "J_used_component_5": penalty_trafo,
        "F1": f1,
        "F2": f2,
    }


# =============================================================================
# PSO HELPERS
# =============================================================================

def build_decision_vector_definition(
    pv_specs: List[Dict[str, Any]],
    critical_candidate: Dict[str, Any],
) -> Tuple[List[Dict[str, Any]], List[float], List[float]]:
    defs: List[Dict[str, Any]] = []
    lb: List[float] = []
    ub: List[float] = []

    for spec in pv_specs:
        qmin = float(spec["Qmin_kvar"])
        qmax = float(spec["Qmax_kvar"])
        defs.append(
            {
                "var_type": "pv_q",
                "name": spec["name"],
                "phase": "",
                "lower": qmin,
                "upper": qmax,
            }
        )
        lb.append(qmin)
        ub.append(qmax)

    pmax = min(float(critical_candidate.get("Pmax_kW") or 0.0), P_STORAGE_PHASE_MAX_KW)
    for ph in PHASES:
        low = -pmax if ALLOW_STORAGE_DISCHARGE else 0.0
        high = pmax if ALLOW_STORAGE_CHARGE else 0.0

        defs.append(
            {
                "var_type": "storage_p",
                "name": critical_candidate.get(f"elem_{ph}", ""),
                "phase": ph,
                "lower": low,
                "upper": high,
            }
        )
        lb.append(low)
        ub.append(high)

    return defs, lb, ub


def unpack_particle(
    x: Sequence[float],
    pv_specs: List[Dict[str, Any]],
) -> Tuple[List[float], List[float]]:
    n_pv = len(pv_specs)
    q_pv = [float(v) for v in x[:n_pv]]
    p_storage = [float(v) for v in x[n_pv:n_pv + 3]]
    return q_pv, p_storage


def storage_rows_from_p_values(p_values: Sequence[float]) -> List[Dict[str, Any]]:
    rows = []
    for ph, p in zip(PHASES, p_values):
        rows.append(
            {
                "phase": ph,
                "p_storage_kw": max(-P_STORAGE_PHASE_MAX_KW, min(P_STORAGE_PHASE_MAX_KW, float(p))),
                "q_storage_kvar": 0.0,
            }
        )
    return rows


def build_best_solution_table(
    x_best: Sequence[float],
    var_defs: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    rows = []
    for val, vd in zip(x_best, var_defs):
        rows.append(
            {
                "variable_name": f"{vd['var_type']}_{vd['name']}_{vd['phase']}".strip("_"),
                "variable_type": vd["var_type"],
                "element_name": vd["name"],
                "phase": vd["phase"],
                "final_value": float(val),
                "lower_bound": float(vd["lower"]),
                "upper_bound": float(vd["upper"]),
            }
        )
    return rows


def build_pso_config_table() -> List[Dict[str, Any]]:
    return [
        {"parameter": "N_PARTICLES", "value": N_PARTICLES},
        {"parameter": "N_ITER", "value": N_ITER},
        {"parameter": "PSO_W", "value": PSO_W},
        {"parameter": "PSO_C1", "value": PSO_C1},
        {"parameter": "PSO_C2", "value": PSO_C2},
        {"parameter": "PSO_RANDOM_SEED", "value": PSO_RANDOM_SEED},
        {"parameter": "NONCONVERGENCE_PENALTY", "value": NONCONVERGENCE_PENALTY},
        {"parameter": "UREF_PU", "value": TARGET_VOLTAGE_PU},
        {"parameter": "U_MIN_ALLOWED_PU", "value": U_MIN_ALLOWED_OBJ},
        {"parameter": "U_MAX_ALLOWED_PU", "value": U_MAX_ALLOWED_OBJ},
        {"parameter": "U_TOL_PU", "value": U_TOL_PU},
        {"parameter": "P_EXPORT_REF_KW", "value": P_EXPORT_REF_KW},
        {"parameter": "K_U", "value": K_U},
        {"parameter": "K_L", "value": K_L},
        {"parameter": "K_T", "value": K_T},
        {"parameter": "TG_PHI_MAX", "value": TG_PHI_MAX},
        {"parameter": "objective_profile", "value": "critical_voltage"},
    ]


def summarize_iteration_evals(iteration: int, evals: List[Dict[str, Any]]) -> Dict[str, Any]:
    if not evals:
        return {
            "iteration": iteration,
            "best_J": math.nan,
            "mean_J": math.nan,
            "worst_J": math.nan,
            "std_J": math.nan,
            "n_feasible": 0,
            "n_nonconverged": 0,
            "best_JU": math.nan,
            "best_JA": math.nan,
            "best_penalty_U": math.nan,
            "best_penalty_lines": math.nan,
            "best_penalty_trafo": math.nan,
            "mean_JU": math.nan,
            "mean_JA": math.nan,
        }

    feasible = [e for e in evals if e.get("converged")]
    all_j = [float(e["J_used_total"]) for e in evals if "J_used_total" in e]

    if feasible:
        best_eval = min(feasible, key=lambda e: float(e["J_used_total"]))
        mean_ju = sum(float(e["JU"]) for e in feasible) / len(feasible)
        mean_ja = sum(float(e["JA"]) for e in feasible) / len(feasible)
    else:
        best_eval = {}
        mean_ju = mean_ja = math.nan

    mean_j = sum(all_j) / len(all_j) if all_j else math.nan
    worst_j = max(all_j) if all_j else math.nan
    std_j = math.sqrt(sum((j - mean_j) ** 2 for j in all_j) / len(all_j)) if all_j else math.nan

    return {
        "iteration": iteration,
        "best_J": min(all_j) if all_j else math.nan,
        "mean_J": mean_j,
        "worst_J": worst_j,
        "std_J": std_j,
        "n_feasible": len(feasible),
        "n_nonconverged": len([e for e in evals if not e.get("converged")]),
        "best_JU": float(best_eval.get("JU", math.nan)),
        "best_JA": float(best_eval.get("JA", math.nan)),
        "best_penalty_U": float(best_eval.get("penalty_U", math.nan)),
        "best_penalty_lines": float(best_eval.get("penalty_lines", math.nan)),
        "best_penalty_trafo": float(best_eval.get("penalty_trafo", math.nan)),
        "mean_JU": mean_ju,
        "mean_JA": mean_ja,
    }


# =============================================================================
# GLOBAL STATE FOR OBJECTIVE
# =============================================================================

GLOBAL_CTX: Dict[str, Any] = {}


def evaluate_particle(x: Sequence[float]) -> float:
    global PARTICLE_EVAL_COUNTER, CURRENT_ITER_EVALS, PSO_PARTICLE_EVAL_ROWS

    PARTICLE_EVAL_COUNTER += 1

    app = GLOBAL_CTX["app"]
    ldf = GLOBAL_CTX["ldf"]
    tr = GLOBAL_CTX["tr"]
    pv_specs = GLOBAL_CTX["pv_specs"]
    critical_candidate = GLOBAL_CTX["critical_candidate"]
    current_iteration = GLOBAL_CTX.get("current_iteration", -1)
    current_particle = GLOBAL_CTX.get("current_particle", -1)

    q_pv, p_storage = unpack_particle(x, pv_specs)

    apply_pv_q_setpoints(pv_specs, q_pv)
    storage_rows = apply_storage_to_candidate(app, critical_candidate, storage_rows_from_p_values(p_storage))

    converged, rc = try_run_loadflow(ldf)
    if not converged:
        eval_row = {
            "iteration": current_iteration,
            "particle": current_particle,
            "eval_no": PARTICLE_EVAL_COUNTER,
            "J_used_total": NONCONVERGENCE_PENALTY,
            "JU": math.nan,
            "JA": math.nan,
            "penalty_U": math.nan,
            "penalty_lines": math.nan,
            "penalty_trafo": math.nan,
            "converged": False,
            "loadflow_result": str(rc),
        }
        CURRENT_ITER_EVALS.append(eval_row)
        PSO_PARTICLE_EVAL_ROWS.append(eval_row)

        if PARTICLE_EVAL_COUNTER <= 20 or PARTICLE_EVAL_COUNTER % 50 == 0:
            log_line(f"[WARN] Niezbieżny loadflow dla particle eval #{PARTICLE_EVAL_COUNTER}, rc={rc}")
        return NONCONVERGENCE_PENALTY

    raw = {
        "node_voltages": collect_node_voltages(app),
        "transformer_phase_results": [collect_transformer_results(tr)],
        "branch_losses": collect_branch_losses(app),
        "pv_setpoints": collect_pv_setpoints(app),
        "storage_setpoints": storage_rows,
    }
    indicators = calculate_indicators(raw)
    obj = calculate_used_objective_critical(indicators)

    eval_row = {
        "iteration": current_iteration,
        "particle": current_particle,
        "eval_no": PARTICLE_EVAL_COUNTER,
        "J_used_total": obj["J_used_total"],
        "JU": float(indicators.get("JU", math.nan)),
        "JA": float(indicators.get("JA", math.nan)),
        "penalty_U": float(indicators.get("penalty_U", math.nan)),
        "penalty_lines": float(indicators.get("penalty_lines", math.nan)),
        "penalty_trafo": float(indicators.get("penalty_trafo", math.nan)),
        "converged": True,
        "loadflow_result": "0",
    }
    CURRENT_ITER_EVALS.append(eval_row)
    PSO_PARTICLE_EVAL_ROWS.append(eval_row)

    return obj["J_used_total"]


class TrackingPSO(PSO):
    def optimize(self):
        global CURRENT_ITER_EVALS, PSO_CONVERGENCE_ROWS

        CURRENT_ITER_EVALS = []
        GLOBAL_CTX["current_iteration"] = 0

        for p in range(self.n_particles):
            GLOBAL_CTX["current_particle"] = p
            val = self._eval_particle(self.X[p], p)
            self.pbest_val[p] = val
            if val < self.gbest_val:
                self.gbest_val = val
                self.gbest = self.X[p].copy()

        self.best_per_iter.append(self.gbest_val)
        PSO_CONVERGENCE_ROWS.append(summarize_iteration_evals(0, CURRENT_ITER_EVALS))

        for it in range(1, self.max_iter + 1):
            self.iter = it
            CURRENT_ITER_EVALS = []
            GLOBAL_CTX["current_iteration"] = it

            for p in range(self.n_particles):
                GLOBAL_CTX["current_particle"] = p

                r1 = np.random.rand(self.dim)
                r2 = np.random.rand(self.dim)

                self.V[p] = (
                    self.w * self.V[p]
                    + self.c1 * r1 * (self.pbest[p] - self.X[p])
                    + self.c2 * r2 * (self.gbest - self.X[p])
                )

                vmax = (self.ub - self.lb) * 0.5
                self.V[p] = np.clip(self.V[p], -vmax, vmax)
                self.X[p] = np.clip(self.X[p] + self.V[p], self.lb, self.ub)

                val = self._eval_particle(self.X[p], p)

                if val < self.pbest_val[p]:
                    self.pbest_val[p] = val
                    self.pbest[p] = self.X[p].copy()

                if val < self.gbest_val:
                    self.gbest_val = val
                    self.gbest = self.X[p].copy()

            self.best_per_iter.append(self.gbest_val)
            PSO_CONVERGENCE_ROWS.append(summarize_iteration_evals(it, CURRENT_ITER_EVALS))

            if self.autosave_every_iters and (it % self.autosave_every_iters == 0):
                self.save_checkpoint()

        if self.autosave_every_iters:
            self.save_checkpoint()

        return {
            "gbest": self.gbest,
            "gbest_val": self.gbest_val,
            "best_per_iter": self.best_per_iter,
        }


# =============================================================================
# WYNIKI
# =============================================================================

def collect_results_with_pso(
    app: Any,
    tr: Any,
    setup_data: Dict[str, List[Dict[str, Any]]],
    storage_rows: List[Dict[str, Any]],
    pso_config: List[Dict[str, Any]],
    pso_convergence: List[Dict[str, Any]],
    pso_particle_evals: List[Dict[str, Any]],
    best_solution: List[Dict[str, Any]],
) -> Dict[str, List[Dict[str, Any]]]:
    node_voltages = collect_node_voltages(app)
    transformer_results = [collect_transformer_results(tr)]
    branch_losses = collect_branch_losses(app)
    pv_setpoints = collect_pv_setpoints(app)

    raw = {
        "node_voltages": node_voltages,
        "transformer_phase_results": transformer_results,
        "branch_losses": branch_losses,
        "pv_setpoints": pv_setpoints,
        "storage_setpoints": storage_rows,
    }
    indicators = calculate_indicators(raw)
    obj = calculate_used_objective_critical(indicators)

    indicators.update(
        {
            "objective_profile": "critical_voltage",
            "J_used_total": obj["J_used_total"],
            "J_used_component_1": indicators.get("JU", math.nan),
            "J_used_component_2": indicators.get("JA", math.nan),
            "J_used_component_3": indicators.get("penalty_U", math.nan),
            "J_used_component_4": indicators.get("penalty_lines", math.nan),
            "J_used_component_5": indicators.get("penalty_trafo", math.nan),
        }
    )

    return {
        "control_config": [{"parameter": k, "value": v} for k, v in CONTROL_CONFIG.items()],
        "inputs_loads": setup_data["loads_input"],
        "inputs_pv": setup_data["pv_input"],
        "inputs_other_generators": setup_data["other_generators_input"],
        "inputs_storage_critical_meta": setup_data["critical_storage_meta"],
        "pso_config": pso_config,
        "pso_convergence": pso_convergence,
        "pso_particle_evals": pso_particle_evals,
        "best_solution": best_solution,
        "node_voltages": node_voltages,
        "node_sequence_components": raw.get("node_sequence_components", []),
        "transformer_phase_results": transformer_results,
        "pv_setpoints": pv_setpoints,
        "storage_setpoints": storage_rows,
        "branch_losses": branch_losses,
        "indicators": [indicators],
    }


# =============================================================================
# EXCEL EXPORT
# =============================================================================

def safe_sheet_name(name: str) -> str:
    bad = "[]:*?/\\"
    for ch in bad:
        name = name.replace(ch, "_")
    return name[:31]


def write_rows(ws: Any, rows: Iterable[Dict[str, Any]]) -> None:
    rows = list(rows)
    if not rows:
        ws.append(["empty"])
        return

    headers: List[str] = []
    for row in rows:
        for key in row.keys():
            if key not in headers:
                headers.append(key)

    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])


def export_to_excel(all_tables: Dict[str, List[Dict[str, Any]]], out_file: str) -> None:
    if Workbook is None:
        raise RuntimeError("Do zapisu wyników do Excela potrzebny jest pakiet openpyxl.")

    out_dir = os.path.dirname(out_file)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, rows in all_tables.items():
        ws = wb.create_sheet(safe_sheet_name(sheet_name))
        write_rows(ws, rows)

    wb.save(out_file)


# =============================================================================
# MAIN
# =============================================================================

def run_pso_global_critical() -> None:
    global EXCEL_CACHE, CONTROL_CONFIG, GLOBAL_CTX
    global PSO_CONVERGENCE_ROWS, PSO_PARTICLE_EVAL_ROWS, CURRENT_ITER_EVALS, PARTICLE_EVAL_COUNTER

    PARTICLE_EVAL_COUNTER = 0
    PSO_CONVERGENCE_ROWS = []
    PSO_PARTICLE_EVAL_ROWS = []
    CURRENT_ITER_EVALS = []

    try:
        log_dir = os.path.dirname(LOG_FILE)
        if log_dir:
            os.makedirs(log_dir, exist_ok=True)
        with open(LOG_FILE, "w", encoding="utf-8") as fh:
            fh.write("=== PSO GLOBAL CRITICAL LOG START ===\n")
    except Exception:
        pass

    log_line("Start pso_global_critical.py")

    EXCEL_CACHE = load_excel_cache(EXCEL_FILE)
    CONTROL_CONFIG = load_control_config()
    apply_runtime_config()

    random.seed(PSO_RANDOM_SEED)
    np.random.seed(PSO_RANDOM_SEED)

    log_line(f"Wczytano Excel: {EXCEL_FILE}")
    log_line(f"Arkusze: {', '.join(sorted(EXCEL_CACHE.keys()))}")
    log_line(
        f"Config: Vmin={VOLTAGE_MIN_PU}, Uref={TARGET_VOLTAGE_PU}, Vmax={VOLTAGE_MAX_PU}, "
        f"Uallowed=[{U_MIN_ALLOWED_OBJ}, {U_MAX_ALLOWED_OBJ}], "
        f"Utol={U_TOL_PU}, PexportRef={P_EXPORT_REF_KW}, "
        f"K=({K_U}, {K_L}, {K_T}), "
        f"PSO particles={N_PARTICLES}, iter={N_ITER}, seed={PSO_RANDOM_SEED}"
    )

    app, ldf = connect_powerfactory()
    tr = get_transformer(app)

    setup_data = set_initial_model(app)
    critical_candidate = setup_data["critical_storage_meta"][0]

    pv_specs = collect_pv_specs_for_pso(app)
    if not pv_specs:
        raise RuntimeError("Nie znaleziono żadnych PV do sterowania Q w PSO.")

    log_line(f"PV do sterowania Q: {len(pv_specs)}")
    log_line(f"Magazyn critical: node={critical_candidate.get('node')}")

    apply_pv_q_setpoints(pv_specs, [0.0] * len(pv_specs))
    base_storage_rows = apply_storage_to_candidate(app, critical_candidate, zero_storage_rows())

    converged_base, rc_base = try_run_loadflow(ldf)
    if not converged_base:
        raise RuntimeError(f"Bazowy loadflow nie jest zbieżny, rc={rc_base}")

    var_defs, lb, ub = build_decision_vector_definition(pv_specs, critical_candidate)

    GLOBAL_CTX = {
        "app": app,
        "ldf": ldf,
        "tr": tr,
        "pv_specs": pv_specs,
        "critical_candidate": critical_candidate,
        "current_iteration": -1,
        "current_particle": -1,
    }

    log_line(f"Uruchamiam PSO: particles={N_PARTICLES}, iter={N_ITER}, dim={len(var_defs)}")

    pso = TrackingPSO(
        func=evaluate_particle,
        n_particles=N_PARTICLES,
        dim=len(var_defs),
        lb=lb,
        ub=ub,
        max_iter=N_ITER,
        w=PSO_W,
        c1=PSO_C1,
        c2=PSO_C2,
        autosave_every_iters=PSO_AUTOSAVE_EVERY,
        autosave_path=PSO_AUTOSAVE_PATH,
        eval_delay=PSO_EVAL_DELAY,
    )

    res = pso.optimize()
    x_best = res["gbest"]
    j_best = float(res["gbest_val"])

    if x_best is None:
        raise RuntimeError("PSO nie zwróciło najlepszego rozwiązania.")

    q_pv_best, p_storage_best = unpack_particle(x_best, pv_specs)

    apply_pv_q_setpoints(pv_specs, q_pv_best)
    final_storage_rows = apply_storage_to_candidate(app, critical_candidate, storage_rows_from_p_values(p_storage_best))
    converged_final, rc_final = try_run_loadflow(ldf)

    if not converged_final:
        raise RuntimeError(f"Końcowy najlepszy wariant PSO nie jest zbieżny, rc={rc_final}")

    pso_config = build_pso_config_table()
    best_solution = build_best_solution_table(x_best, var_defs)

    result_tables = collect_results_with_pso(
        app=app,
        tr=tr,
        setup_data=setup_data,
        storage_rows=final_storage_rows,
        pso_config=pso_config,
        pso_convergence=PSO_CONVERGENCE_ROWS,
        pso_particle_evals=PSO_PARTICLE_EVAL_ROWS,
        best_solution=best_solution,
    )
    export_to_excel(result_tables, OUT_FILE)

    ind = result_tables["indicators"][0]
    log_line(
        "Zakończono. "
        f"F1_best={j_best:.6f}, "
        f"F2_report={ind['F2']:.6f}, "
        f"Umax={ind['Umax_pu']:.4f} pu, "
        f"Umin={ind['Umin_pu']:.4f} pu, "
        f"P_export={ind['P_export_total_kW']:.2f} kW, "
        f"P_import={ind['P_import_total_kW']:.2f} kW, "
        f"Fcelu_2={ind['Fcelu_2_A']:.4f} A, "
        f"Fcelu_3={ind['Fcelu_3_weighted']:.6f}"
    )
    log_line(f"Zapisano wyniki do: {OUT_FILE}")


if __name__ == "__main__":
    run_pso_global_critical()